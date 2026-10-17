"""第九阶段对象模板的实际行为验收；所有自动API均模拟，无真实论文。"""
from pathlib import Path
from copy import deepcopy
from tempfile import TemporaryDirectory
from unittest.mock import patch
import unittest,json,hashlib
from lxml import etree as E
from core import check_document,fix_document,load_rules,load_custom_rules,effective_rules
from core.templates import FIELDS,ALL_FIELDS,EXT_FIELDS,defaults,validate_value,validate_draft,validate_template,make_template,upgrade_template,object_review_notes
from core.norms import selected_segments
from core.package import Package,q,NS,DemoError,write_new
from core.engine import _ensure
from core.extended import formatting_snapshot,snapshots_equal
from web.session import WebSession
from ai.template_spec import parse_template_requirements
from ai.short_spec import Settings

ROOT=Path(__file__).resolve().parents[1]
PAPER=ROOT/'samples/extended-correct.docx'
TYPES=json.loads(PAPER.with_suffix('.types.json').read_text())
B={'paper':'A4 portrait','margin_top':30,'margin_bottom':20,'margin_left':35,'margin_right':30,
   'font_cn':'楷体','font_western':'Arial','font_size':14,'line_spacing':2,'first_line_indent':1,
   'heading1_font_cn':'宋体','heading1_font_size':20,'heading1_alignment':'left','heading1_line_spacing':1.25,
   'heading2_font_cn':'仿宋','heading2_font_size':17,'heading2_alignment':'right','heading2_line_spacing':1.5,
   'heading3_font_cn':'楷体','heading3_font_size':13.5,'heading3_alignment':'both','heading3_line_spacing':2,
   'figure_caption_font_cn':'宋体','figure_caption_font_size':9,'table_caption_font_cn':'黑体','table_caption_font_size':12}


def draft_for(segments,values=B):
    # Independent expected numeric values; quotes map to named object, not engine output.
    out={n:None for n in ALL_FIELDS}
    titles={'heading1':'一级标题','heading2':'二级标题','heading3':'三级标题','figure_caption':'图题','table_caption':'表题'}
    for n,(rid,label,unit) in ALL_FIELDS.items():
        if n not in values:continue
        prefix=next((v for k,v in titles.items() if n.startswith(k+'_')),None)
        if prefix:line=next(s for s in segments if s['text'].startswith(prefix))
        elif n=='paper':line=next(s for s in segments if s['text'].startswith('纸张'))
        elif n.startswith('margin_'):
            side={'margin_top':'上','margin_bottom':'下','margin_left':'左','margin_right':'右'}[n]
            line=next(s for s in segments if s['text'].startswith(side+'边距') or s['text'].startswith('上下左右'))
        else:
            part={'font_cn':'正文中文','font_western':'正文西文','font_size':'正文普通','line_spacing':'正文采用','first_line_indent':'正文首行'}[n]
            line=next(s for s in segments if s['text'].startswith(part))
        out[n]={'value':values[n],'unit':unit,'quote':line['text'],'location':line['location'],'support':'supported'}
    out['unsupported']=[{'quote':s['text'],'location':s['location'],'reason':'printing' if '打印' in s['text'] else 'unsupported'} for s in segments if '打印' in s['text'] or '题注编号' in s['text']]
    return out


def mocked(segments,values=B):
    d=draft_for(segments,values)
    return parse_template_requirements(segments,extended=True,config=Settings('mock-only-key'),transport=lambda req,t:json.dumps({'choices':[{'finish_reason':'stop','message':{'content':json.dumps(d,ensure_ascii=False)}}]}).encode())


def raw_value(pkg,p,r,tag,attr):
    # Independent literal XML read for this known synthetic fixture, not Resolver output.
    layers=[pkg.styles.find('w:docDefaults/w:rPrDefault/w:rPr',NS)]
    styles={s.get(q('styleId')):s for s in pkg.styles.findall(q('style'))}
    for ref in (p.find('w:pPr/w:pStyle',NS),r.find('w:rPr/w:rStyle',NS)):
        sid=ref.get(q('val')) if ref is not None else None;chain=[]
        while sid:
            style=styles[sid];chain.insert(0,style.find(q('rPr')))
            base=style.find(q('basedOn'));sid=base.get(q('val')) if base is not None else None
        layers.extend(chain)
    layers.append(r.find(q('rPr')));value=None
    for layer in layers:
        if layer is None:continue
        e=layer.find(q(tag))
        if e is not None and e.get(q(attr)) is not None:value=e.get(q(attr))
    return value


class ObjectTemplateAcceptance(unittest.TestCase):
    def setUp(self):
        self.temp=TemporaryDirectory();self.dir=Path(self.temp.name);self.base=load_rules();self.model=WebSession(self.dir)
    def tearDown(self):self.temp.cleanup()
    def applied(self,values=B):
        m=self.model;m.set_mode('extended');m.set_norm_text((ROOT/'samples/norm9-B.txt').read_text());m.confirm_norm_text()
        self.assertTrue(m.parse_template(parser=lambda s:mocked(s,values)),m.template_error)
        self.assertTrue(m.apply_template(),m.template_error);self.assertTrue(m.load(PAPER.read_bytes(),'合成.docx'),m.error)
        m.set_types(TYPES['paragraph_types']);m.confirm();self.assertTrue(m.check(),m.error)
        return m
    def blank(self,text):
        d={n:None for n in ALL_FIELDS};d['unsupported']=[];return d,selected_segments(text)
    def field(self,d,name,value,text,loc='摘录:L1'):
        d[name]={'value':value,'unit':ALL_FIELDS[name][2],'quote':text,'location':loc,'support':'supported'}

    def test_A_and_B_all_objects_and_default_comparison(self):
        s=selected_segments((ROOT/'samples/norm9-A.txt').read_text());a=mocked(s,defaults(self.base,extended=True));self.assertEqual(a['receipt']['kind'],'mock_api');self.assertEqual(a['draft']['figure_caption_font_size']['value'],10.5)
        m=self.applied();self.assertEqual(m.active_template['structure_version'],3)
        self.assertEqual(m.active_template['final_values'],B)
        fixed=check_document(PAPER,mode='extended',classification_manifest=TYPES)
        self.assertEqual(fixed['summary']['problem_count'],0);self.assertGreater(m.report['summary']['problem_count'],20)
        self.assertEqual(m.report['summary']['rule_count'],12)
        self.assertEqual(m.report['rules']['template_structure_version'],3)

    def test_distinct_captions_expectations_and_sources(self):
        m=self.applied();figure=[f for f in m.report['findings'] if f.get('object_type')=='figure_caption'];table=[f for f in m.report['findings'] if f.get('object_type')=='table_caption']
        for group,size,font in ((figure,9,'宋体'),(table,12,'黑体')):
            self.assertTrue(group)
            self.assertTrue(all(f['expected']['pt']==size for f in group if f['locator']['property']=='size'))
            self.assertTrue(all(f['expected']['font']==font for f in group if f['locator']['property']=='font_eastAsia'))
            for f in group:
                if f['locator']['property'] in ('size','font_eastAsia'):
                    self.assertIsNone(f['source']['physical_page']);self.assertIn('图题' if group is figure else '表题',f['source']['quote'])

    def test_all_selected_actual_xml_and_independent_recheck(self):
        m=self.applied();raw=PAPER.read_bytes();old=Package(PAPER);m.set_selection([f['key'] for f in m.eligible()]);self.assertTrue(m.generate(),m.error)
        r=m.result;new=Package(r['output']);self.assertEqual(r['after']['summary']['problem_count'],0);self.assertEqual(r['after']['summary']['review_count'],0)
        independent=check_document(r['output'],mode='extended',classification_manifest=r['output_classification_manifest'],custom_rules_path=m.custom_rules_path)
        self.assertEqual(independent['summary']['problem_count'],0);self.assertEqual(PAPER.read_bytes(),raw)
        self.assertEqual(old.parts['word/numbering.xml'],new.parts['word/numbering.xml'])
        self.assertEqual(old.root.xpath('//w:t/text()',namespaces=NS),new.root.xpath('//w:t/text()',namespaces=NS))
        for n,t in TYPES['paragraph_types'].items():
            if t.startswith('heading') or t.endswith('caption'):
                p=new.paragraphs[int(n)-1];prefix=t
                for run in p.findall(q('r')):
                    if any('\u4e00'<=c<='\u9fff' for c in ''.join(run.xpath('.//w:t/text()',namespaces=NS))):
                        self.assertEqual(raw_value(new,p,run,'rFonts','eastAsia'),B[prefix+'_font_cn'])
                    self.assertEqual(float(raw_value(new,p,run,'sz','val'))/2,B[prefix+'_font_size'])
                if t.startswith('heading'):
                    self.assertEqual(p.find('w:pPr/w:jc',NS).get(q('val')),B[t+'_alignment']);self.assertEqual(int(p.find('w:pPr/w:spacing',NS).get(q('line'))),round(240*B[t+'_line_spacing']))
        self.assertTrue(r['preservation']['other_zip_member_bytes_equal'])

    def test_partial_only_heading_line_changes(self):
        m=self.applied();m.set_selection([f['key'] for f in m.eligible() if f['rule_id']=='R09' and f['locator']['property']=='line']);self.assertTrue(m.generate(),m.error)
        old,new=Package(PAPER),Package(m.result['output'])
        self.assertEqual(len(m.result['changes']),2)
        for a,b in zip(old.paragraphs,new.paragraphs):self.assertEqual([E.tostring(x) for x in a.findall(q('r'))],[E.tostring(x) for x in b.findall(q('r'))])
        self.assertGreater(m.result['after']['summary']['problem_count'],0)
        self.assertTrue(all(f['status']=='pass' for f in m.result['after']['findings'] if f['rule_id']=='R09' and f['locator'].get('property')=='line'))

    def test_caption_only_font_size_no_position_indent_number_or_latin_changes(self):
        m=self.applied();m.set_selection([f['key'] for f in m.eligible() if f['rule_id']=='R12']);self.assertTrue(m.generate(),m.error)
        a,b=Package(PAPER),Package(m.result['output'])
        for n,t in TYPES['paragraph_types'].items():
            x,y=a.paragraphs[int(n)-1],b.paragraphs[int(n)-1]
            if t.endswith('caption'):
                self.assertEqual(E.tostring(x.find(q('pPr'))),E.tostring(y.find(q('pPr'))))
                for r1,r2 in zip(x.findall(q('r')),y.findall(q('r'))):
                    self.assertEqual(raw_value(a,x,r1,'rFonts','ascii'),raw_value(b,y,r2,'rFonts','ascii'))
            else:self.assertEqual(E.tostring(x),E.tostring(y))
        self.assertEqual(a.parts['word/numbering.xml'],b.parts['word/numbering.xml'])

    def test_shared_explicit_all_levels_and_generic_title_ambiguous(self):
        d,s=self.blank('各级标题均黑体。')
        for i in (1,2,3):self.field(d,f'heading{i}_font_cn','黑体','各级标题均黑体。')
        validate_draft(d,s)
        for text in ('标题小二号。','标题1.1小二号。','1.1 黑体小二号。'):
            d,s=self.blank(text);self.field(d,'heading1_font_size',18,text)
            with self.subTest(text=text),self.assertRaises(DemoError):validate_draft(d,s)
        self.assertTrue(any(n['reason']=='ambiguous' for n in object_review_notes(selected_segments('标题小二号。'))))

    def test_body_headings_captions_never_cross_objects(self):
        cases=[('font_size',16,'正文小四号，标题三号。'),('heading1_font_size',12,'正文小四号。'),('figure_caption_font_size',12,'表题小四号。'),('heading2_font_cn','黑体','一级标题黑体。'),('heading1_font_size',14,'摘要四号。')]
        for n,v,text in cases:
            d,s=self.blank(text);self.field(d,n,v,text)
            with self.subTest(n=n),self.assertRaises(DemoError):validate_draft(d,s)
        d,s=self.blank('图题五号、表题小四号。');self.field(d,'figure_caption_font_size',10.5,'图题五号');self.field(d,'table_caption_font_size',12,'表题小四号。');validate_draft(d,s)
        d,s=self.blank('标题1中文字体为黑体，字号18pt。');self.field(d,'heading1_font_size',18,s[0]['text']);validate_draft(d,s)

    def test_same_object_conflicts_not_last_wins(self):
        for n,v,text in [('heading1_font_size',20,'一级标题18pt。\n一级标题20pt。'),('table_caption_font_cn','宋体','表题黑体。\n表题宋体。'),('heading2_alignment','center','二级标题居左。\n二级标题居中。')]:
            d,s=self.blank(text);self.field(d,n,v,s[-1]['text'],'摘录:L2')
            with self.subTest(n=n),self.assertRaises(DemoError):validate_draft(d,s)
            self.assertTrue(any(note['reason']=='conflict' for note in object_review_notes(s)))
            d[n]=None;validate_draft(d,s)

    def test_missing_requires_explicit_default_or_manual(self):
        m=self.model;m.set_mode('extended');m.set_norm_text('图题五号。');m.confirm_norm_text();d,s=self.blank(m.norm_text);self.field(d,'figure_caption_font_size',10.5,m.norm_text)
        self.assertTrue(m.parse_template(parser=lambda _: {'draft':d,'receipt':{'kind':'mock_api','model':'deepseek-flash','timestamp_utc':'2026-10-14T00:00:00+00:00'}}));self.assertFalse(m.apply_template());self.assertIsNone(m.template_draft['heading1_font_size'])
        m.choose_missing_defaults();self.assertTrue(m.apply_template(),m.template_error)
        self.assertEqual(m.active_template['field_sources']['heading1_font_size']['kind'],'default_inherited');self.assertEqual(m.active_template['field_sources']['figure_caption_font_size']['kind'],'ai_extracted')

    def test_unsupported_line_font_numbering_and_size_not_executable(self):
        cases=[('heading1_line_spacing',2,'一级标题固定20磅行距。'),('heading1_font_size',20,'一级标题固定20磅行距。'),('heading1_font_cn','黑体','一级标题方正黑体。'),('heading1_font_size',18,'一级标题建议小二号。')]
        for n,v,text in cases:
            d,s=self.blank(text);self.field(d,n,v,text)
            with self.subTest(n=n),self.assertRaises(DemoError):validate_draft(d,s)
        d,s=self.blank('一级标题编号采用1.1。');d['numbering']={'value':1}
        with self.assertRaises(DemoError):validate_draft(d,s)

    def test_numeric_types_ranges_halfpoints_and_enums(self):
        for n in EXT_FIELDS:
            values=[True,None,[],{},float('nan'),float('inf')]
            if n.endswith('font_size'):values += [7.5,36.5 if n.startswith('heading') else 24.5,13.3,'18']
            elif n.endswith('font_cn'):values+=['微软雅黑','黑体加粗']
            elif n.endswith('alignment'):values+=['justify','居中',1]
            else:values += [1.13,5.05,'2']
            for v in values:
                with self.subTest(n=n,v=str(v)),self.assertRaises(DemoError):validate_value(n,v)
        for v in (8,24.5,36):self.assertEqual(validate_value('heading1_font_size',v),v)

    def test_units_quotes_locations_unknown_fields_and_self_confirmation(self):
        m=self.applied();data=deepcopy(m.template_draft)
        for kind in ('unit','quote','location','unknown','confirm','confidence','path'):
            d=deepcopy(data)
            if kind=='unit':d['heading1_font_size']['unit']='mm'
            if kind=='quote':d['heading1_font_size']['quote']='一级标题40pt。'
            if kind=='location':d['heading1_font_size']['location']='摘录:L999'
            if kind=='unknown':d['heading4_font_size']=18
            if kind=='confirm':d['confirmed']=True
            if kind=='confidence':d['confidence']=.95
            if kind=='path':d['unsupported'].append({'location':'摘录:L1','quote':'/Users/private/file','reason':'unsupported'})
            with self.subTest(kind=kind),self.assertRaises(DemoError):validate_draft(d,m.norm_segments)

    def test_user_edit_changes_source_without_immediate_effect(self):
        m=self.applied();active=deepcopy(m.active_template);r=deepcopy(m.report)
        m.edit_template('heading1_font_size','ai',21,touched=True)
        self.assertEqual(m.active_template,active);self.assertEqual(m.report,r)
        p=m.proposed_template();self.assertEqual(p['field_sources']['heading1_font_size']['kind'],'user_confirmed');self.assertIn('20pt',p['field_sources']['heading1_font_size']['quote'])
        self.assertTrue(m.apply_template());self.assertIsNone(m.report);self.assertIsNone(m.classification_confirmed);self.assertFalse(m.downloads())

    def test_v3_roundtrip_same_checking_without_api(self):
        m=self.applied();raw=m.template_download();second=WebSession(self.dir/'second');second.set_mode('extended')
        with patch('ai.short_spec.http_post',side_effect=AssertionError('No real API')):
            self.assertTrue(second.import_template(raw));self.assertIsNone(second.active_template);self.assertTrue(second.apply_template());second.load(PAPER.read_bytes(),'合成.docx');second.set_types(TYPES['paragraph_types']);second.confirm();second.check()
        self.assertEqual(second.active_template['final_values'],B)
        self.assertEqual([(f['key'],f['expected'],f['status']) for f in m.report['findings']],[(f['key'],f['expected'],f['status']) for f in second.report['findings']])
        self.assertNotEqual(second.task_dir,m.task_dir);self.assertNotEqual(second.norm_dir,m.norm_dir)
        self.assertNotIn(str(ROOT).encode(),raw);self.assertNotIn(b'API_KEY',raw)

    def test_v2_kept_and_explicit_upgrade_fixed_extension(self):
        values=defaults(self.base);values['font_size']=14
        sources={n:{'kind':'default_inherited','quote':'','location':None} for n in FIELDS};sources['font_size']['kind']='user_confirmed'
        old=make_template('旧版课程演示要求',values,sources,[],self.base);self.assertEqual(old['structure_version'],2)
        m=self.model;m.set_mode('extended');self.assertTrue(m.import_template(json.dumps(old,ensure_ascii=False).encode()));self.assertEqual(m.template_import['structure_version'],2);self.assertTrue(m.apply_template());legacy=effective_rules(custom_rules_path=m.custom_rules_path,mode='extended')
        m.upgrade_template_draft();self.assertEqual(m.template_import['structure_version'],3)
        self.assertEqual({n:m.template_values[n] for n in FIELDS},values)
        self.assertTrue(all(m.template_import['field_sources'][n]['kind']=='legacy_fixed' for n in EXT_FIELDS))
        self.assertTrue(m.apply_template());up=effective_rules(custom_rules_path=m.custom_rules_path,mode='extended')
        for rid in ('R08','R09','R10'):self.assertEqual(up['rules'][rid]['expected'],legacy['rules'][rid]['expected'])
        self.assertEqual(up['rules']['R12']['expected']['figure_caption'],legacy['rules']['R12']['expected'])
        self.assertEqual(m.active_template['compatibility'],{'from_version':2,'method':'local_default_extension'})

    def test_invalid_v3_template_version_values_or_extra_execution_rejected(self):
        m=self.applied();data=m.active_template
        for kind in ('version','newkey','bool','range','source','extensionhash','compatibility','R11'):
            bad=deepcopy(data)
            if kind=='version':bad['structure_version']=4
            if kind=='newkey':bad['final_values']['title_before']=20
            if kind=='bool':bad['final_values']['heading1_font_size']=True
            if kind=='range':bad['final_values']['table_caption_font_size']=36
            if kind=='source':bad['field_sources']['heading1_font_size']['quote']='伪造20pt'
            if kind=='extensionhash':bad['extension_base_sha256']='0'*64
            if kind=='compatibility':bad['compatibility']={'from_version':2,'method':'AI confirmed'}
            if kind=='R11':bad['heading_style']='Anything'
            with self.subTest(kind=kind),self.assertRaises(DemoError):validate_template(bad,self.base)

    def test_automatic_numbering_R11_block_and_selected_title_size_safe(self):
        m=self.applied();pkg=Package(PAPER);p=pkg.paragraphs[6];pp=_ensure(p,'pPr');_ensure(pp,'pStyle').set(q('val'),'Normal');numpr=_ensure(pp,'numPr');_ensure(numpr,'numId').set(q('val'),'99');_ensure(numpr,'ilvl').set(q('val'),'0')
        num=E.Element(q('numbering'),nsmap={'w':NS['w']});ab=E.SubElement(num,q('abstractNum'),{q('abstractNumId'):'99'});lvl=E.SubElement(ab,q('lvl'),{q('ilvl'):'0'});E.SubElement(lvl,q('numFmt'),{q('val'):'decimal'});E.SubElement(lvl,q('lvlText'),{q('val'):'%1.'});inst=E.SubElement(num,q('num'),{q('numId'):'99'});E.SubElement(inst,q('abstractNumId'),{q('val'):'99'});pkg.parts['word/numbering.xml']=E.tostring(num)
        path=self.dir/'numbered.docx';write_new(path,pkg.bytes(True));manifest=deepcopy(TYPES);manifest['input_sha256']=hashlib.sha256(path.read_bytes()).hexdigest()
        before=check_document(path,mode='extended',classification_manifest=manifest,custom_rules_path=m.custom_rules_path);style=next(f for f in before['findings'] if f['key']=='R11:paragraph=7');self.assertEqual(style['status'],'problem');self.assertFalse(style['fixable'])
        out=fix_document(path,self.dir/'numbered-fixed.docx',None,mode='extended',classification_manifest=manifest,selected_keys=['R08:paragraph=7:run=1:property=size',style['key']],custom_rules_path=m.custom_rules_path)
        self.assertEqual(Package(out['output']).parts['word/numbering.xml'],pkg.parts['word/numbering.xml']);self.assertEqual(len(out['changes']),1);self.assertEqual(next(f for f in out['after']['findings'] if f['key']=='R08:paragraph=7:run=1:property=size')['actual'],{'pt':20})

    def test_R11_preserves_unselected_custom_effective_properties(self):
        m=self.applied();pkg=Package(PAPER);p=pkg.paragraphs[12]
        clone=deepcopy(pkg.styles.find("w:style[@w:styleId='Heading2']",NS));clone.set(q('styleId'),'ManualTitle');clone.set(q('customStyle'),'1');clone.find(q('name')).set(q('val'),'ManualTitle');pkg.styles.append(clone);pkg.parts['word/styles.xml']=E.tostring(pkg.styles)
        _ensure(_ensure(p,'pPr'),'pStyle').set(q('val'),'ManualTitle')
        path=self.dir/'fake-heading.docx';write_new(path,pkg.bytes(True));manifest=deepcopy(TYPES);manifest['input_sha256']=hashlib.sha256(path.read_bytes()).hexdigest();snapshot=formatting_snapshot(pkg,p)
        r=fix_document(path,self.dir/'styled.docx',None,mode='extended',classification_manifest=manifest,selected_keys=['R11:paragraph=13'],custom_rules_path=m.custom_rules_path);fixed=Package(r['output'])
        self.assertTrue(snapshots_equal(snapshot,formatting_snapshot(fixed,fixed.paragraphs[12])));self.assertEqual(r['changes'][0]['rule_id'],'R11')
        self.assertTrue(any(f['status']=='problem' and f['rule_id']=='R09' for f in r['after']['findings']))

    def test_state_switch_preview_download_rules_and_sessions(self):
        m=self.applied();m.set_selection([f['key'] for f in m.eligible()]);self.assertTrue(m.generate());raw=Path(m.result['output']).read_bytes();self.assertTrue(m.downloads())
        m.edit_template('table_caption_font_size','manual',13,touched=True);self.assertTrue(m.apply_template());self.assertFalse(m.downloads());self.assertIsNone(m.report);self.assertIsNone(m.classification_confirmed)
        second=WebSession(self.dir/'second');self.assertIsNone(second.active_template);self.assertIsNone(second.template_draft);self.assertIsNone(second.input_path);self.assertNotEqual(second.rules_hash,m.rules_hash)
        m.use_default_rules();m.confirm();m.check();self.assertEqual(m.report['summary']['problem_count'],0)

    def test_request_one_selected_excerpt_no_paper_no_key_log(self):
        s=selected_segments((ROOT/'samples/norm9-B.txt').read_text());calls=[]
        def transport(req,timeout):calls.append((json.loads(req.data),timeout));return json.dumps({'choices':[{'finish_reason':'stop','message':{'content':json.dumps(draft_for(s),ensure_ascii=False)}}]}).encode()
        r=parse_template_requirements(s,extended=True,config=Settings('DO_NOT_LOG_THIS'),transport=transport)
        self.assertEqual(len(calls),1);self.assertEqual(calls[0][1],30);self.assertNotIn(str(PAPER),json.dumps(calls));self.assertNotIn('DO_NOT_LOG_THIS',json.dumps(r));self.assertNotIn('paragraph_types',json.dumps(calls));self.assertEqual(calls[0][0]['response_format'],{'type':'json_object'})

    def test_failed_api_preserves_v3_active_and_default(self):
        m=self.applied();active=deepcopy(m.active_template);m.set_norm_text('一级标题18pt。');m.confirm_norm_text()
        calls=[]
        def fail(*args,**kwargs):calls.append(1);raise TimeoutError()
        with patch('ai.short_spec.http_post',side_effect=fail),patch('ai.template_spec.settings',return_value=Settings('test-only')):self.assertFalse(m.parse_template())
        self.assertEqual(len(calls),1);self.assertEqual(m.active_template,active);m.confirm();self.assertTrue(m.check());m.use_default_rules();m.confirm();self.assertTrue(m.check());self.assertEqual(m.report['summary']['problem_count'],0)

    def test_basic_mode_does_not_execute_custom_title_or_caption(self):
        m=self.applied();m.set_mode('basic');m.set_body([int(n) for n,t in TYPES['paragraph_types'].items() if t=='body']);m.confirm();m.check();self.assertEqual(m.report['summary']['rule_count'],7);self.assertFalse(any(f['rule_id'] in ('R08','R09','R10','R11','R12') for f in m.report['findings']))

    def test_DOCX_object_clauses_keep_local_locations_and_apply(self):
        from docx import Document
        from io import BytesIO
        d=Document()
        for line in (ROOT/'samples/norm9-B.txt').read_text().splitlines():d.add_paragraph(line)
        data=BytesIO();d.save(data)
        m=self.model;m.set_mode('extended')
        self.assertTrue(m.load_norm(data.getvalue(),'格式要求.docx'))
        m.select_norm_segments([s['location'] for s in m.norm_inventory['segments']]);m.confirm_norm_text()
        self.assertTrue(m.parse_template(parser=mocked),m.template_error);self.assertTrue(m.apply_template(),m.template_error)
        self.assertEqual(m.active_template['final_values'],B)
        self.assertEqual(m.active_template['field_sources']['heading1_font_size']['location'],'DOCX:P12')
        self.assertEqual(m.active_template['field_sources']['table_caption_font_size']['location'],'DOCX:P16')

    def test_blank_lines_and_trailing_newline_do_not_drop_requirement_text(self):
        m=self.model;m.set_norm_text('\n一级标题小二号。\n\n表题小四号。\n');m.confirm_norm_text()
        self.assertEqual(m.norm_segments,[{'location':'摘录:L2','text':'一级标题小二号。'},{'location':'摘录:L4','text':'表题小四号。'}])
        from core.norms import marked_text
        self.assertEqual(m.norm_confirmed_hash,hashlib.sha256(marked_text(m.norm_segments).encode()).hexdigest())

    def test_legacy_body_API_cannot_expand_its_requested_schema(self):
        s=selected_segments((ROOT/'samples/norm9-B.txt').read_text());calls=[]
        def transport(req,timeout):
            calls.append(1)
            return json.dumps({'choices':[{'finish_reason':'stop','message':{'content':json.dumps(draft_for(s),ensure_ascii=False)}}]}).encode()
        with self.assertRaises(DemoError):parse_template_requirements(s,config=Settings('mock-only'),transport=transport)
        self.assertEqual(calls,[1])

if __name__=='__main__':unittest.main()
