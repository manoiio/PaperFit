"""Author-defined spans + raw XML assertions, separate from the new verifier."""
from pathlib import Path
from tempfile import TemporaryDirectory
from copy import deepcopy
from unittest.mock import patch
import unittest,json
from docx import Document
from lxml import etree as E
from core.package import Package,q,NS,DemoError
from core import dynamic
from scripts.build_smart_dynamic_fixture import build
from web.session import WebSession
from web.dynamic_session import preview,adopt,valid,cancel

class DynamicAcceptance(unittest.TestCase):
 def setUp(self):
  self.tmp=TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.root=Path(self.tmp.name);self.path,self.m=build(self.root/'fixture')
 def mutate(self,fn):
  p=Package(self.path);fn(p);self.path.write_bytes(p.bytes(True));self.m['input_sha256']=Package(self.path).hash
 def session(self):
  s=WebSession(self.root/'session');s.set_mode('academic');s.load(self.path.read_bytes(),'动态合成.docx');s.set_types(self.m['paragraph_types']);s.confirm_and_check();return s
 def test_explicit_complete_and_no_word_or_automatic_flags(self):
  with self.assertRaises(DemoError):dynamic.propose(self.path,self.m)
  before=self.path.read_bytes();plan=dynamic.propose(self.path,self.m,complete=True);self.assertEqual(before,self.path.read_bytes());self.assertEqual(len(plan['edits']),5)
  self.assertFalse(plan['word']['available']);self.assertEqual(plan['word']['status'],'not_run')
  with patch('core.dynamic.sys.platform','win32'):self.assertFalse(dynamic.word_capability()['available'])
 def test_cross_run_exact_digits_only_raw_xml_independent(self):
  old=self.path.read_bytes();p=dynamic.propose(self.path,self.m,complete=True);r=dynamic.rehearse(self.path,self.root/'out.docx',self.m,p);a,b=Package(self.path),Package(r['output'])
  expected=json.loads((self.root/'fixture/independent-errors.json').read_text());self.assertEqual(dynamic.text(b.paragraphs[3]),expected['expected_body'])
  self.assertEqual([t.text for t in b.paragraphs[3].xpath('./w:r/w:t',namespaces=NS)],['依据图1-','1与图2-','1，可比较合成结果。再次参见图1-1。'])
  # Hand-written affected text-node slots, not parser-produced spans.
  targets={4: ['依据图1-','1与图2-','1，可比较合成结果。再次参见图1-1。'],5:['图1-1 示例曲线'],13:['图2-1 示例曲线']}
  root=E.fromstring(a.parts['word/document.xml']);ps=root.find(q('body')).findall(q('p'))
  for n,values in targets.items():
   ts=ps[n-1].xpath('./w:r/w:t',namespaces=NS);self.assertEqual(len(ts),len(values))
   for t,value in zip(ts,values):t.text=value
  self.assertEqual(E.tostring(root,method='c14n'),E.tostring(b.root,method='c14n'))
  self.assertEqual({k:v for k,v in a.parts.items() if k!='word/document.xml'},{k:v for k,v in b.parts.items() if k!='word/document.xml'});self.assertEqual(old,self.path.read_bytes())
  again=dynamic.propose(r['output'],r['output_classification_manifest'],complete=True);self.assertFalse(again['edits'])
 def test_chain_replacement_and_number_group_uniqueness(self):
  d=Document();d.add_paragraph('第1章 合成',style='Heading 1');d.add_paragraph('参见图1-2以及图1-3，重复图1-2。');d.add_paragraph('图1-2 甲');d.add_paragraph('图1-3 乙');p=self.root/'chain.docx';d.save(p)
  m={'input_sha256':Package(p).hash,'paragraph_types':{'1':'heading1','2':'body','3':'figure_caption','4':'figure_caption'},'confirmed':True,'confirmation':'作者手工确认'}
  plan=dynamic.propose(p,m,complete=True);out=dynamic.rehearse(p,self.root/'chain-out.docx',m,plan)
  self.assertEqual(dynamic.text(Package(out['output']).paragraphs[1]),'参见图1-1以及图1-2，重复图1-1。')
  # Duplicate old number blocks the entire linked chapter/type group.
  self.mutate(lambda pkg:setattr(pkg.paragraphs[12].find('w:r/w:t',NS),'text','图1-2 重复编号'))
  duplicate=dynamic.propose(self.path,self.m,complete=True);self.assertFalse(duplicate['edits']);self.assertTrue(duplicate['blocked_groups'])
 def test_unknown_chapter_complex_caption_and_reference_block_groups(self):
  self.mutate(lambda p:setattr(p.paragraphs[2].find('w:r/w:t',NS),'text','无可靠章号'))
  p=dynamic.propose(self.path,self.m,complete=True);self.assertTrue(any('章号' in str(g) for g in p['blocked_groups']));self.assertFalse(any(e['paragraph']==5 for e in p['edits']))
  # A field/complex wrapper in an associated body reference prevents partial caption-only updates.
  self.mutate(lambda pkg:E.SubElement(pkg.paragraphs[3],q('fldSimple')).set(q('instr'),'REF Unknown'))
  p=dynamic.propose(self.path,self.m,complete=True);self.assertFalse(p['edits']);self.assertTrue(p['blocked_groups'])
 def test_unselected_format_tamper_refused_and_original_guard_unchanged(self):
  p=dynamic.propose(self.path,self.m,complete=True);r=dynamic.rehearse(self.path,self.root/'out.docx',self.m,p)
  obj=Package(r['output']);obj.paragraphs[3].find('w:r/w:rPr/w:color',NS).set(q('val'),'FF0000');Path(r['output']).write_bytes(obj.bytes(True))
  with self.assertRaises(DemoError):dynamic.validate_output(self.path,r['output'],p,self.m)
  from core.engine import _preservation
  self.assertNotEqual(_preservation(Package(self.path),[]),_preservation(Package(r['output']),[]))
 def test_adoption_idempotency_archive_failure_and_session_isolation(self):
  s=self.session();state={};preview(s,state,complete=True);self.assertFalse(s.history);self.assertIsNone(s.result)
  with patch.object(WebSession,'_archive_round',side_effect=OSError('synthetic IO failure')):
   with self.assertRaises(OSError):adopt(s,state)
  self.assertFalse(s.history);self.assertEqual(s.generation_count,0);self.assertIsNone(s.result)
  self.assertTrue(adopt(s,state));self.assertTrue(adopt(s,state));self.assertEqual(len(s.history),1);self.assertEqual(s.current_round,0)
  raw=s.prepare_project();other=WebSession(self.root/'other');other.stage_project(raw);new=other.restore_project(acknowledged=True);self.assertEqual(len(new.history),1);self.assertIsNone(new.result)
  self.assertIsNone(new.rehearsal);self.assertFalse(hasattr(new,'dynamic_preview'))
 def test_external_preview_change_and_cancel_owned_only(self):
  s=self.session();state={};preview(s,state,complete=True);sentinel=s.task_dir/'other-task';sentinel.mkdir();(sentinel/'keep').write_text('keep')
  p=Path(state['result']['output']);p.write_bytes(p.read_bytes()+b'changed')
  with self.assertRaises(DemoError):adopt(s,state)
  self.assertFalse(s.history);cancel(s,state);self.assertTrue((sentinel/'keep').is_file());self.assertFalse(state)
 def test_format_output_is_dynamic_parent_and_no_auto_input_switch(self):
  # Ordinary edit first, then the independent numeric transaction, each archived.
  self.mutate(lambda pkg:E.SubElement(pkg.paragraphs[3].findall(q('r'))[0].find(q('rPr')),q('sz')).set(q('val'),'28'))
  s=self.session();key=next(f['key'] for f in s.eligible() if f['rule_id']=='R05' and f['locator'].get('paragraph')==4);s.set_selection([key]);self.assertTrue(s.rehearse());self.assertTrue(s.adopt_rehearsal(s.rehearsal['id'],acknowledged=True))
  intermediate=s.result['after']['input_sha256'];state={};rr=preview(s,state,complete=True);self.assertEqual(rr['before']['input_sha256'],intermediate);adopt(s,state)
  self.assertEqual(s.result['round']['input_round'],1);self.assertEqual(s.generation_count,2);self.assertEqual(s.current_round,0);self.assertEqual(Package(s.history[0]['result']['output']).hash,intermediate)

 def test_readback_rule_binding_change_refuses_and_removes_preview(self):
  from core import check_document
  calls=[]
  def changed(*args,**kwargs):
   r=check_document(*args,**kwargs);calls.append(1)
   if len(calls)==2:r['rules']['sha256']='changed during readback'
   return r
  p=dynamic.propose(self.path,self.m,complete=True);out=self.root/'refused.docx';original=self.path.read_bytes()
  with patch('core.check_document',side_effect=changed):
   with self.assertRaises(DemoError):dynamic.rehearse(self.path,out,self.m,p)
  self.assertFalse(out.exists());self.assertEqual(original,self.path.read_bytes())
