"""L02/L03属性行为与保护；模拟预览不作视觉证据。全部是合成内容。"""
from pathlib import Path
from tempfile import TemporaryDirectory
from copy import deepcopy
import json,unittest
from lxml import etree as E
from docx import Document
from docx.shared import Pt
from core.layout import suggest_pagination_aids,apply_pagination_layout,_flag
from core.package import Package,DemoError,q,NS,sha256,write_new
from core.engine import _ensure
from core.classification import suggest_paragraph_types
from web.session import WebSession

class PaginationAcceptance(unittest.TestCase):
 def setUp(self):
  self.tmp=TemporaryDirectory();self.root=Path(self.tmp.name)
  d=Document();d.add_paragraph('软件演示样本');h=d.add_paragraph('1 示例标题',style='Heading 1');h.paragraph_format.keep_with_next=False;h.paragraph_format.keep_together=False
  d.add_paragraph('这是软件演示正文。内容足够清楚，仅用于属性保护验证。'*4)
  d.add_paragraph('表1 合成数据');t=d.add_table(rows=4,cols=2);t.cell(0,0).merge(t.cell(0,1)).text='合并单元格'
  for i in range(1,4):t.cell(i,0).text=f'项目{i}';t.cell(i,1).text='数据'
  d.add_paragraph('尾部内容。');self.p=self.root/'in.docx';d.save(self.p)
  self.m={'input_sha256':sha256(self.p.read_bytes()),'paragraph_types':{'1':'other','2':'heading1','3':'body','4':'table_caption','5':'body'},'confirmed':True,'confirmation':'人工合成样本确认'}
 def tearDown(self):self.tmp.cleanup()
 def proposals(self,p=None,m=None):return suggest_pagination_aids(p or self.p,m or self.m)['pairs']
 def edit(self,fn):
  pkg=Package(self.p);fn(pkg);p=self.root/'edited.docx';write_new(p,pkg.bytes(True));m=deepcopy(self.m);m['input_sha256']=sha256(p.read_bytes());return p,m
 def apply(self,keys,p=None,m=None,name='out.docx'):
  return apply_pagination_layout(p or self.p,self.root/name,classification_manifest=m or self.m,confirmed_keys=keys)
 def model(self):
  s=WebSession(self.root/'web');s.set_mode('extended');s.load(self.p.read_bytes(),'合成.docx');s.set_types(self.m['paragraph_types']);s.confirm();s.check();return s
 def test_table_association_actual_xml_and_grid_overview(self):
  a=[x for x in self.proposals() if x['id']=='L02'][0];self.assertEqual((a['status'],a['table_position'],a['rows'],a['columns']),('available',1,4,2));self.assertEqual(a['table_body_object'],5)
 def test_intervening_paragraph_not_skipped(self):
  def fn(pkg):
   p=E.Element(q('p'));E.SubElement(E.SubElement(p,q('r')),q('t')).text='中间内容';pkg.paragraphs[3].addnext(p)
  p,m=self.edit(fn);m['paragraph_types']['6']='body';m['paragraph_types']['5']='other'
  a=[x for x in self.proposals(p,m) if x['id']=='L02'][0];self.assertEqual(a['status'],'review')
  with self.assertRaises(DemoError):self.apply(['L02:4:1'],p,m)
 def test_explicit_confirmation_and_hash_required(self):
  for keys in [[],['L02:4:99'],[True],['L03:2:3','L03:2:3']]:
   with self.assertRaises(DemoError):self.apply(keys)
  m=deepcopy(self.m);m['confirmed']=False
  with self.assertRaises(DemoError):self.apply(['L02:4:1'],m=m)
 def test_caption_only_changes_keepnext_table_and_merge_unchanged(self):
  old=Package(self.p);r=self.apply(['L02:4:1']);new=Package(r['output'])
  self.assertTrue(_flag(new,new.paragraphs[3],'keepNext'));self.assertEqual(len(r['changes']),1)
  self.assertEqual(E.tostring(old.root.find('w:body/w:tbl',NS)),E.tostring(new.root.find('w:body/w:tbl',NS)))
  for i in (0,1,2,4):self.assertEqual(E.tostring(old.paragraphs[i]),E.tostring(new.paragraphs[i]))
  self.assertTrue(all(r['preservation'][k] for k in ('original_sha256_unchanged','other_zip_member_bytes_equal','document_except_selected_format_attributes_equal')))
 def test_heading_selected_without_body_chain_or_style_change(self):
  old=Package(self.p);r=self.apply(['L03:2:3']);new=Package(r['output'])
  self.assertTrue(_flag(new,new.paragraphs[1],'keepNext'));self.assertTrue(_flag(new,new.paragraphs[1],'keepLines'))
  for i in (0,2,3,4):self.assertEqual(E.tostring(old.paragraphs[i]),E.tostring(new.paragraphs[i]))
  self.assertEqual(old.parts['word/styles.xml'],new.parts['word/styles.xml']);self.assertEqual(old.root.xpath('//w:t/text()',namespaces=NS),new.root.xpath('//w:t/text()',namespaces=NS))
 def test_empty_body_not_chained(self):
  p,m=self.edit(lambda pkg:pkg.paragraphs[2].remove(pkg.paragraphs[2].find(q('r'))))
  self.assertEqual([x for x in self.proposals(p,m) if x['id']=='L03'][0]['status'],'review')
 def test_heading_followed_by_table_or_unconfirmed_type_review(self):
  m=deepcopy(self.m);m['paragraph_types']['3']='pending';self.assertEqual([x for x in self.proposals(m=m) if x['id']=='L03'][0]['status'],'review')
 def test_forced_break_retained_and_not_executed(self):
  for tag in ('pageBreakBefore','sectPr'):
   with self.subTest(tag=tag):
    with TemporaryDirectory() as tmp:
     pkg=Package(self.p);_ensure(_ensure(pkg.paragraphs[1],'pPr'),tag).set(q('val'),'1');p=Path(tmp)/'in.docx';write_new(p,pkg.bytes(True));m=deepcopy(self.m);m['input_sha256']=sha256(p.read_bytes());data=p.read_bytes()
     with self.assertRaises(DemoError):self.apply(['L03:2:3'],p,m)
     self.assertEqual(p.read_bytes(),data)
 def test_next_body_keepnext_chain_review(self):
  p,m=self.edit(lambda pkg:_ensure(_ensure(pkg.paragraphs[2],'pPr'),'keepNext').set(q('val'),'1'))
  with self.assertRaises(DemoError):self.apply(['L03:2:3'],p,m)
 def test_nested_table_review(self):
  def fn(pkg):pkg.root.find('w:body/w:tbl/w:tr/w:tc',NS).append(E.Element(q('tbl')))
  p,m=self.edit(fn)
  with self.assertRaises(DemoError):self.apply(['L02:4:1'],p,m)
 def test_table_field_or_floating_review(self):
  for tag in ('fldChar','tblpPr'):
   with self.subTest(tag=tag),TemporaryDirectory() as tmp:
    pkg=Package(self.p);t=pkg.root.find('w:body/w:tbl',NS)
    (t.find(q('tblPr')) if tag=='tblpPr' else t.find('.//'+q('r'))).append(E.Element(q(tag)))
    p=Path(tmp)/'in.docx';write_new(p,pkg.bytes(True));m=deepcopy(self.m);m['input_sha256']=sha256(p.read_bytes())
    with self.assertRaises(DemoError):self.apply(['L02:4:1'],p,m)
 def test_long_table_and_title_no_visual_success_claim(self):
  p,m=self.edit(lambda pkg:pkg.root.find('w:body/w:tbl',NS).extend([deepcopy(pkg.root.find('w:body/w:tbl/w:tr',NS)) for _ in range(60)]))
  r=self.apply(['L02:4:1','L03:2:3'],p,m);self.assertEqual(r['layout']['visual_status'],'not_reviewed');self.assertIn('长表',r['layout']['limitation'])
  self.assertEqual(Package(r['output']).root.xpath('//w:trPr/w:cantSplit',namespaces=NS),[])
 def test_idempotent_and_preserve_original_previous(self):
  original=self.p.read_bytes();r=self.apply(['L02:4:1','L03:2:3']);first=Path(r['output']);data=first.read_bytes()
  r2=self.apply(['L02:4:1','L03:2:3'],first,r['output_classification_manifest'],'second.docx')
  self.assertEqual(r2['changes'],[]);self.assertEqual(Path(r2['output']).read_bytes(),data);self.assertEqual(first.read_bytes(),data);self.assertEqual(self.p.read_bytes(),original)
 def test_existing_output_or_original_refused(self):
  with self.assertRaises(DemoError):self.apply(['L02:4:1'],name='in.docx')
  self.apply(['L02:4:1'])
  with self.assertRaises(DemoError):self.apply(['L02:4:1'])
 def test_session_combined_rounds_and_bound_download(self):
  s=self.model();self.assertTrue(s.generate_layout('input',['L02:4:1']));first=Path(s.result['output']);data=first.read_bytes();self.assertTrue(s.generate_layout('input',['L02:4:1']));self.assertEqual(s.generation_count,1)
  self.assertTrue(s.generate_layout('round-1',['L03:2:3']));self.assertEqual(s.downloads()['round-input.docx'],data);self.assertEqual(first.read_bytes(),data);self.assertEqual(s.result['round']['input_round'],1)
 def test_sessions_selection_mode_input_invalidate(self):
  s=self.model();other=self.model();s.generate_layout('input',['L02:4:1']);self.assertIsNone(other.result);self.assertNotEqual(s.task_dir,other.task_dir)
  s.set_types({**self.m['paragraph_types'],'4':'other'});self.assertEqual(s.downloads(),{});self.assertEqual(s.preview_targets(),{})
  s.confirm();s.check();s.generate_layout('input',['L03:2:3']);s.set_mode('basic');self.assertEqual(s.downloads(),{});self.assertEqual(s.preview_targets(),{})
 def test_failed_selection_clears_current_result_retains_prior(self):
  s=self.model();s.generate_layout('input',['L02:4:1']);p=Path(s.result['output']);data=p.read_bytes();self.assertFalse(s.generate_layout('input',['bad']));self.assertIsNone(s.result);self.assertEqual(p.read_bytes(),data)
 def test_numbering_clue_ambiguous_not_automatic_confirmation(self):
  d=Document();p=d.add_paragraph('合成章节标题');pp=_ensure(p._p,'pPr');num=_ensure(pp,'numPr');_ensure(num,'ilvl').set(q('val'),'0');_ensure(num,'numId').set(q('val'),'1');d.add_paragraph('这是合成叙述正文，仅用于验证编号线索，不包含真实课程资料。'*5)
  nr=d.part.numbering_part.element;num=next(e for e in nr.findall(q('num')) if e.get(q('numId'))=='1');aid=num.find(q('abstractNumId')).get(q('val'));a=next(e for e in nr.findall(q('abstractNum')) if e.get(q('abstractNumId'))==aid);_ensure(a,'multiLevelType').set(q('val'),'multilevel');a.find('w:lvl/w:numFmt',NS).set(q('val'),'decimal');a.find('w:lvl/w:lvlText',NS).set(q('val'),'%1.')
  path=self.root/'num.docx';d.save(path);s=suggest_paragraph_types(path);self.assertEqual(s['paragraphs'][0]['suggested_type'],'heading1');self.assertTrue(s['paragraphs'][0]['ambiguous']);self.assertFalse(s['confirmed']);self.assertEqual(s['paragraphs'][1]['suggested_type'],'body')
 def test_normal_linebreak_supported_but_pagebreak_rejected(self):
  p,m=self.edit(lambda pkg:pkg.paragraphs[2].find(q('r')).append(E.Element(q('br'))));self.assertEqual([x for x in self.proposals(p,m) if x['id']=='L03'][0]['status'],'available')
  pkg=Package(p);pkg.paragraphs[2].find('w:r/w:br',NS).set(q('type'),'page');p2=self.root/'forced.docx';write_new(p2,pkg.bytes(True));m['input_sha256']=sha256(p2.read_bytes())
  with self.assertRaises(DemoError):self.apply(['L03:2:3'],p2,m)
 def test_long_heading_still_only_property_result(self):
  p,m=self.edit(lambda pkg:setattr(pkg.paragraphs[1].find('w:r/w:t',NS),'text','过长的合成标题'*300))
  r=self.apply(['L03:2:3'],p,m);self.assertEqual(r['layout']['visual_status'],'not_reviewed');self.assertIn('过长标题',r['layout']['limitation'])
 def test_table_first_row_chain_review(self):
  def fn(pkg):_ensure(_ensure(pkg.root.find('w:body/w:tbl/w:tr/w:tc/w:p',NS),'pPr'),'keepNext').set(q('val'),'1')
  p,m=self.edit(fn)
  with self.assertRaises(DemoError):self.apply(['L02:4:1'],p,m)
 def test_heading_revision_or_floating_review(self):
  for tag in ('pPrChange','framePr'):
   with self.subTest(tag=tag),TemporaryDirectory() as tmp:
    pkg=Package(self.p);_ensure(pkg.paragraphs[1],'pPr').append(E.Element(q(tag)));p=Path(tmp)/'in.docx';write_new(p,pkg.bytes(True));m=deepcopy(self.m);m['input_sha256']=sha256(p.read_bytes())
    with self.assertRaises(DemoError):self.apply(['L03:2:3'],p,m)
 def test_inherited_pagebreak_not_erased(self):
  def fn(pkg):
   _ensure(_ensure(next(s for s in pkg.styles.findall(q('style')) if s.get(q('styleId'))=='Heading1'),'pPr'),'pageBreakBefore').set(q('val'),'1')
   pkg.parts['word/styles.xml']=E.tostring(pkg.styles)
  p,m=self.edit(fn)
  self.assertTrue(_flag(Package(p),Package(p).paragraphs[1],'pageBreakBefore'))
  with self.assertRaises(DemoError):self.apply(['L03:2:3'],p,m)
 def test_unselected_media_sections_remain_with_mixed_aids(self):
  src=Path(__file__).resolve().parents[1]/'samples/extended-correct.docx';m=json.loads(src.with_suffix('.types.json').read_text());a=suggest_pagination_aids(src,m);keys=[x['key'] for x in a['pairs'] if x['status']=='available' and x['id'] in ('L02','L03')];self.assertTrue(keys)
  old=Package(src);r=self.apply(keys,src,m);new=Package(r['output']);self.assertEqual(old.parts['word/media/image1.png'],new.parts['word/media/image1.png']);self.assertEqual([E.tostring(x) for x in old.sections],[E.tostring(x) for x in new.sections]);self.assertEqual(old.root.xpath('//w:drawing',namespaces=NS)[0].getroottree().getpath(old.root.xpath('//w:drawing',namespaces=NS)[0]),new.root.xpath('//w:drawing',namespaces=NS)[0].getroottree().getpath(new.root.xpath('//w:drawing',namespaces=NS)[0]));self.assertTrue(r['preservation']['document_except_selected_format_attributes_equal'])
 def test_round_change_old_suggestion_cannot_bypass_new_body_type(self):
  s=self.model();s.generate_layout('input',['L02:4:1']);s.next_round();s.set_types({**s.paragraph_types,'3':'other'});s.confirm();s.check();self.assertFalse(s.generate_layout('input',['L03:2:3']));self.assertIsNone(s.result)
 def test_short_plain_or_bullet_list_not_heading(self):
  d=Document();d.add_paragraph('普通短句');d.add_paragraph('列表项目',style='List Bullet');d.add_paragraph('这是一段叙述。'*30);p=self.root/'list.docx';d.save(p);s=suggest_paragraph_types(p);self.assertFalse(any(x['suggested_type'].startswith('heading') for x in s['paragraphs']))
 def test_identical_round_bytes_still_bind_selected_input_round(self):
  s=self.model();s.generate_layout('input',['L02:4:1']);s.generate_layout('round-1',['L02:4:1']);second=Path(s.result['output']);data=second.read_bytes();self.assertEqual(s.result['round']['number'],2)
  self.assertTrue(s.generate_layout('round-2',['L02:4:1']));self.assertEqual(s.result['round']['input_round'],2);self.assertEqual(s.result['round']['number'],3);self.assertEqual(s.downloads()['round-input.docx'],data);self.assertEqual(second.read_bytes(),data)
 def test_conditional_table_style_pagination_review(self):
  def fn(pkg):
   t=pkg.root.find('w:body/w:tbl',NS);_ensure(_ensure(t,'tblPr'),'tblStyle').set(q('val'),'TestTable')
   style=E.SubElement(pkg.styles,q('style'),{q('styleId'):'TestTable',q('type'):'table'});conditional=E.SubElement(style,q('tblStylePr'),{q('type'):'firstRow'});_ensure(_ensure(conditional,'pPr'),'pageBreakBefore').set(q('val'),'1');pkg.parts['word/styles.xml']=E.tostring(pkg.styles)
  p,m=self.edit(fn)
  with self.assertRaises(DemoError):self.apply(['L02:4:1'],p,m)
 def test_legacy_fangsong_alias_preview_only_not_docx(self):
  from core.preview import _fonts,CONFIG
  from unittest.mock import patch
  import subprocess
  def fn(pkg):_ensure(_ensure(pkg.paragraphs[2].find(q('r')),'rPr'),'rFonts').set(q('eastAsia'),'仿宋_GB2312')
  p,m=self.edit(fn);data=p.read_bytes();job=self.root/'fonts';job.mkdir()
  def probe(cmd,**kw):return subprocess.CompletedProcess(cmd,0,'FangSong\nTimes New Roman\nCalibri\n' if 'fc-list' in cmd[0] else 'FangSong','')
  with patch('core.preview.subprocess.run',side_effect=probe):env,fonts,notes=_fonts(Package(p),job,CONFIG)
  self.assertEqual(CONFIG['font_aliases']['仿宋_GB2312'],'FangSong');self.assertIn('仿宋_GB2312',CONFIG['strong_font_aliases']);self.assertIn('<family>仿宋_GB2312</family><prefer><family>FangSong</family>',(job/'fonts.conf').read_text());self.assertTrue(any('仿宋_GB2312' in n and 'FangSong' in n for n in notes));self.assertEqual(p.read_bytes(),data)
