"""独立布局几何及保护验收：全部合成，渲染不在自动测试冒充视觉验收。"""
from pathlib import Path
from tempfile import TemporaryDirectory
from copy import deepcopy
import unittest,json
from io import BytesIO
from docx import Document
from docx.shared import Mm
from docx.enum.section import WD_SECTION_START
from lxml import etree as E
from core.widths import check_widths,apply_width_adaptation,WP,A,PIC,_scale,width_markdown
from core.package import Package,q,NS,DemoError,write_new,sha256
from core.engine import effective_rules,_ensure
from web.session import WebSession
ROOT=Path(__file__).resolve().parents[1]

class WidthAcceptance(unittest.TestCase):
 def setUp(self):
  self.tmp=TemporaryDirectory();self.d=Path(self.tmp.name)
  doc=Document();s=doc.sections[0];s.page_width=Mm(210);s.page_height=Mm(297);s.left_margin=Mm(35);s.right_margin=Mm(30)
  doc.add_paragraph('软件演示样本，仅用于宽度验证。')
  t=doc.add_table(rows=3,cols=2);t.style='Table Grid'
  for col in t.columns:col.width=Mm(80)
  for row in t.rows:
   for c in row.cells:c.width=Mm(80);c.text='合成表格文字，保留内容与自动行高。'
  p=Package(ROOT/'samples/extended-correct.docx');media=next(b for n,b in p.parts.items() if n.startswith('word/media/'))
  self.media=media;doc.add_picture(BytesIO(media),width=Mm(180));doc.add_paragraph('图1 软件演示图')
  self.p=self.d/'in.docx';doc.save(self.p)
  self.rules=effective_rules()['sha256']
 def tearDown(self):self.tmp.cleanup()
 def report(self,p=None):return check_widths(p or self.p,rules_sha256=self.rules)
 def edit(self,fn):
  pkg=Package(self.p);fn(pkg);p=self.d/('e-'+str(len(list(self.d.glob('e-*'))))+'.docx');write_new(p,pkg.bytes(True));return p
 def table(self,pkg):return pkg.root.find('w:body/w:tbl',NS)
 def fix(self,p=None,keys=('W01:table=1','W02:paragraph=2'),fixed=('W01:table=1',),name='out.docx',plan=None):
  p=p or self.p;return apply_width_adaptation(p,self.d/name,[],plan=plan or self.report(p),selected_keys=keys,confirmed_fixed_keys=fixed)
 def model(self):
  s=WebSession(self.d/'web');self.assertTrue(s.load(self.p.read_bytes(),'../合成.docx'));s.set_body([1]);s.confirm();self.assertTrue(s.check());return s
 def test_b_table_actual_risk_and_excess(self):
  r=check_widths(ROOT/'work/phase9/final/B-fixed.docx')['objects'][-1]
  self.assertEqual(r['status'],'risk');self.assertTrue(r['fixable']);self.assertAlmostEqual(r['excess_mm'],15.01069,places=4);self.assertTrue(r['plan']['autofit_to_fixed'])
 def test_simple_image_and_table_risk(self):
  r=self.report();self.assertEqual(r['summary'],{'risks':2,'reviews':0,'passes':0});self.assertAlmostEqual(r['objects'][1]['declared_width_mm'],180,places=4)
 def test_corresponding_sections_use_xml_ending_section(self):
  doc=Document(self.p);s=doc.add_section(WD_SECTION_START.NEW_PAGE);s.left_margin=Mm(20);s.right_margin=Mm(20)
  doc.add_picture(BytesIO(self.media),width=Mm(160));p=self.d/'sections.docx';doc.save(p)
  r=self.report(p)['objects'];self.assertEqual([x['section'] for x in r],[1,1,2]);self.assertEqual(r[-1]['status'],'pass');self.assertAlmostEqual(r[-1]['available_mm'],170,places=1)
 def test_margin_change_recomputes_risk(self):
  p=self.edit(lambda pkg:[sp.find(q('pgMar')).set(q('left'),'1000') for sp in pkg.sections]);self.assertEqual(self.report(p)['objects'][0]['status'],'pass')
 def test_normal_gutter_and_table_indent_subtracted(self):
  def fn(pkg):
   pkg.sections[0].find(q('pgMar')).set(q('gutter'),'100');_ensure(self.table(pkg).find(q('tblPr')),'tblInd').set(q('type'),'dxa');self.table(pkg).find('w:tblPr/w:tblInd',NS).set(q('w'),'200')
  r=self.report(self.edit(fn))['objects'][0];self.assertEqual(r['available_twips'],r['basis']['body_width_twips']-200);self.assertEqual(r['basis']['gutter_twips'],100)
 def test_fixed_width_coherent_and_conflict_review(self):
  def fn(pkg):
   t=self.table(pkg);w=t.find('w:tblPr/w:tblW',NS);w.set(q('type'),'dxa');w.set(q('w'),'9072');_ensure(t.find(q('tblPr')),'tblLayout').set(q('type'),'fixed')
  p=self.edit(fn);r=self.report(p)['objects'][0];self.assertEqual(r['status'],'risk');self.assertFalse(r['plan']['autofit_to_fixed'])
  p=self.edit(lambda pkg:self.table(pkg).find('w:tblPr/w:tblW',NS).set(q('type'),'dxa'))
  self.assertEqual(self.report(p)['objects'][0]['status'],'review')
 def test_percentage_coherent_readonly_and_conflict(self):
  def fn(pkg):
   t=self.table(pkg);w=t.find('w:tblPr/w:tblW',NS);w.set(q('type'),'pct');w.set(q('w'),'6000')
   total=round((11906-1984-1701)*1.2);cols=_scale([1,1],total)
   for e,c in zip(t.findall('w:tblGrid/w:gridCol',NS),cols):e.set(q('w'),str(c))
   for tr in t.findall(q('tr')):
    for tc,c in zip(tr.findall(q('tc')),cols):tc.find('w:tcPr/w:tcW',NS).set(q('w'),str(c))
  r=self.report(self.edit(fn))['objects'][0];self.assertEqual(r['status'],'risk');self.assertFalse(r['fixable']);self.assertIn('120%',r['reason'])
  p=self.edit(lambda pkg:self.table(pkg).find('w:tblPr/w:tblW',NS).set(q('type'),'pct'));self.assertEqual(self.report(p)['objects'][0]['status'],'review')
 def test_grid_cell_conflict_review_cannot_forced(self):
  p=self.edit(lambda pkg:self.table(pkg).find('w:tr/w:tc/w:tcPr/w:tcW',NS).set(q('w'),'1000'));self.assertEqual(self.report(p)['objects'][0]['status'],'review')
  with self.assertRaises(DemoError):self.fix(p,keys=['W01:table=1'])
 def test_autofit_conversion_requires_explicit_confirmation(self):
  with self.assertRaisesRegex(DemoError,'明确确认'):self.fix(fixed=[])
 def test_table_sizes_sum_and_all_cell_definitions_saved(self):
  r=self.fix();pkg=Package(r['output']);t=self.table(pkg);cols=[int(x.get(q('w'))) for x in t.findall('w:tblGrid/w:gridCol',NS)]
  self.assertEqual(sum(cols),8221);self.assertEqual(t.find('w:tblPr/w:tblW',NS).get(q('w')),'8221');self.assertEqual(t.find('w:tblPr/w:tblLayout',NS).get(q('type')),'fixed')
  for tr in t.findall(q('tr')):self.assertEqual([int(tc.find('w:tcPr/w:tcW',NS).get(q('w'))) for tc in tr.findall(q('tc'))],cols)
  self.assertEqual(r['widths']['after']['summary']['risks'],0)
 def test_integer_rounding_total_exact(self):
  for cols,target in [([4536,4536],8221),([901,321,1889],2701),([2,2,2],5)]:
   self.assertEqual(sum(_scale(cols,target)),target);self.assertTrue(all(abs(a/b-target/sum(cols))<1/b for a,b in zip(_scale(cols,target),cols)))
 def test_table_contents_borders_rows_unchanged_no_fixed_row_height(self):
  old=Package(self.p);r=self.fix();new=Package(r['output']);a=self.table(old);b=self.table(new)
  self.assertEqual(a.xpath('.//w:t/text()',namespaces=NS),b.xpath('.//w:t/text()',namespaces=NS));self.assertEqual(len(a.findall(q('tr'))),len(b.findall(q('tr'))));self.assertFalse(b.xpath('.//w:trHeight',namespaces=NS))
  self.assertEqual(a.find('w:tblPr/w:tblBorders',NS) is None,b.find('w:tblPr/w:tblBorders',NS) is None)
  self.assertTrue(all(r['preservation'][k] for k in ('original_sha256_unchanged','document_except_selected_format_attributes_equal','other_zip_member_bytes_equal')))
 def test_image_ratio_inner_outer_and_media_unchanged(self):
  old=Package(self.p);r=self.fix();new=Package(r['output']);a=old.paragraphs[1].find('.//{'+WP+'}extent');b=new.paragraphs[1].find('.//{'+WP+'}extent')
  self.assertAlmostEqual(int(a.get('cx'))/int(a.get('cy')),int(b.get('cx'))/int(b.get('cy')),places=5)
  x=new.paragraphs[1].find('.//{'+PIC+'}spPr/{'+A+'}xfrm/{'+A+'}ext');self.assertEqual(dict(b.attrib),dict(x.attrib))
  self.assertEqual(int(b.get('cx')),8221*635)
  for n in old.parts:
   if n.startswith('word/media/'):self.assertEqual(old.parts[n],new.parts[n])
 def test_image_crop_keepnext_preserved(self):
  def fn(pkg):
   p=pkg.paragraphs[1];_ensure(_ensure(p,'pPr'),'keepNext').set(q('val'),'1');fill=p.find('.//{'+PIC+'}blipFill');E.SubElement(fill,'{'+A+'}srcRect').set('l','100')
  p=self.edit(fn);r=self.fix(p);new=Package(r['output']);self.assertEqual(new.paragraphs[1].find('w:pPr/w:keepNext',NS).get(q('val')),'1');self.assertEqual(new.paragraphs[1].find('.//{'+A+'}srcRect').get('l'),'100')
 def test_unselected_object_unchanged_and_partial_risk(self):
  old=Package(self.p);r=self.fix(keys=['W01:table=1']);new=Package(r['output']);self.assertEqual(E.tostring(old.paragraphs[1]),E.tostring(new.paragraphs[1]));self.assertEqual(r['widths']['after']['summary']['risks'],1)
 def test_original_and_prior_round_and_numbering_sections_unchanged(self):
  raw=self.p.read_bytes();r=self.fix(keys=['W01:table=1']);p=Path(r['output']);prior=p.read_bytes();r2=self.fix(p,keys=['W02:paragraph=2'],fixed=[],name='out2.docx');new=Package(r2['output']);old=Package(self.p)
  self.assertEqual(raw,self.p.read_bytes());self.assertEqual(prior,p.read_bytes());self.assertEqual(old.parts['word/numbering.xml'],new.parts['word/numbering.xml']);self.assertEqual(E.tostring(old.sections[0]),E.tostring(new.sections[0]))
 def test_no_unnecessary_repeated_resize_or_normal_objects(self):
  r=self.fix();p=Path(r['output']);raw=p.read_bytes();plan=self.report(p);self.assertEqual(plan['summary']['passes'],2)
  with self.assertRaises(DemoError):self.fix(p,name='again.docx',plan=plan)
  self.assertFalse((self.d/'again.docx').exists());self.assertEqual(p.read_bytes(),raw)
 def test_nested_merge_and_floating_table_review(self):
  for fn in [lambda t:E.SubElement(t.find('w:tr/w:tc',NS),q('tbl')),lambda t:E.SubElement(t.find('w:tr/w:tc/w:tcPr',NS),q('gridSpan')).set(q('val'),'2'),lambda t:E.SubElement(t.find(q('tblPr')),q('tblpPr'))]:
   r=self.report(self.edit(lambda pkg:fn(self.table(pkg))))['objects'][0];self.assertEqual(r['status'],'review');self.assertFalse(r['fixable'])
 def test_floating_vml_multi_rotated_image_review(self):
  def rotate(pkg):pkg.paragraphs[1].find('.//{'+PIC+'}spPr/{'+A+'}xfrm').set('rot','60000')
  for fn in [lambda pkg:setattr(pkg.paragraphs[1].find('.//{'+WP+'}inline'),'tag','{'+WP+'}anchor'),lambda pkg:E.SubElement(pkg.paragraphs[1].find(q('r')),q('pict')),lambda pkg:pkg.paragraphs[1].append(deepcopy(pkg.paragraphs[1].find(q('r')))),rotate]:
   r=self.report(self.edit(fn))['objects'][1];self.assertEqual(r['status'],'review')
 def test_inner_extent_conflict_review(self):
  r=self.report(self.edit(lambda pkg:pkg.paragraphs[1].find('.//{'+PIC+'}spPr/{'+A+'}xfrm/{'+A+'}ext').set('cx','500')))['objects'][1];self.assertEqual(r['status'],'review')
 def test_mirror_multicolumn_and_missing_section_review(self):
  for fn in [lambda pkg:_ensure(pkg.sections[0],'cols').set(q('num'),'2'),lambda pkg:pkg.root.find(q('body')).remove(pkg.sections[0]),lambda pkg:pkg.sections[0].find(q('pgMar')).attrib.pop(q('left'))]:
   self.assertEqual(self.report(self.edit(fn))['summary']['reviews'],2)
 def test_character_indent_image_review_points_supported(self):
  p=self.edit(lambda pkg:_ensure(_ensure(pkg.paragraphs[1],'pPr'),'ind').set(q('firstLineChars'),'200'));self.assertEqual(self.report(p)['objects'][1]['status'],'review')
  p=self.edit(lambda pkg:_ensure(_ensure(pkg.paragraphs[1],'pPr'),'ind').set(q('left'),'1440'));self.assertAlmostEqual(self.report(p)['objects'][1]['available_mm'],119.6093,places=3)
 def test_too_narrow_columns_blocked(self):
  p=self.edit(lambda pkg:pkg.sections[0].find(q('pgMar')).set(q('right'),'9000'));r=self.report(p)['objects'][0];self.assertEqual(r['status'],'risk');self.assertFalse(r['fixable']);self.assertIn('10mm',r['reason'])
 def test_output_protection_and_plan_tampering(self):
  with self.assertRaises(DemoError):apply_width_adaptation(self.p,self.p,[],plan=self.report(),selected_keys=['W01:table=1'])
  self.fix()
  with self.assertRaises(DemoError):self.fix()
  p=deepcopy(self.report());p['objects'][0]['plan']['target_twips']=1
  with self.assertRaisesRegex(DemoError,'不匹配'):self.fix(name='bad.docx',plan=p)
  self.assertFalse((self.d/'bad.docx').exists())
 def test_rules_hash_and_input_changes_reject_plan(self):
  plan=self.report();plan['rules_sha256']='wrong'
  with self.assertRaises(DemoError):self.fix(plan=plan)
  p=self.edit(lambda pkg:pkg.sections[0].find(q('pgMar')).set(q('left'),'1000'))
  with self.assertRaises(DemoError):self.fix(p,plan=self.report())
 def test_forged_selection_and_duplicate_rejected(self):
  for keys in [[],['W01:table=99'],[True],['W01:table=1','W01:table=1']]:
   with self.assertRaises(DemoError):self.fix(keys=keys,fixed=[])
 def test_web_duplicate_download_rerun_does_not_regenerate(self):
  s=self.model();p=s.width_report('input');self.assertTrue(s.generate_width('input',p,['W01:table=1'],['W01:table=1']),s.width_error);n=s.generation_count;data=s.downloads()['fixed.docx']
  self.assertTrue(s.generate_width('input',p,['W01:table=1'],['W01:table=1']));self.assertEqual(s.generation_count,n);self.assertEqual(s.downloads()['fixed.docx'],data);self.assertIn('width-recheck.json',s.downloads())
 def test_web_failure_clears_success_and_preserves_input(self):
  s=self.model();raw=s.input_path.read_bytes();plan=s.width_report('input');self.assertTrue(s.generate_width('input',plan,['W01:table=1'],['W01:table=1']));bad=deepcopy(plan);bad['version']='bad'
  self.assertFalse(s.generate_width('input',bad,['W01:table=1'],['W01:table=1']));self.assertIsNone(s.result);self.assertEqual(s.downloads(),{});self.assertEqual(raw,s.input_path.read_bytes())
 def test_web_mode_file_and_scope_invalidation(self):
  for action in ('mode','file','body'):
   s=self.model();p=s.width_report('input');self.assertTrue(s.generate_width('input',p,['W01:table=1'],['W01:table=1']))
   if action=='mode':s.set_mode('extended')
   elif action=='file':s.load(self.p.read_bytes(),'另一个.docx')
   else:s.set_body([3])
   self.assertIsNone(s.result);self.assertEqual(s.downloads(),{});self.assertIsNone(s.width_signature)
 def test_web_two_sessions_and_multi_round_outputs(self):
  s=self.model();t=self.model();self.assertNotEqual(s.task_dir,t.task_dir);p=s.width_report('input');self.assertTrue(s.generate_width('input',p,['W01:table=1'],['W01:table=1']));self.assertEqual(t.history,[]);self.assertIsNone(t.result)
  raw=Path(s.result['output']).read_bytes();self.assertTrue(s.next_round());plan=s.width_report('input');self.assertTrue(s.generate_width('input',plan,['W02:paragraph=2']));self.assertEqual(s.result['round']['input_round'],1);self.assertEqual(s.result['round']['number'],2);self.assertEqual(Path(s.history[0]['result']['output']).read_bytes(),raw)
 def test_download_binding_rejects_changed_round_input(self):
  s=self.model();p=s.width_report('input');self.assertTrue(s.generate_width('input',p,['W01:table=1'],['W01:table=1']));s.input_path.write_bytes(b'changed')
  with self.assertRaises(DemoError):s.downloads()
 def test_geometry_not_visual_or_font_validation(self):
  r=self.fix();self.assertEqual(r['widths']['visual_status'],'not_reviewed');self.assertIn('不证明',width_markdown(r['widths']['after']));self.assertEqual(r['after']['summary']['rule_count'],7)
 def test_custom_body_title_rules_unchanged_by_width_adaptation(self):
  p=ROOT/'work/phase9/final/B-fixed.docx';template=ROOT/'work/phase9/final/template-B-active.json';m=json.loads((ROOT/'work/phase9/final/B-types.json').read_text());rh=effective_rules(custom_rules_path=template,mode='extended')['sha256'];plan=check_widths(p,rules_sha256=rh)
  r=apply_width_adaptation(p,self.d/'B.docx',None,custom_rules_path=template,mode='extended',classification_manifest=m,plan=plan,selected_keys=['W01:table=1'],confirmed_fixed_keys=['W01:table=1'])
  self.assertEqual(r['after']['summary']['problem_count'],0);self.assertEqual(r['after']['summary']['review_count'],0);self.assertEqual(r['after']['rules']['sha256'],rh);self.assertEqual(r['after']['summary']['rule_count'],12)

 def test_row_geometry_exception_and_fixed_height_review(self):
  for fn in [lambda t:E.SubElement(t.find(q('tr')),q('tblPrEx')),lambda t:_ensure(_ensure(t.find(q('tr')),'trPr'),'trHeight').set(q('hRule'),'exact')]:
   self.assertEqual(self.report(self.edit(lambda pkg:fn(self.table(pkg))))['objects'][0]['status'],'review')
 def test_default_table_style_geometry_is_not_ignored(self):
  def fn(pkg):
   t=self.table(pkg);t.find(q('tblPr')).remove(t.find('w:tblPr/w:tblStyle',NS))
   style=Resolver(pkg).styles['TableNormal'];_ensure(_ensure(style,'tblPr'),'tblpPr')
   pkg.parts['word/styles.xml']=E.tostring(pkg.styles)
  from core.formats import Resolver
  self.assertEqual(self.report(self.edit(fn))['objects'][0]['status'],'review')
 def test_table_layout_inserted_before_tbllook(self):
  r=self.fix();pr=self.table(Package(r['output'])).find(q('tblPr'));tags=[e.tag for e in pr]
  self.assertLess(tags.index(q('tblW')),tags.index(q('tblLayout')));self.assertLess(tags.index(q('tblLayout')),tags.index(q('tblLook')))
