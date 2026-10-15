"""Stage17 actual attributes, stories, PAGE, independent source and state acceptance."""
import unittest,json,zipfile
from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
from lxml import etree as E
from core import check_document,fix_document
from core.package import Package,q,NS,parse,sha256,DemoError
from core.front import inventory,section_inventory,fields,validate_scope,IDS,locate
from core.front_template import school,loads,validate,validate_draft,FIELDS as CONFIG,value
from core.engine import _ensure
from web.session import WebSession
ROOT=Path(__file__).resolve().parents[1];GOOD=ROOT/'samples/front-correct.docx';BAD=ROOT/'samples/front-errors.docx'
M=json.loads(GOOD.with_suffix('.types.json').read_text())
class FrontAcceptance(unittest.TestCase):
 def setUp(self):self.tmp=TemporaryDirectory();self.root=Path(self.tmp.name)
 def tearDown(self):self.tmp.cleanup()
 def manifest(self,p=GOOD):
  m=deepcopy(M);h=sha256(Path(p).read_bytes());m['input_sha256']=m['front_scope']['input_sha256']=h;return m
 def check(self,p=GOOD,m=None):return check_document(p,mode='academic',classification_manifest=m or self.manifest(p))
 def front(self,r):return [f for f in r['findings'] if f['rule_id'] in IDS]
 def alter(self,fn):
  pkg=Package(GOOD);fn(pkg)
  if pkg.styles is not None:pkg.parts['word/styles.xml']=E.tostring(pkg.styles,encoding='UTF-8')
  p=self.root/'altered.docx';p.write_bytes(pkg.bytes(True));return p
 def fix(self,p=BAD,m=None,keys=None,ids=('R24','R26','R27'),name='fixed.docx'):
  return fix_document(p,self.root/name,None,mode='academic',classification_manifest=m or self.manifest(p),selected_rules=ids,selected_keys=keys or [])
 def model(self):
  s=WebSession(self.root);s.set_mode('academic');self.assertTrue(s.load(BAD.read_bytes(),'软件演示.docx'));s.set_types(M['paragraph_types']);s.enable_front(True);s.edit_front(self.manifest(BAD)['front_scope']);s.confirm_front();self.assertTrue(s.confirm_and_check(),s.error);return s
 def test_school_title_two_is_22_not_body_12_or_heading18(self):self.assertEqual(school()['values']['cover_title__font_size'],22)
 def test_correct_control_new_rules_no_problem(self):self.assertFalse([f for f in self.front(self.check()) if f['status']=='problem'])
 def test_independent_injection_seven_and_late_run(self):
  fs=[f for f in self.front(self.check(BAD)) if f['status']=='problem'];gold=json.loads((ROOT/'samples/front-expected-errors.json').read_text());self.assertEqual(len(fs),len(gold['errors']));self.assertEqual(len(fs),7);self.assertEqual(sum(f['locator'].get('run')==2 for f in fs),3)
 def test_candidates_do_not_fill_or_infer_actual_name(self):
  inv=inventory(BAD);self.assertIn('name',next(r for r in inv['cover_objects'] if r['id']=='t:1:1:1:1')['candidates']);r=self.check(BAD);self.assertTrue(any('占位' in str(f['actual']) for f in r['findings'] if f['rule_id']=='R25'));self.assertFalse(any(f['fixable'] for f in r['findings'] if f['rule_id']=='R25'))
 def test_duplicate_name_and_absent_scope_only_review(self):
  fs=self.check()['findings'];self.assertTrue(any(f.get('property')=='candidates_name' and f['actual']==2 for f in fs));self.assertFalse(any(f['rule_id']=='R25' and f['status']=='pass' for f in fs))
 def test_actual_cell_cover_locator_and_no_table_rebuild(self):
  m=self.manifest();m['front_scope']['objects']={'t:1:2:2:1':'title'};out=self.fix(GOOD,m,ids=('R24',));a=Package(GOOD);b=Package(out['output']);self.assertEqual([c.text for c in a.root.iter(q('t'))],[c.text for c in b.root.iter(q('t'))]);self.assertTrue(any(f['locator'].get('table')==1 for f in out['changes']))
 def test_table_width_and_borders_preserved(self):
  m=self.manifest();m['front_scope']['objects']={'t:1:2:2:1':'title'};out=self.fix(GOOD,m,ids=('R24',));a=Package(GOOD);b=Package(out['output']);self.assertEqual(E.tostring(a.root.find('w:body/w:tbl/w:tblPr',NS)),E.tostring(b.root.find('w:body/w:tbl/w:tblPr',NS)));self.assertEqual(E.tostring(a.root.find('w:body/w:tbl/w:tblGrid',NS)),E.tostring(b.root.find('w:body/w:tbl/w:tblGrid',NS)))
 def test_style_inheritance_and_all_ordinary_runs(self):
  def change(p):
   rp=p.paragraphs[1].findall(q('r'))[1].find(q('rPr'));rp.remove(rp.find(q('sz')))
   st=p.styles.find("w:style[@w:styleId='Normal']",NS);_ensure(_ensure(st,'rPr'),'sz').set(q('val'),'24')
  p=self.alter(change);f=next(f for f in self.check(p)['findings'] if f['rule_id']=='R24' and f['locator'].get('run')==2 and f.get('property')=='font_size');self.assertEqual(f['actual'],12);self.assertEqual(f['status'],'problem')
 def test_special_cover_link_protected(self):
  p=self.alter(lambda p:p.paragraphs[1].append(E.Element(q('hyperlink'))));self.assertFalse(any(f['fixable'] for f in self.front(self.check(p)) if f['rule_id']=='R24'))
 def test_cover_formula_and_revision_protected(self):
  for tag in ('ins','fldSimple'):
   p=self.alter(lambda p:p.paragraphs[1].append(E.Element(q(tag))));self.assertTrue(any(f['rule_id']=='R24' and f['status']=='review' for f in self.check(p)['findings']));p.unlink()
 def test_real_section_object_range(self):
  inv=inventory(GOOD);secs=inv['sections'];self.assertEqual(len(secs),4);self.assertEqual(secs[0]['body_objects']['start'],1);self.assertEqual(secs[1]['body_objects']['start'],secs[0]['body_objects']['end']+1);self.assertEqual(secs[-1]['body_objects']['end'],len(list(Package(GOOD).root.find(q('body')))))
 def test_shared_story_effective_link(self):
  s=inventory(GOOD)['sections'];self.assertTrue(s[3]['references']['header:default']['linked_previous']);self.assertEqual(s[3]['references']['header:default']['part'],s[2]['references']['header:default']['part'])
 def test_only_selected_distance_changes(self):
  r=self.check(BAD);key=next(f['key'] for f in r['findings'] if f['rule_id']=='R26' and f['locator']['section']==2);out=self.fix(keys=[key],ids=());self.assertEqual(len(out['changes']),1);a,b=Package(BAD),Package(out['output']);self.assertEqual(a.sections[2].find(q('pgMar')).attrib,b.sections[2].find(q('pgMar')).attrib)
 def test_shared_unselected_section_blocks_story_font(self):
  m=self.manifest();m['front_scope']['section_roles']['4']='pending';m['front_scope']['template']['name']='用户确认的演示要求';n='header_text__font_size';m['front_scope']['template']['values'][n]=12;m['front_scope']['template']['sources'][n]={'kind':'user_confirmed','quote':'','location':None}
  fs=[f for f in self.check(m=m)['findings'] if f['rule_id']=='R27' and f['locator'].get('section')==3 and f.get('property')=='font_size'];self.assertTrue(fs);self.assertTrue(all(not f['fixable'] and 4 in f['affected_sections'] for f in fs))
 def test_shared_selected_simple_story_can_fix_all_impacted_sections(self):
  m=self.manifest();m['front_scope']['section_roles']['4']='body';n='header_text__font_size';m['front_scope']['template']['name']='用户确认的演示要求';m['front_scope']['template']['values'][n]=12;m['front_scope']['template']['sources'][n]={'kind':'user_confirmed','quote':'','location':None};fs=[f for f in self.check(m=m)['findings'] if f['rule_id']=='R27' and f['locator'].get('section')==3 and f.get('property')=='font_size'];self.assertTrue(all(f['fixable'] for f in fs));out=self.fix(GOOD,m,keys=[fs[0]['key']],ids=());self.assertEqual(out['changes'][0]['verification'],'pass')
 def test_selected_story_does_not_touch_other_package_members(self):
  out=self.fix();a,b=Package(BAD),Package(out['output']);changed=[n for n in a.parts if a.parts[n]!=b.parts[n]];self.assertEqual(set(changed),{'word/document.xml',inventory(BAD)['sections'][1]['references']['header:default']['part']});self.assertTrue(out['preservation']['unselected_zip_member_bytes_equal']);self.assertFalse(out['preservation']['other_zip_member_bytes_equal'])
 def test_fields_media_numbering_and_relationships_byte_equal(self):
  out=self.fix();a,b=Package(BAD),Package(out['output']);self.assertTrue(all(b.parts[n]==v for n,v in a.parts.items() if n.endswith('.rels') or n.startswith('word/media/') or n.endswith('numbering.xml') or 'footer' in n))
 def test_simple_page_field_read_only(self):
  f=next(f for f in self.check()['findings'] if f['rule_id']=='R29' and f.get('property')=='page_field' and f['locator'].get('section')==2);self.assertEqual(f['actual']['representation'],'simple');self.assertEqual(f['actual']['saved_result'],'I');self.assertFalse(f['fixable']);self.assertEqual(f['status'],'review')
 def test_complex_page_field_balanced(self):
  f=next(f for f in self.check()['findings'] if f['rule_id']=='R29' and f.get('property')=='page_field' and f['locator'].get('section')==3);self.assertEqual(f['actual']['representation'],'complex');self.assertTrue(f['actual']['balanced']);self.assertTrue(f['actual']['is_page'])
 def test_not_text_search_page_and_not_numpages(self):
  root=E.fromstring(f'<w:ftr xmlns:w="{NS["w"]}"><w:p><w:r><w:t>PAGE</w:t></w:r><w:fldSimple w:instr="NUMPAGES"><w:r><w:t>3</w:t></w:r></w:fldSimple></w:p></w:ftr>');self.assertFalse(any(f['is_page'] for f in fields(root)))
 def test_unbalanced_complex_fields_review(self):
  root=E.fromstring(f'<w:p xmlns:w="{NS["w"]}"><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText> PAGE </w:instrText></w:r></w:p>');self.assertFalse(fields(root)[0]['balanced'])
 def test_roman_arabic_format_separate_from_actual_pages(self):
  fs=self.check()['findings'];a=[f for f in fs if f.get('property')=='number_format'];self.assertEqual([f['status'] for f in a],['pass','pass']);self.assertTrue(all(f['status']=='review' for f in fs if f.get('property')=='number_start'))
 def test_mixed_use_does_not_auto_split(self):
  out=self.fix();self.assertEqual(len(Package(out['output']).sections),4);self.assertTrue(any(f.get('property')=='mixed_use' and f['status']=='review' for f in out['after']['findings']))
 def test_missing_page_scope_not_declared_definitely_missing(self):
  m=self.manifest();m['front_scope']['section_roles']['1']='body';fs=self.check(m=m)['findings'];self.assertTrue(any(f.get('property')=='page_presence' and f['status']=='review' for f in fs))
 def test_page_position_conflict_kept_visible(self):
  self.assertTrue(any(f.get('property')=='page_position' and '冲突' in f['actual'] for f in self.check()['findings']))
 def test_full_fix_and_independent_recheck(self):
  out=self.fix();again=self.check(out['output'],out['output_classification_manifest']);self.assertEqual(again,out['after']);self.assertFalse(any(f['status']=='problem' for f in self.front(again)));self.assertTrue(all(f['verification']=='pass' for f in out['changes']))
 def test_repeated_fix_no_extra_changes_or_bytes(self):
  a=self.fix();b=self.fix(Path(a['output']),a['output_classification_manifest'],name='two.docx');self.assertEqual(b['changes'],[]);self.assertEqual(Path(a['output']).read_bytes(),Path(b['output']).read_bytes())
 def test_partial_fix_preserves_unselected_title_format(self):
  key=next(f['key'] for f in self.check(BAD)['findings'] if f['rule_id']=='R24' and f.get('property')=='font_size' and f['fixable']);out=self.fix(keys=[key],ids=());self.assertEqual(len(out['changes']),1);self.assertEqual(sum(f['status']=='problem' for f in self.front(out['after'])),6)
 def test_original_and_existing_output_protected(self):
  h=sha256(BAD.read_bytes());out=self.fix();self.assertEqual(h,sha256(BAD.read_bytes()))
  with self.assertRaises(DemoError):self.fix()
  with self.assertRaises(DemoError):fix_document(BAD,BAD,None,mode='academic',classification_manifest=self.manifest(BAD),selected_rules=['R24'])
 def test_old_manifest_counts_unchanged(self):
  m=self.manifest();m.pop('front_scope');self.assertEqual(self.check(m=m)['summary']['rule_count'],22)
 def test_bad_scope_hash_unknown_object_and_unconfirmed(self):
  for mutate in [lambda s:s.update(input_sha256='0'*64),lambda s:s.update(confirmed=False),lambda s:s['objects'].update({'t:999:1:1:1':'title'}),lambda s:s['section_roles'].update({'1':'bogus'})]:
   m=self.manifest();mutate(m['front_scope'])
   with self.assertRaises(DemoError):self.check(m=m)
 def test_template_roundtrip_custom_actual_title(self):
  m=self.manifest();t=m['front_scope']['template'];t['name']='课程封面演示';t['values']['cover_title__font_size']=24;t['sources']['cover_title__font_size']={'kind':'user_confirmed','quote':'','location':None};t=loads(json.dumps(t,ensure_ascii=False).encode());m['front_scope']['template']=t;out=self.fix(GOOD,m,ids=('R24',));self.assertTrue(any(f['expected']==24 and f['actual']==24 for f in out['after']['findings'] if f.get('property')=='font_size' and f['rule_id']=='R24'))
 def test_template_invalid_unknown_bool_nan_range_and_source(self):
  for n,v in [('cover_title__font_size',True),('cover_title__font_size',float('nan')),('cover_title__font_size',36.5),('cover_title__bold','true'),('header_distance_mm',-1),('header_text__font_cn','未知')]:
   with self.assertRaises(DemoError):value(n,v)
  t=school();t['values']['cover_title__font_size']=24
  with self.assertRaises(DemoError):validate(t)
  t=school();t['execute']='read'
  with self.assertRaises(DemoError):validate(t)
 def test_AI_missing_fields_and_exact_object_evidence(self):
  text='封面论文标题二号黑体，加粗，居中。';d={n:None for n in CONFIG};d['unsupported']=[];self.assertIsNone(validate_draft(d,text)['footer_text__font_size']);d['cover_title__font_size']={'value':22,'unit':'pt','quote':'封面论文标题二号黑体，加粗，居中','location':'TEXT:1','support':'supported'};self.assertEqual(validate_draft(d,text)['cover_title__font_size']['value'],22)
 def test_AI_conflict_wrong_quote_and_body_title_rejected(self):
  for text,qte in [('正文标题二号黑体','正文标题二号黑体'),('封面标题二号。封面标题四号。','封面标题二号'),('封面标题二号','封面标题三号')]:
   d={n:None for n in CONFIG};d['unsupported']=[];d['cover_title__font_size']={'value':22,'unit':'pt','quote':qte,'location':'TEXT:1','support':'supported'}
   with self.assertRaises(DemoError):validate_draft(d,text)
 def test_web_confirm_and_partial_then_round_preserves_binding(self):
  s=self.model();s.set_selection([f['key'] for f in s.eligible() if f['rule_id']=='R24']);self.assertTrue(s.generate(),s.error);first=Path(s.result['output']);h=sha256(first.read_bytes());self.assertTrue(s.next_round(),s.error);self.assertEqual(s.front_manifest()['input_sha256'],s.inventory['input_sha256']);s.set_selection([f['key'] for f in s.eligible()]);self.assertTrue(s.generate(),s.error);self.assertEqual(sha256(first.read_bytes()),h);self.assertFalse(any(f['status']=='problem' for f in self.front(s.result['after'])))
 def test_web_scope_change_invalidates_report_downloads(self):
  s=self.model();d=deepcopy(s.front_draft);d['section_roles']['1']='pending';s.edit_front(d);self.assertIsNone(s.report);self.assertIsNone(s.front_confirmed);self.assertEqual(s.downloads(),{})
 def test_two_session_front_drafts_isolate(self):
  a=self.model();b=WebSession(self.root);b.set_mode('academic');b.load(GOOD.read_bytes(),'other.docx');self.assertFalse(b.front_enabled);self.assertIsNone(b.front_draft);self.assertIsNotNone(a.front_draft)
 def test_project_restores_scope_template_and_fresh_report(self):
  s=self.model();raw=s.prepare_project();other=WebSession(self.root);other.stage_project(raw);new=other.restore_project(acknowledged=True);self.assertEqual(new.front_draft['objects'],s.front_draft['objects']);self.assertEqual(new.front_draft['template'],s.front_draft['template']);self.assertEqual(new.report['summary'],s.report['summary']);self.assertEqual(new.report['coverage']['front_scope']['input_sha256'],new.inventory['input_sha256'])
 def test_project_unconfirmed_draft_not_promoted(self):
  s=self.model();s.invalidate();raw=s.prepare_project();other=WebSession(self.root);other.stage_project(raw);new=other.restore_project(acknowledged=True);self.assertTrue(new.front_enabled);self.assertIsNone(new.report);self.assertIsNone(new.front_confirmed)
 def test_zip_includes_supplement_and_no_private_path(self):
  s=self.model();s.set_selection([f['key'] for f in s.eligible()]);self.assertTrue(s.generate(),s.error);raw=s.result_package();z=zipfile.ZipFile(__import__('io').BytesIO(raw));self.assertIn('front-supplement.json',z.namelist());self.assertEqual(json.loads(z.read('front-supplement.json')),s.result['after']['coverage']['front_scope']['template']);self.assertFalse(any('/Volumes/' in n or '.env' in n for n in z.namelist()))
 def test_mock_API_only_norm_once_no_paper_or_personal_data(self):
  from ai.front_spec import parse_requirements
  from ai.short_spec import Settings
  text='封面论文标题二号黑体，加粗，居中。页眉距离1.5cm。'
  d={n:None for n in CONFIG};d['unsupported']=[];d['cover_title__font_size']={'value':22,'unit':'pt','quote':'封面论文标题二号黑体，加粗，居中','location':'TEXT:1','support':'supported'};d['header_distance_mm']={'value':15,'unit':'mm','quote':'页眉距离1.5cm','location':'TEXT:1','support':'supported'};requests=[]
  def transport(req,timeout):
   requests.append((json.loads(req.data),timeout));return json.dumps({'choices':[{'finish_reason':'stop','message':{'content':json.dumps(d,ensure_ascii=False)}}]})
  out=parse_requirements(text,transport=transport,config=Settings('not-real-secret','https://api.deepseek.com','deepseek-flash'));self.assertEqual(len(requests),1);self.assertEqual(requests[0][1],30);self.assertEqual(requests[0][0]['messages'][-1]['content'],text);self.assertNotIn('not-real-secret',json.dumps(requests[0][0]));self.assertEqual(out['receipt']['kind'],'mock_api')
 def test_front_AI_failure_does_not_remove_fixed_check(self):
  s=self.model();before=deepcopy(s.report);s.set_front_spec('封面标题二号黑体')
  with self.assertRaises(DemoError):s.parse_front(parser=lambda text:(_ for _ in ()).throw(DemoError('timeout')))
  self.assertEqual(before,s.report);self.assertIsNone(s.front_ai_draft)
 def test_front_AI_draft_stays_missing_until_user_adopts(self):
  s=self.model();s.set_front_spec('封面论文标题二号。');d={n:None for n in CONFIG};d['unsupported']=[];s.parse_front(parser=lambda text:{'draft':d,'receipt':{'kind':'mock_api','model':'mock','timestamp_utc':'2026-10-15'}});self.assertEqual(s.front_draft['template']['values']['header_distance_mm'],15);s.use_front_ai_draft();self.assertIsNone(s.front_draft['template']['values']['header_distance_mm']);self.assertIsNone(s.report)
 def test_vml_story_not_rewritten_and_kept(self):
  def change(pkg):
   part=section_inventory(pkg)[1]['references']['header:default']['part'];root=parse(pkg.parts[part]);p=root.findall(q('p'))[0];p.append(E.Element(q('pict')));pkg.parts[part]=E.tostring(root,encoding='UTF-8')
  path=self.alter(change);m=self.manifest(path);r=self.check(path,m);self.assertFalse(any(f['fixable'] for f in r['findings'] if f['rule_id']=='R27' and f['locator'].get('section')==2));out=self.fix(path,m,ids=('R26',));part=section_inventory(Package(path))[1]['references']['header:default']['part'];self.assertEqual(Package(path).parts[part],Package(out['output']).parts[part])
 def test_external_file_change_blocks_web_outputs(self):
  s=self.model();s.set_selection([f['key'] for f in s.eligible()]);self.assertTrue(s.generate());Path(s.result['output']).write_bytes(GOOD.read_bytes())
  with self.assertRaises(DemoError):s.downloads()
 def test_project_two_rounds_again_and_history_stable(self):
  s=self.model();s.set_selection([f['key'] for f in s.eligible() if f['rule_id']=='R24']);self.assertTrue(s.generate());self.assertTrue(s.next_round());s.set_selection([f['key'] for f in s.eligible()]);self.assertTrue(s.generate());raw=s.prepare_project();o=WebSession(self.root);o.stage_project(raw);r=o.restore_project(acknowledged=True);self.assertEqual(len(r.history),2);self.assertEqual(r.front_draft['template'],s.front_draft['template']);r.prepare_project();self.assertEqual(len(r.history_rows()),2)
 def test_mode_file_change_invalidates_front_confirmation(self):
  s=self.model();s.set_mode('extended');self.assertIsNone(s.front_confirmed);self.assertIsNone(s.report);self.assertTrue(s.load(GOOD.read_bytes(),'new.docx'));self.assertFalse(s.front_enabled);self.assertIsNone(s.front_draft)
 def test_preview_target_rebinds_front_scope_to_original(self):
  s=self.model();s.set_selection([f['key'] for f in s.eligible()]);self.assertTrue(s.generate());self.assertTrue(s.next_round());targets=s.preview_targets();self.assertEqual(targets['original']['classification']['front_scope']['input_sha256'],sha256(BAD.read_bytes()));self.assertEqual(targets['input']['classification']['front_scope']['input_sha256'],s.inventory['input_sha256'])
 def test_corrupt_story_and_duplicate_relationships_review_no_repair(self):
  def change(pkg):
   part=section_inventory(pkg)[1]['references']['header:default']['part'];pkg.parts[part]=b'<broken'
  p=self.alter(change);r=self.check(p);self.assertTrue(any(f['rule_id']=='R28' and f.get('property')=='relationships' and f['status']=='review' for f in r['findings']));self.assertFalse(any(f['fixable'] for f in r['findings'] if f['rule_id']=='R27' and f['locator'].get('section')==2))
 def test_footer_ordinary_text_font_can_fix_without_touching_PAGE(self):
  def change(pkg):
   part=section_inventory(pkg)[1]['references']['footer:default']['part'];root=parse(pkg.parts[part]);p=E.SubElement(root,q('p'));pp=E.SubElement(p,q('pPr'));E.SubElement(pp,q('jc')).set(q('val'),'center');r=E.SubElement(p,q('r'));rp=E.SubElement(r,q('rPr'));E.SubElement(rp,q('sz')).set(q('val'),'24');E.SubElement(r,q('t')).text='软件演示页脚';pkg.parts[part]=E.tostring(root)
  p=self.alter(change);m=self.manifest(p);n='footer_text__font_size';t=m['front_scope']['template'];t['name']='用户确认的补充演示';t['values'][n]=10.5;t['sources'][n]={'kind':'user_confirmed','quote':'','location':None};fs=[f for f in self.check(p,m)['findings'] if f['rule_id']=='R27' and f['locator'].get('section')==2 and f.get('property')=='font_size'];self.assertEqual(len(fs),1);out=self.fix(p,m,keys=[fs[0]['key']],ids=());self.assertEqual(out['changes'][0]['verification'],'pass');part=section_inventory(Package(p))[1]['references']['footer:default']['part'];a=parse(Package(p).parts[part]);b=parse(Package(out['output']).parts[part]);self.assertEqual(E.tostring(a.find('.//'+q('fldSimple'))),E.tostring(b.find('.//'+q('fldSimple'))))

 def test_school_header_format_not_silently_applied_to_front(self):
  m=self.manifest(BAD);m['front_scope']['template']=school();fs=self.front(self.check(BAD,m));self.assertFalse(any(f['rule_id']=='R27' and f['locator'].get('section')==2 and f.get('property')=='alignment' for f in fs));self.assertTrue(any(f.get('property')=='header_outside_body' for f in fs))
 def test_actual_embedded_media_and_drawing_preserved_with_story_edit(self):
  from io import BytesIO
  from docx import Document
  from docx.shared import Mm
  asset=Package(ROOT/'samples/correct.docx');media=next(v for n,v in asset.parts.items() if n.startswith('word/media/') and n.endswith('.png'))
  d=Document(BAD);d.paragraphs[-1].add_run().add_picture(BytesIO(media),width=Mm(12));p=self.root/'with-image.docx';d.save(p);out=self.fix(p,self.manifest(p));a,b=Package(p),Package(out['output']);names=[n for n in a.parts if n.startswith('word/media/')];self.assertTrue(names);self.assertTrue(all(a.parts[n]==b.parts[n] for n in names));self.assertEqual([E.tostring(e) for e in a.root.iter(q('drawing'))],[E.tostring(e) for e in b.root.iter(q('drawing'))]);self.assertTrue(out['preservation']['unselected_zip_member_bytes_equal'])

 def test_switch_extended_can_confirm_without_hidden_new_scopes(self):
  s=self.model();s.set_table_enabled(True);s.set_table_types({'1':'cover'});s.confirm_tables();s.set_mode('extended');self.assertTrue(s.confirm_and_check(),s.error);self.assertEqual(s.report['summary']['rule_count'],12);self.assertNotIn('front_scope',s.classification_confirmed);self.assertNotIn('table_scope',s.classification_confirmed)

 def test_imported_AI_source_requires_matching_quote_not_self_assertion(self):
  t=school();t['name']='导入历史规范';n='cover_title__font_size';t['sources'][n]={'kind':'ai_extracted','quote':'','location':None}
  with self.assertRaises(DemoError):validate(t)
  t['sources'][n]={'kind':'ai_extracted','quote':'正文小四','location':'TEXT:1'}
  with self.assertRaises(DemoError):validate(t)
  t['sources'][n]={'kind':'ai_extracted','quote':'封面标题二号','location':'TEXT:1'};self.assertEqual(validate(t)['values'][n],22)

 def test_dynamic_even_header_saved_text_not_static_content_pass(self):
  def change(pkg):
   part=section_inventory(pkg)[2]['references']['header:even']['part'];root=parse(pkg.parts[part]);p=root.findall(q('p'))[0]
   for child in list(p):p.remove(child)
   f=E.SubElement(p,q('fldSimple'));f.set(q('instr'),' STYLEREF 1 ');r=E.SubElement(f,q('r'));E.SubElement(r,q('t')).text='华南理工大学学士学位论文';pkg.parts[part]=E.tostring(root)
  path=self.alter(change);fs=self.check(path)['findings'];f=next(f for f in fs if f.get('property')=='even_header_content');self.assertEqual(f['status'],'review');self.assertFalse(f['fixable'])

 def test_empty_or_placeholder_title_not_treated_as_formatted_title(self):
  p=self.alter(lambda pkg:[r.getparent().remove(r) for r in list(pkg.paragraphs[1].findall(q('r')))])
  fs=self.check(p)['findings'];self.assertTrue(any(f.get('property')=='title_content' and f['status']=='review' for f in fs));self.assertFalse(any(f['fixable'] for f in fs if f['rule_id']=='R24'))

 def test_cover_objects_cannot_overlap_body_or_heading_classification(self):
  for typ in ('body','heading1','abstract_cn_body'):
   m=self.manifest();m['paragraph_types']['2']=typ
   with self.assertRaisesRegex(DemoError,'重叠'):self.check(m=m)
 def test_cover_cell_cannot_also_be_content_table_scope(self):
  m=self.manifest();m['table_scope']={'input_sha256':m['input_sha256'],'table_types':{'1':'content'},'confirmed':True,'confirmation':'测试显式但互相冲突的范围'}
  with self.assertRaisesRegex(DemoError,'封面对象所在表格'):self.check(m=m)
