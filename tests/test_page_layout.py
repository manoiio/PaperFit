from pathlib import Path
from tempfile import TemporaryDirectory
from copy import deepcopy
from unittest.mock import patch
import unittest,json,zipfile
from lxml import etree as E
from core.package import Package,DemoError,q,NS
from core import page_layout as pages
from core.page_policy import DEFAULT,extract
from scripts.build_page_fixture import build
from web.session import WebSession

class PageLayoutTests(unittest.TestCase):
 def setUp(self):
  self.tmp=TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.root=Path(self.tmp.name);self.path,self.roles=build(self.root/'fixture');self.policy={**DEFAULT,'pages':'front-body','footer':'chapter','dynamic':True}
 def manifest(self,path=None):
  path=path or self.path;return {'input_sha256':Package(path).hash,'paragraph_types':self.roles,'confirmed':True,'confirmation':'作者手写范围'}
 def run_pages(self,path=None):
  path=path or self.path;m=self.manifest(path);r=pages.automatic(Package(path),m,self.policy);plan=pages.propose(path,m,self.policy,r);out=self.root/('out-'+path.name);result=pages.rehearse(path,out,m,plan);return out,result
 def test_existing_and_single_preserve_content_and_inherited_parts(self):
  for source in (self.path,self.root/'fixture/single.docx'):
   out,r=self.run_pages(source);a,b=Package(source),Package(out)
   self.assertEqual(len(a.paragraphs),len(b.paragraphs));self.assertEqual(len(b.sections),4)
   self.assertEqual([''.join(p.itertext()) for p in a.paragraphs],[''.join(p.itertext()) for p in b.paragraphs])
   for x,y in zip(a.paragraphs,b.paragraphs):self.assertEqual([E.tostring(r) for r in x.findall(q('r'))],[E.tostring(r) for r in y.findall(q('r'))])
   for n in a.parts:
    if n not in ('word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'):self.assertEqual(a.parts[n],b.parts[n])
   inv=__import__('core.front',fromlist=['section_inventory']).section_inventory(b)
   self.assertEqual([s['page_numbering'].get('start') for s in inv],[None,'1','1',None]);self.assertEqual([s['page_numbering'].get('fmt') for s in inv],[None,'lowerRoman','decimal','decimal'])
   for i,s in enumerate(inv):
    for variant in ('default','first','even'):
     root=E.fromstring(b.parts[s['references']['footer:'+variant]['part']]);fields=root.findall('.//'+q('fldSimple'))
     self.assertEqual(len(fields),0 if i==0 else 1)
     if i>=2:self.assertIn('第'+str(i-1)+'章',''.join(root.itertext()))
   p2=pages.propose(out,r['output_classification_manifest'],self.policy,r['page_layout']['plan']['regions']);self.assertFalse(p2['records']);rep,add,_,_=pages.build(b,p2['regions']);self.assertFalse(rep);self.assertFalse(add)
 def test_structural_whitelist_and_unknown_tamper_refuse(self):
  out,r=self.run_pages();pkg=Package(out);pkg.paragraphs[5].find(q('r')).find(q('t')).text+='changed';out.write_bytes(pkg.bytes(True))
  with self.assertRaises(DemoError):pages.validate_output(self.path,out,self.manifest(),r['page_layout']['plan'])
  with self.assertRaises(DemoError):Package(self.path).structure_bytes({}, {'word/evil.xml':b'bad'})
 def test_keep_region_and_variants_remain_independent(self):
  pkg=Package(self.path);r=pages.automatic(pkg,self.manifest(),self.policy);r[-1]['role']='keep';plan=pages.propose(self.path,self.manifest(),self.policy,r);out=self.root/'keep.docx';pages.rehearse(self.path,out,self.manifest(),plan)
  from core.front import section_inventory
  old,new=section_inventory(pkg),section_inventory(Package(out))
  for v in ('default','first','even'):self.assertEqual(old[-1]['references']['footer:'+v]['part'],new[-1]['references']['footer:'+v]['part'])
  # Chapter-footer-only authorization must not add PAGE or alter cover/front stories.
  only={**self.policy,'pages':'chapter-footers'};regions=pages.automatic(pkg,self.manifest(),only);self.assertEqual(regions[0]['role'],'keep')
  plan=pages.propose(self.path,self.manifest(),only,regions);output=self.root/'chapter-only.docx';pages.rehearse(self.path,output,self.manifest(),plan);actual=section_inventory(Package(output))
  for i in (0,1):
   self.assertEqual(old[i]['page_numbering'],actual[i]['page_numbering'])
   for v in ('default','first','even'):self.assertEqual(old[i]['references']['footer:'+v]['part'],actual[i]['references']['footer:'+v]['part'])
 def test_table_before_later_section_keeps_existing_restart_and_scope_identity(self):
  from docx import Document
  d=Document(self.path);table=d.add_table(rows=1,cols=1);table.cell(0,0).text='未选表格内容';anchor=d.paragraphs[5]._p;anchor.addprevious(table._tbl);d.sections[-1]._sectPr.append(E.Element(q('pgNumType'),{q('fmt'):'decimal',q('start'):'7'}));path=self.root/'body-table.docx';d.save(path)
  pkg=Package(path);policy={**self.policy,'pages':'chapter-footers'};m=self.manifest(path);regions=pages.automatic(pkg,m,policy);plan=pages.propose(path,m,policy,regions);out=self.root/'body-table-pages.docx';pages.rehearse(path,out,m,plan);saved=Package(out)
  self.assertEqual(saved.sections[-1].find(q('pgNumType')).get(q('start')),'7');self.assertEqual(E.tostring(pkg.root.find('w:body/w:tbl',NS)),E.tostring(saved.root.find('w:body/w:tbl',NS)))
 def test_complex_header_page_and_restart_notes_protection(self):
  pkg=Package(self.path);name=__import__('core.front',fromlist=['section_inventory']).section_inventory(pkg)[0]['references']['header:default']['part'];root=E.fromstring(pkg.parts[name]);E.SubElement(root.find(q('p')),q('fldSimple'),{q('instr'):' PAGE '});pkg.parts[name]=E.tostring(root)
  path=self.root/'header.docx';path.write_bytes(pkg.structure_bytes({},{}))
  # Use exact ZIP edit for this protection input.
  import io
  stream=io.BytesIO()
  with zipfile.ZipFile(stream,'w') as z:
   for n,b in pkg.parts.items():z.writestr(n,b)
  path.write_bytes(stream.getvalue())
  with self.assertRaises(DemoError):pages.propose(path,self.manifest(path),self.policy,pages.automatic(Package(path),self.manifest(path),self.policy))
  # Existing shared paragraph safety also applies to a new structural boundary.
  revised=Package(self.path);pp=revised.paragraphs[1].find(q('pPr'));E.SubElement(pp,q('pPrChange'));rp=self.root/'boundary-revision.docx';rp.write_bytes(revised.bytes(True))
  with self.assertRaises(DemoError):pages.propose(rp,self.manifest(rp),self.policy,pages.automatic(Package(rp),self.manifest(rp),self.policy))
 def test_smart_complete_same_local_no_network_and_idempotent(self):
  from web.onepass_session import start,current
  s=WebSession(self.root/'smart');s.set_mode('academic');s.load(self.path.read_bytes(),'公开合成.docx');s.set_types(self.roles);s.confirm_classification();s.page_layout_draft={'enabled':True,'policy':self.policy,'input_sha256':s.inventory['input_sha256'],'regions':[]};state={}
  with patch('urllib.request.OpenerDirector.open',side_effect=AssertionError('no network')):start(s,state,structure_runner=lambda *_:self.fail('confirmed scope no AI'))
  self.assertEqual(state['status'],'done',state.get('error'));self.assertEqual([x['kind'] for x in state['audit']['assembly_steps']],['format','dynamic_text','section_pages']);self.assertEqual(s.current_round,0);self.assertTrue(current(s,state,''));self.assertEqual(len(s.history),3)
  old=s.result['after']['input_sha256'];start(s,state,'');self.assertEqual(len(s.history),3);self.assertEqual(old,s.result['after']['input_sha256']);self.assertTrue(s.downloads())
  from streamlit.testing.v1 import AppTest
  app=AppTest.from_file(str(Path(__file__).resolve().parents[1]/'app.py'),default_timeout=35);app.session_state['workflow']=s;app.session_state['workspace-step']=5;app.run()
  self.assertFalse(app.exception);self.assertFalse(any('完整性结果异常' in x.value for x in app.error));self.assertTrue(any(x.label=='下载修正版DOCX' for x in app.get('download_button')))
 def test_local_complete_preset_defaults_to_numeric_sync_without_changing_old_default(self):
  from streamlit.testing.v1 import AppTest
  x=WebSession(self.root/'local-preset');x.set_mode('academic');x.load(self.path.read_bytes(),'公开合成.docx');x.set_types(self.roles);x.confirm_and_check()
  app=AppTest.from_file(str(Path(__file__).resolve().parents[1]/'app.py'),default_timeout=35);app.session_state['workflow']=x;app.run();app.button(key='workspace-nav-5').click().run()
  self.assertIsNone(x.page_layout_draft)
  next(c for c in app.selectbox if c.label=='本次整理方案').set_value('complete').run()
  self.assertFalse(app.exception);self.assertTrue(x.page_layout_draft['policy']['dynamic']);self.assertTrue(next(c for c in app.checkbox if c.label=='按本次方案同步安全编号与纯文本引用').value)
  self.assertFalse(x.history);self.assertIsNone(x.result)
 def test_local_exact_requirements_and_bundle_project_compatibility(self):
  p,rest,proof=extract('封面不显示页码，摘要和目录用小写罗马数字，正文从1开始连续编号。各章页脚显示章节名称和页码。整理图表编号并同步正文中的对应引用。');self.assertEqual(p,self.policy);self.assertFalse(rest);self.assertEqual(len(proof),3)
  with self.assertRaises(DemoError):extract('封面不显示页码，摘要和目录用小写罗马数字，正文从1开始连续编号；封面不显示页码，摘要和目录用大写罗马数字，正文从1开始连续编号')
 def test_local_and_smart_equal_atomic_failure_and_restore(self):
  from web.onepass_session import start
  from web.assembly_session import execute_local
  def model(name):
   x=WebSession(self.root/name);x.set_mode('academic');x.load(self.path.read_bytes(),'公开合成.docx');x.set_types(self.roles);x.confirm_classification();x.page_layout_draft={'enabled':True,'policy':deepcopy(self.policy),'input_sha256':x.inventory['input_sha256'],'regions':[]};return x
  smart=model('smart-equal');state={};start(smart,state,structure_runner=lambda *_:self.fail('no AI'));self.assertEqual(state['status'],'done',state.get('error'))
  local=model('local-equal')
  with patch('urllib.request.OpenerDirector.open',side_effect=AssertionError('no network')):execute_local(local,{},acknowledged=True)
  self.assertEqual(smart.downloads()['fixed.docx'],local.downloads()['fixed.docx'])
  bundle=smart.rule_bundle_download();reuse=model('reuse');reuse.import_rule_bundle(bundle);reuse.apply_workbench(acknowledged=True);self.assertEqual(reuse.current_rule_bundle()['assembly'],self.policy)
  raw=smart.prepare_project();reuse.stage_project(raw);restored=reuse.restore_project(acknowledged=True);self.assertEqual(len(restored.history),3);self.assertIsNone(restored.result);self.assertEqual(restored.page_layout_draft['policy'],self.policy)
  failed=model('failed')
  with patch('web.page_session.adopt',side_effect=OSError('synthetic storage failure')):
   with self.assertRaises(OSError):execute_local(failed,{},acknowledged=True)
  self.assertEqual(failed.history,[]);self.assertIsNone(failed.result);self.assertEqual(failed.generation_count,0)
 def test_automatic_dynamic_ref_field_isolation_and_unknown_scope(self):
  from core import dynamic
  pkg=Package(self.path);p=pkg.paragraphs[5];E.SubElement(p,q('fldSimple'),{q('instr'):' REF HiddenFigure '});path=self.root/'ref.docx';path.write_bytes(pkg.bytes(True));m=self.manifest(path);plan=dynamic.propose(path,m,complete=True,automatic=True)
  self.assertFalse(plan['edits']);self.assertTrue(plan['automatic_scope']['unresolved'])
  # Old manual interface still has its original opt-in requirement.
  with self.assertRaises(DemoError):dynamic.propose(path,m)
 def test_external_preview_and_draft_change_refuse_adoption(self):
  from web.page_session import preview,adopt
  x=WebSession(self.root/'external');x.set_mode('academic');x.load(self.path.read_bytes(),'合成.docx');x.set_types(self.roles);x.confirm_and_check();r=pages.automatic(Package(self.path),self.manifest(),self.policy);st={};preview(x,st,self.policy,r)
  Path(st['result']['output']).write_bytes(b'changed')
  with self.assertRaises(DemoError):adopt(x,st)
  self.assertFalse(x.history)
  st={};preview(x,st,self.policy,r);x.page_layout_draft={'enabled':True,'policy':deepcopy(self.policy),'input_sha256':x.inventory['input_sha256'],'regions':r}
  with self.assertRaises(DemoError):adopt(x,st)
  self.assertFalse(x.history)
 def test_text_and_requirement_files_use_exact_local_strategy_and_zero_parse_requests(self):
  from web.onepass_session import start,continue_command
  from web.requirements_input import file_input
  from docx import Document
  import io
  text='封面不显示页码，摘要和目录用小写罗马数字，正文从1开始连续编号。各章页脚显示章节名称和页码。整理图表编号并同步正文中的对应引用。'
  for kind in ('text','txt','docx'):
   x=WebSession(self.root/('requirements-'+kind));x.set_mode('academic');x.load(self.path.read_bytes(),'合成.docx');x.set_types(self.roles);x.confirm_classification();state={};kwargs={}
   if kind=='txt':kwargs['requirements']=file_input(x,text.encode(),'公开要求.txt')
   if kind=='docx':
    d=Document();d.add_paragraph(text);stream=io.BytesIO();d.save(stream);kwargs['requirements']=file_input(x,stream.getvalue(),'公开要求.docx')
   with patch('urllib.request.OpenerDirector.open',side_effect=AssertionError('no network')):
    start(x,state,text if kind=='text' else '',file_parser=lambda *_:self.fail('exact strategy must not request API'),spec_parser=lambda *_:self.fail('exact strategy must not request API'),structure_runner=lambda *_:self.fail('human scope no AI'),**kwargs)
   self.assertEqual(state['status'],'done',state.get('error'));self.assertEqual(x.current_rule_bundle()['assembly'],self.policy);self.assertEqual(state['audit']['API']['requirements']['requests'],0)
 def test_unknown_prefix_is_not_silently_classified_as_cover(self):
  from web.assembly_session import prepare,continue_regions
  pkg=Package(self.path);pkg.paragraphs[0].find('w:r/w:t',NS).text='用途尚未确定的内容';path=self.root/'unknown-prefix.docx';path.write_bytes(pkg.bytes(True))
  x=WebSession(self.root/'unknown-prefix');x.set_mode('academic');x.load(path.read_bytes(),'公开合成.docx');x.set_types({n:'pending' for n in self.roles});proposal={'draft':{'paragraph_types':deepcopy(x.paragraph_types)}};state={'assembly':self.policy,'audit':{},'proposal':proposal}
  questions=prepare(x,state,proposal);self.assertEqual({q['id'] for q in questions},{'page_body','page_front','page_cover'})
  continue_regions(x,state,{'page_body':'5','page_front':'3','page_cover':'keep'});self.assertEqual(state['page_regions'][0]['role'],'keep')
  with self.assertRaises(DemoError):pages.automatic(Package(path),self.manifest(path),self.policy,front_start=9,body_start=5)
 def test_upper_roman_explicit_restart_right_text_and_footer_conflicts(self):
  p={**self.policy,'roman':'upperRoman','body_start':5,'chapter_restart':True,'footer':'text','text':'合成研究','alignment':'right'};m=self.manifest();r=pages.automatic(Package(self.path),m,p);plan=pages.propose(self.path,m,p,r);out=self.root/'right.docx';pages.rehearse(self.path,out,m,plan)
  from core.front import section_inventory
  inv=section_inventory(Package(out));self.assertEqual([s['page_numbering'].get('start') for s in inv],[None,'1','5','5']);self.assertEqual(inv[1]['page_numbering']['fmt'],'upperRoman')
  root=E.fromstring(Package(out).parts[inv[2]['references']['footer:default']['part']]);self.assertTrue(root.xpath('./w:p/w:pPr/w:jc[@w:val="right"]',namespaces=NS));self.assertIn('合成研究',''.join(root.itertext()))
  # Existing mixed page/text or graphic footer is not cleared/rebuilt.
  pkg=Package(self.path);name=section_inventory(pkg)[0]['references']['footer:default']['part'];fr=E.fromstring(pkg.parts[name]);page=fr.findall(q('p'))[-1];rr=E.SubElement(page,q('r'));E.SubElement(rr,q('t')).text='非授权普通文字';pkg.parts[name]=E.tostring(fr)
  import io
  stream=io.BytesIO()
  with zipfile.ZipFile(stream,'w') as z:
   for n,v in pkg.parts.items():z.writestr(n,v)
  bad=self.root/'mixed-footer.docx';bad.write_bytes(stream.getvalue());original=bad.read_bytes()
  with self.assertRaises(DemoError):pages.propose(bad,self.manifest(bad),self.policy,pages.automatic(Package(bad),self.manifest(bad),self.policy))
  self.assertEqual(original,bad.read_bytes())
 def test_explicit_start_feedback_and_complex_footer_exception_continue_zero_API(self):
  from web.onepass_session import start,continue_command
  x=WebSession(self.root/'unknown-start');x.set_mode('academic');x.load(self.path.read_bytes(),'公开合成.docx');x.page_layout_draft={'enabled':True,'policy':self.policy,'input_sha256':x.inventory['input_sha256'],'regions':[]};st={}
  start(x,st,structure_runner=lambda *_:(_ for _ in ()).throw(DemoError('AI unavailable')));self.assertEqual(st['status'],'waiting');self.assertEqual([q['id'] for q in st['questions']],['page_body'])
  with patch('urllib.request.OpenerDirector.open',side_effect=AssertionError('continue must not request')):continue_command(x,st,'',{'page_body':'5'})
  self.assertEqual(st['status'],'done',st.get('error'));self.assertEqual(st['audit']['explicit_body_boundary_reanalysis'],5);self.assertEqual(x.current_round,0)
  # A real legal picture footer is preserved by refusing the structural step.
  from docx import Document
  from docx.shared import Inches
  d=Document(self.path);picture=next((Path(__file__).resolve().parents[1]/'samples/phase23').glob('*.png'),None)
  if picture is None:
   import base64
   picture=self.root/'pixel.png';picture.write_bytes(base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII='))
  d.sections[0].footer.add_paragraph().add_run().add_picture(str(picture),width=Inches(.2));path=self.root/'picture.docx';d.save(path)
  y=WebSession(self.root/'picture-session');y.set_mode('academic');y.load(path.read_bytes(),'公开图片页脚.docx');y.set_types(self.roles);y.confirm_classification();y.page_layout_draft={'enabled':True,'policy':self.policy,'input_sha256':y.inventory['input_sha256'],'regions':[]};state={};old=Package(path)
  start(y,state,structure_runner=lambda *_:self.fail('no structure request'));self.assertEqual(state['status'],'waiting');self.assertEqual([q['id'] for q in state['questions']],['page_conflict'])
  continue_command(y,state,'',{'page_conflict':'keep'});self.assertEqual(state['status'],'done',state.get('error'));self.assertIn('page_skipped',state['audit'])
  out=Package(y.result['output'])
  for n,v in old.parts.items():
   if n!='word/document.xml':self.assertEqual(v,out.parts[n])
if __name__=='__main__':unittest.main()
