"""Independent synthetic labels, original validators and transactions; no paid HTTP."""
from pathlib import Path
from tempfile import TemporaryDirectory
from copy import deepcopy
import unittest,io,json
from unittest.mock import patch
from docx import Document
from core.classification import suggest_paragraph_types
from core.package import DemoError,Package,NS
from core.page_policy import extract,DEFAULT,merge,validate
from web.session import WebSession
from web.structure_session import prepare,coordinate
from web.smart_session import suggest
from ai.semantic_spec import request_for,adapt,VERSION
ROOT=Path(__file__).resolve().parents[1]
class SharedStructureTests(unittest.TestCase):
 def setUp(self):self.tmp=TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.root=Path(self.tmp.name)
 def doc(self,rows):
  d=Document()
  for text,style in rows:d.add_paragraph(text,style)
  p=self.root/'source.docx';d.save(p);return p
 def model(self,path):
  s=WebSession(self.root/str(len(list(self.root.iterdir()))));s.set_mode('academic');s.load(path.read_bytes(),'公开合成.docx');return s
 def test_abs_without_keywords_and_en_direct_body(self):
  for marker,expected in [('摘要','abstract_cn_body'),('Abstract','abstract_en_body')]:
   p=self.doc([(marker,None),('这是摘要中明确的叙述文字。',None),('第1章 研究方法','Heading 1'),('正文叙述保持本来用途，不再延续此前摘要区域。',None),('第2章 结果','Heading 1'),('结果段落。',None)])
   types=[r['suggested_type'] for r in suggest_paragraph_types(p,include_sections=True)['paragraphs']]
   self.assertEqual(types,[expected.replace('_body','_title'),expected,'heading1','body','heading1','body'])
 def test_numbered_special_markers_end_main_zone_with_shared_evidence(self):
  for marker,title,body in [('第2章 参考文献','reference_title','reference_item'),('第3章 附录','appendix_title','appendix_text'),('第4章 致谢','ack_title','ack_text'),('第2章 目录','other','other')]:
   p=self.doc([('第1章 研究','Heading 1'),('这里是正文叙述。',None),(marker,'Heading 1'),('[1] 这是一段公开合成的完整长文字，用于验证特殊区域不会错误延续正文状态。',None)])
   rows=suggest_paragraph_types(p,include_sections=True)['paragraphs'];self.assertEqual(rows[2]['suggested_type'],title);self.assertEqual(rows[3]['suggested_type'],body);self.assertNotEqual(rows[3]['suggested_type'],'body')
  p=self.doc([('第1章 中文摘要','Heading 1'),('这是摘要普通文字。',None),('第1章 正文','Heading 1'),('正文叙述。',None)]);rows=suggest_paragraph_types(p,include_sections=True)['paragraphs'];self.assertEqual(rows[0]['suggested_type'],'abstract_cn_title');self.assertEqual(rows[1]['suggested_type'],'abstract_cn_body');self.assertEqual(rows[-1]['suggested_type'],'body')
 def test_directory_marker_without_abstract_supplies_same_front_boundary(self):
  from web.onepass_session import start
  from core.package import q
  pkg=Package(ROOT/'samples/phase48/existing.docx');pkg.paragraphs[2].find('w:r/w:t',NS).text='第2章 目录';path=self.root/'directory.docx';path.write_bytes(pkg.bytes(True));s=self.model(path);s.page_layout_draft={'enabled':True,'policy':{**DEFAULT,'pages':'front-body','footer':'chapter'},'regions':[],'input_sha256':s.inventory['input_sha256']};state={}
  start(s,state,structure_runner=lambda *_:(_ for _ in ()).throw(DemoError('no API')));self.assertEqual(state['status'],'done',state.get('error'));self.assertFalse(state['questions']);self.assertEqual(state['audit']['structure']['body_start'],5)
  sections=__import__('core.front',fromlist=['section_inventory']).section_inventory(Package(s.result['output']));self.assertEqual(sections[1]['page_numbering'].get('fmt'),'lowerRoman')
 def test_toc_lists_mentions_and_protected_boundaries(self):
  p=self.doc([('摘要',None),('摘要说明。',None),('目录',None),('第1章 研究方法......1','TOC Heading'),('2025年实验结果。','Heading 1'),('1 完成数据准备。',None),('本文在绪论中提及研究方法。',None),('第1章 研究方法','Heading 1'),('这里才是正文说明。',None)])
  rows=suggest_paragraph_types(p,include_sections=True)['paragraphs']
  self.assertFalse(any(r['structure_boundary'] for r in rows[:7]));self.assertTrue(rows[7]['structure_boundary']);self.assertEqual(rows[8]['suggested_type'],'body')
  from core.structure_evidence import heading
  for special in ('书签','修订','字段'):
   self.assertIsNone(heading({**rows[7],'special':special}))
 def test_one_body_decision_reclassifies_only_unconfirmed_and_keeps_exclusions(self):
  p=self.doc([('摘要',None),('摘要叙述。',None),('研究方法','Heading 1'),('这一部分实际是普通正文，具有完整叙述和足够的上下文长度。',None),('x = [1, 2];',None)])
  s=self.model(p);s.material_exclusions={'5':'人工排除程序代码'};proposal=suggest(s);coordinate(s,proposal,body_start=3,decided=['3']);proposal['draft']['paragraph_types']['3']='heading1'
  self.assertEqual(proposal['draft']['paragraph_types']['4'],'body');self.assertEqual(proposal['draft']['paragraph_types']['5'],'other');self.assertEqual(proposal['structure']['body_start'],3)
  original=deepcopy(s.paragraph_types);s.set_types({**s.paragraph_types,'4':'other'});proposal=suggest(s);coordinate(s,proposal,body_start=3);self.assertEqual(proposal['draft']['paragraph_types']['4'],'other');self.assertEqual(s.material_exclusions,{'5':'人工排除程序代码'})
 def test_shared_skeleton_budget_privacy_and_changed_scope(self):
  s=self.model(ROOT/'samples/phase44/errors.docx');p=prepare(s,include_tables=True);sk=p['request']['skeleton'];self.assertTrue(sk['context_only']);self.assertLessEqual(len(sk['anchors']),16);self.assertFalse(any('姓名' in x['text'] for x in sk['anchors']))
  from ai.structure_ids import wire_request,VERSION as STRUCTURE_VERSION
  self.assertEqual(wire_request(p['request'])['version'],STRUCTURE_VERSION)
  s.material_exclusions['2']='保护'
  from web.structure_session import verify
  with self.assertRaises(DemoError):verify(s,p)
 def test_equivalent_intents_target_only_explicit_sections_and_conflicts(self):
  p,rest,e=extract('封面不要页码，摘要和目录用i、ii，正文从1开始。');self.assertFalse(rest);self.assertEqual(p['targets'],['cover','front','body']);self.assertEqual(p['roman'],'lowerRoman')
  for text in ('正文从5开始，后面每章连续。','正文页码从5开始连续编号。'):
   p,rest,e=extract(text);self.assertFalse(rest);self.assertEqual(p['body_start'],5);self.assertEqual(p['targets'],['body'])
  p,rest,e=extract('保留现在的页码，只统一正文格式。');self.assertEqual(p['pages'],'keep');self.assertFalse(p['dynamic'])
  for text in ('正文从1开始；正文从5开始','保留现在的页码；正文从1开始'):
   with self.assertRaises(DemoError):extract(text)
  p,rest,e=extract('各章换不同页脚，显示章名，正文页码接着排。');self.assertFalse(rest);self.assertEqual(p['pages'],'chapter-footers');self.assertFalse(p['chapter_restart'])
  with self.assertRaises(DemoError):validate({**p,'targets':[{}]})
 def test_policy_targets_keep_unmentioned_and_original_transactions(self):
  from web.assembly_session import regions_from
  p,_,proof=extract('正文从5开始，后面每章连续。');s=self.model(ROOT/'samples/phase48/existing.docx');roles=json.loads((ROOT/'samples/phase48/roles.json').read_text());s.set_types(roles);s.confirm_classification()
  rows=regions_from(s,roles,p,body_start=5);self.assertEqual(rows[0]['role'],'keep');self.assertEqual(rows[1]['start'],5);self.assertEqual(rows[1]['restart'],5);self.assertEqual(rows[-1]['restart'],5 if len(rows)==2 else None)
  old={**DEFAULT,'pages':'front-body','dynamic':True,'roman':'upperRoman'};out=merge(old,p,proof);self.assertEqual(out['roman'],'upperRoman');self.assertTrue(out['dynamic'])
 def test_one_answer_shared_by_classification_pages_and_footer_no_second_request(self):
  from web.onepass_session import start,continue_command
  from core.package import q
  pkg=Package(ROOT/'samples/phase48/existing.docx')
  for n in (5,9):
   p=pkg.paragraphs[n-1];style=p.find('w:pPr/w:pStyle',NS);style.getparent().remove(style);p.find('w:r/w:t',NS).text='研究方法' if n==5 else '研究结果'
  path=self.root/'uncertain.docx';path.write_bytes(pkg.bytes(True));s=self.model(path);s.page_layout_draft={'enabled':True,'policy':{**DEFAULT,'pages':'front-body','footer':'chapter'},'regions':[],'input_sha256':s.inventory['input_sha256']};state={}
  start(s,state,structure_runner=lambda *_:(_ for _ in ()).throw(DemoError('no API')));self.assertEqual(state['status'],'waiting');self.assertEqual([q['id'] for q in state['questions']],['page_body'])
  with patch('urllib.request.OpenerDirector.open',side_effect=AssertionError('no new request')):continue_command(s,state,'',{'page_body':'5'})
  self.assertEqual(state['status'],'done',state.get('error'));self.assertEqual(state['audit']['structure']['body_start'],5);self.assertEqual(s.result['output_classification_manifest']['paragraph_types']['5'],'heading1');self.assertEqual(s.result['output_classification_manifest']['paragraph_types']['6'],'body')
 def test_earlier_unsettled_heading_and_unspecified_restart_are_not_guessed(self):
  p=self.doc([('摘要',None),('摘要说明。',None),('研究方法','Heading 1'),('这一段真实用途尚不明确，需要人决定该标题是否已经进入正文。',None),('第2章 结果','Heading 1'),('这里是结果说明。',None)])
  s=self.model(p);proposal=suggest(s);index=coordinate(s,proposal);self.assertIsNone(index['body_start']);self.assertTrue(index['conflicts'])
  from web.onepass_session import start,continue_command
  s=self.model(ROOT/'samples/phase48/existing.docx');state={};start(s,state,'前面用大写罗马，正文用阿拉伯数字重新编号。',structure_runner=lambda *_:(_ for _ in ()).throw(DemoError('no API')))
  self.assertEqual(state['status'],'waiting');self.assertEqual([q['id'] for q in state['questions']],['page_number_start']);self.assertFalse(s.history)
  continue_command(s,state,'前面用大写罗马，正文用阿拉伯数字重新编号。',{'page_number_start':'1'});self.assertEqual(state['status'],'done',state.get('error'));self.assertEqual(state['assembly']['body_start'],1)
 def test_semantic_page_ids_full_sources_and_native_format_validation(self):
  text='正文五号；正文页码起算值为5；各章页脚包含本章名称和页码。';req=request_for(text);fields={v['name']:k for k,v in req['fields'].items()};data={'version':VERSION,'items':[{'field':fields['font_size'],'value':10.5,'source':'S1'},{'field':'P03','value':5,'source':'S2'},{'field':'P05','value':'chapter','source':'S3'}],'issues':[]}
  r=adapt(data,text);self.assertEqual(r['draft']['fields']['font_size']['value'],10.5);self.assertEqual(r['page_strategy']['policy']['body_start'],5);self.assertEqual(r['page_strategy']['policy']['footer'],'chapter');self.assertEqual(r['page_strategy']['proof'][0]['quote'],'正文页码起算值为5')
  for change in ({'value':6},{'source':'S1'},{'value':'5'}):
   bad=deepcopy(data);bad['items'][1].update(change)
   with self.assertRaises(DemoError):adapt(bad,text)
  with self.assertRaises(DemoError):request_for(text+'；摘要润色一下')
  with self.assertRaises(DemoError):adapt({**data,'items':data['items'][:-1]},text)
 def test_explicit_preface_purposes_and_conflicting_declarations(self):
  from web.onepass_session import start
  p=self.doc([('摘要',None),('摘要说明。',None),('前言','Heading 1'),('这里是前言的普通叙述，在不同任务中用途由原要求明确指定。',None),('第1章 研究','Heading 1'),('这里是正文说明。',None)])
  for purpose,expected in [('正文',3),('前置部分',5)]:
   s=self.model(p);state={};start(s,state,'前言属于'+purpose,structure_runner=lambda *_:(_ for _ in ()).throw(DemoError('local fallback')))
   self.assertEqual(state['status'],'done',state.get('error'));self.assertEqual(state['audit']['structure']['body_start'],expected)
  with self.assertRaises(DemoError):extract('前言属于正文；前言属于前置部分')
 def test_plain_footer_text_and_alignment_share_same_request(self):
  text='页脚写“合成研究”并附页码；页脚靠右。';req=request_for(text);self.assertEqual(set(req['fields']),{'P08','P06'})
  result=adapt({'version':VERSION,'items':[{'field':'P08','value':'合成研究','source':'S1'},{'field':'P06','value':'right','source':'S2'}],'issues':[]},text)
  self.assertEqual(result['page_strategy']['policy']['text'],'合成研究');self.assertEqual(result['page_strategy']['policy']['alignment'],'right');self.assertEqual(result['page_strategy']['policy']['pages'],'chapter-footers')
  only,_,_=extract('页脚居右');self.assertEqual(only['pages'],'chapter-footers')
  with self.assertRaises(DemoError):adapt({'version':VERSION,'items':[{'field':'P08','value':'模型编造文字','source':'S1'},{'field':'P06','value':'right','source':'S2'}],'issues':[]},text)
 def test_alignment_alone_does_not_authorize_new_page_fields(self):
  from web.onepass_session import start,continue_command
  p=self.doc([('第1章 研究','Heading 1'),('这是明确的正文叙述，原稿没有任何页脚或PAGE域。',None)])
  s=self.model(p);state={};start(s,state,'页脚居右',structure_runner=lambda *_:(_ for _ in ()).throw(DemoError('no API')))
  self.assertEqual(state['status'],'waiting');self.assertEqual([q['id'] for q in state['questions']],['page_conflict'])
  continue_command(s,state,'页脚居右',{'page_conflict':'keep'});self.assertEqual(state['status'],'done',state.get('error'));self.assertFalse(any('footerPF' in n for n in Package(s.result['output']).parts))
 def test_declined_partial_file_cannot_leave_its_body_boundary(self):
  from web.onepass_session import start,continue_command
  from web.requirements_input import file_input
  from ai.smart_spec import VERSION as NATIVE
  p=self.doc([('摘要',None),('摘要叙述。',None),('前言','Heading 1'),('前言中这一段不应因已放弃的文件要求而套用正文格式。',None),('第1章 研究','Heading 1'),('明确正文说明。',None)])
  s=self.model(p);state={};requirements=file_input(s,'前言属于正文；摘要润色一下'.encode(),'合成要求.txt')
  def parser(ex):return {'draft':{'version':NATIVE,'fields':{},'issues':[]},'segments':ex['segments'],'notices':[{'location':ex['segments'][0]['location'],'text':'摘要润色一下','reason':'不支持内容改写'}],'receipt':{'kind':'mock_api','model':'deepseek-flash','timestamp_utc':'2026-01-01T00:00:00Z','request_count':1}}
  start(s,state,requirements=requirements,file_parser=parser,structure_runner=lambda *_:(_ for _ in ()).throw(DemoError('no API')))
  self.assertEqual(state['status'],'waiting');self.assertEqual([q['id'] for q in state['questions']],['requirements_partial']);continue_command(s,state,'',{'requirements_partial':'template'},requirements=requirements)
  self.assertEqual(state['status'],'done',state.get('error'));self.assertTrue(state['audit']['file_boundary_decisions_discarded']);self.assertEqual(s.result['output_classification_manifest']['paragraph_types']['3'],'pending')
 def test_keep_page_and_body_only_does_not_modify_titles_or_paper(self):
  from web.onepass_session import start
  from lxml import etree as E
  from core.package import q
  source=ROOT/'samples/phase48/existing.docx';s=self.model(source);state={};start(s,state,'保留现在的页码，只统一正文格式',structure_runner=lambda *_:(_ for _ in ()).throw(DemoError('no API')))
  self.assertEqual(state['status'],'done',state.get('error'));self.assertEqual(state['audit']['format_roles'],['body']);before=Package(source);after=Package(s.result['output'])
  for n in (1,2,3,4,5,7,9,11):self.assertEqual(E.tostring(before.paragraphs[n-1]),E.tostring(after.paragraphs[n-1]))
  for a,b in zip(before.sections,after.sections):self.assertEqual(E.tostring(a),E.tostring(b))
  for name in before.parts:
   if name!='word/document.xml':self.assertEqual(before.parts[name],after.parts[name])
 def test_semantic_file_page_ids_and_entire_input_conflicts(self):
  from core.norms import extract_norm
  from ai.file_spec import request_for as freq,adapt as fadapt,VERSION as FV
  p=self.root/'要求.txt';p.write_text('正文页码起算值为5；各章页脚包含章名及页号。')
  ex=extract_norm(p.read_bytes(),'txt');req=freq(ex);result=fadapt({'version':FV,'items':[{'field':'P03','value':5,'source':'S1'},{'field':'P05','value':'chapter','source':'S2'}],'issues':[],'background':[]},ex);self.assertFalse(result['notices']);self.assertEqual(result['page_strategy']['policy']['body_start'],5)
  from ai.page_intent import compile
  with self.assertRaises(DemoError):compile([{'id':'P03','value':5,'source_text':'正文页码起算值为5'}],'正文从1开始；正文页码起算值为5')
if __name__=='__main__':unittest.main()
