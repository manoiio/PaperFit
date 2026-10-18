from pathlib import Path
from tempfile import TemporaryDirectory
from copy import deepcopy
import json,unittest
from lxml import etree as E
from unittest.mock import patch
from core.package import Package,q,NS,DemoError,sha256
from core.formats import Resolver,Unresolved
from core.text_effects import read,value,candidate
from core.engine import check_document,fix_document,load_rules
from core import rule_bundle as rb
from web.session import WebSession
from web.onepass_session import start
from ai.semantic_spec import adapt,request_for,parse_requirements
from ai.short_spec import Settings
S=Path(__file__).resolve().parents[1]/'samples/phase53';TEXT=(S/'requirements.txt').read_text().strip();EXPECTED=json.loads((S/'expected.json').read_text())['fields'];RESPONSE=json.loads((S/'public-response.json').read_text())
def parser(t):return parse_requirements(t,config=Settings('mock'),transport=lambda *_:json.dumps({'choices':[{'finish_reason':'stop','message':{'content':json.dumps(RESPONSE,ensure_ascii=False)}}]}))
def model(folder):
 s=WebSession(folder);s.set_mode('academic');s.load((S/'source.docx').read_bytes(),'公开合成.docx');s.set_types(json.loads((S/'roles.json').read_text()));s.confirm_classification();return s
class VisualNormalization(unittest.TestCase):
 def test_semantic_proof_units_conflicts_and_no_defaults(self):
  d=adapt(RESPONSE,TEXT);self.assertEqual({n:v['value'] for n,v in d['fields'].items()},EXPECTED)
  self.assertFalse(any(n in rb.flat('main',rb.defaults(load_rules())['main']) and rb.flat('main',rb.defaults(load_rules())['main'])[n] is not None for n in EXPECTED if '__' in n))
  for n,v in [('body__shadow',0),('body__scale',True),('body__scale',99.5),('body__character_spacing',.03),('body__baseline',.1)]:
   with self.assertRaises(DemoError):value(n,v)
  from ai.smart_spec import coverage
  for t in ('全文清除格式','正文正常','正文艺术字','标题不使用阴影'):
   with self.assertRaises(DemoError):coverage(t)
  broken=deepcopy(RESPONSE);broken['items'][0]['source']='S99'
  with self.assertRaises(DemoError):adapt(broken,TEXT)
 def test_effects_real_transaction_exact_mask_and_direct_path(self):
  with TemporaryDirectory() as d:
   s=model(Path(d)/'smart');raw=s.input_path.read_bytes();state={};start(s,state,TEXT,spec_parser=parser,structure_runner=lambda *_:self.fail('No AI for confirmed scope'))
   self.assertEqual(state['status'],'done',state.get('error'));self.assertEqual(s.input_path.read_bytes(),raw)
   output=Package(s.result['output']);source=Package(S/'source.docx');r=output.paragraphs[5].find(q('r'));rp=r.find(q('rPr'))
   for tag,v in [('shadow','0'),('outline','0'),('w','100'),('spacing','0'),('position','0')]:self.assertEqual(rp.find(q(tag)).get(q('val')),v)
   self.assertEqual(output.paragraphs[6].find('w:r/w:rPr/w:emboss',NS).get(q('val')),'0');self.assertEqual(output.paragraphs[7].find('w:r/w:rPr/w:imprint',NS).get(q('val')),'0')
   self.assertEqual(E.tostring(source.paragraphs[5].findall(q('r'))[-1]),E.tostring(output.paragraphs[5].findall(q('r'))[-1]))
   self.assertEqual(source.parts['word/numbering.xml'],output.parts['word/numbering.xml']);self.assertEqual(source.parts['word/styles.xml'],output.parts['word/styles.xml'])
   self.assertEqual(source.root.xpath('//w:t/text()',namespaces=NS),output.root.xpath('//w:t/text()',namespaces=NS))
   for n in source.parts:
    if n!='word/document.xml':self.assertEqual(source.parts[n],output.parts[n])
   # Same public engine selection, not the new adapter's own assertion.
   direct=fix_document(s.input_path,Path(d)/'direct.docx',s.confirmed,selected_keys=state['audit']['selection'] if 'selection' in state['audit'] else [c['key'] for c in s.result['changes']],**s._base_rules_arguments(),mode=s.mode,classification_manifest=s.classification_confirmed)
   self.assertEqual(Path(direct['output']).read_bytes(),s.downloads()['fixed.docx'])
   self.assertTrue(s.result['preservation']['other_zip_member_bytes_equal'])
 def test_single_effect_only_and_project_bundle_recovery(self):
  with TemporaryDirectory() as d:
   s=model(Path(d)/'a');s.start_workbench_draft();s.workbench_policy('keep',acknowledged=True);s.workbench_edit('main','body__shadow','manual',False);s.apply_workbench(acknowledged=True);s.set_types(json.loads((S/'roles.json').read_text()));s.confirm_and_check()
   fs=[f for f in s.eligible() if f['locator'].get('property')=='text_shadow'];self.assertTrue(fs);s.set_selection([fs[0]['key']]);self.assertTrue(s.generate(),s.error)
   old=Package(s.input_path);new=Package(s.result['output']);p=fs[0]['locator']['paragraph']-1;i=fs[0]['locator']['run']-1;rr=new.paragraphs[p].findall(q('r'))[i].find(q('rPr'));self.assertIsNone(rr.find(q('outline')));self.assertIsNone(rr.find(q('w')))
   t=model(Path(d)/'b');t.import_rule_bundle(s.rule_bundle_download());t.apply_workbench(acknowledged=True);self.assertEqual(t.current_rule_bundle()['modules'],s.current_rule_bundle()['modules'])
   t.stage_project(s.prepare_project());restored=t.restore_project(acknowledged=True);self.assertEqual(restored.current_rule_bundle()['modules'],s.current_rule_bundle()['modules']);self.assertEqual(restored.report['input_sha256'],restored.inventory['input_sha256']);self.assertEqual(restored.preview_cache,{})
 def test_duplicate_properties_and_superscript_protected(self):
  p=Package(S/'source.docx');r=Resolver(p);x=p.paragraphs[5];layers=r.layers(x,x.find(q('r')));self.assertTrue(read(r,layers,'shadow')['value'])
  extra=deepcopy(layers[-2] if len(layers)>1 else layers[0]);E.SubElement(extra,q('shadow'));E.SubElement(extra,q('shadow'))
  with self.assertRaises(Unresolved):read(r,[extra],'shadow')
 def test_theme_slot_japanese_and_same_layer_override(self):
  p=Package(S/'source.docx');r=Resolver(p);x=p.paragraphs[14];layers=r.layers(x,x.find(q('r')));info=r.font_info(layers,'eastAsia');self.assertEqual(info['theme_basis']['script'],'Jpan');self.assertEqual(info['font'],'ＭＳ ゴシック')
  from core.engine import _ensure
  rf=_ensure(_ensure(x.find(q('r')),'rPr'),'rFonts');rf.set(q('eastAsia'),'黑体');self.assertEqual(r.font_info(r.layers(x,x.find(q('r'))),'eastAsia')['font'],'黑体')
 def test_txt_docx_sources_and_partial_file_cannot_hide_effect_requirements(self):
  from core.norms import extract_norm
  from ai.file_spec import request_for as request,adapt as file_adapt,background
  from docx import Document
  from io import BytesIO
  doc=Document();doc.add_paragraph(TEXT);buf=BytesIO();doc.save(buf)
  for raw,kind in ((TEXT.encode(),'txt'),(buf.getvalue(),'docx')):
   with TemporaryDirectory() as td:
    path=Path(td)/'requirements.docx';path.write_bytes(raw)
    extraction=extract_norm(raw,kind,path=path if kind=='docx' else None)
   r=request(extraction)
   items=[]
   for fid,field in r['fields'].items():
    n=field['name']
    if n not in EXPECTED:continue
    obj='正文' if n.startswith('body__') else '二级标题' if n.startswith('heading2') else '三级标题'
    items.append({'field':fid,'value':EXPECTED[n],'source':next(x['id'] for x in r['sources'] if obj in x['text'])})
   out=file_adapt({'version':r['version'],'items':items,'background':[],'issues':[]},extraction)
   self.assertEqual({n:f['value'] for n,f in out['draft']['fields'].items()},EXPECTED)
   self.assertTrue(all(('TXT:L1' if kind=='txt' else 'DOCX:P1') in f['location'] for f in out['draft']['fields'].values()))
  self.assertFalse(background('本规范不使用阴影，用于测试'))
  self.assertFalse(background('本规范艺术字用于测试'))
 def test_old_editor_optional_fields_restore_as_disabled_no_authorization(self):
  from core.text_effects import FIELDS
  with TemporaryDirectory() as d:
   s=model(Path(d));s.start_workbench_draft();old=deepcopy(s.workbench)
   for n in FIELDS:old['draft']['fields']['main'].pop(n)
   normalized=rb.validate_state(old,load_rules())
   self.assertEqual({n:f for n,f in normalized['draft']['fields']['main'].items() if n not in FIELDS},old['draft']['fields']['main'])
   self.assertTrue(all(normalized['draft']['fields']['main'][n]=={'mode':'keep','value':None,'candidate':None} for n in FIELDS))
   self.assertFalse(any(n in old['draft']['fields']['main'] for n in FIELDS))
   broken=deepcopy(old);broken['draft']['fields']['main'].pop('font_size')
   with self.assertRaises(DemoError):rb.validate_state(broken,load_rules())
 def test_position_never_masks_hidden_text_or_other_special_protection(self):
  with TemporaryDirectory() as d:
   p=Package(S/'source.docx');r=p.paragraphs[5].find(q('r'));E.SubElement(r.find(q('rPr')),q('vanish'))
   s=WebSession(Path(d));s.set_mode('academic');s.load(p.bytes(True),'保护合成.docx');s.start_workbench_draft();s.workbench_policy('keep',acknowledged=True)
   for prop,v in [('shadow',False),('baseline',0)]:s.workbench_edit('main','body__'+prop,'manual',v)
   s.apply_workbench(acknowledged=True);s.set_types(json.loads((S/'roles.json').read_text()));self.assertTrue(s.confirm_and_check(),s.error)
   fs=[f for f in s.report['findings'] if f['locator'].get('paragraph')==6 and f['locator'].get('property','').startswith('text_') and f['locator'].get('run')==1]
   self.assertEqual(len(fs),2);self.assertTrue(all(f['status']=='review' and not f['fixable'] for f in fs));self.assertIn('隐藏',fs[0]['detail'])
 def test_mutually_exclusive_effects_never_guessed_or_enabled_together(self):
  from core.text_effects import validate_requested
  with self.assertRaises(DemoError):validate_requested({'body__shadow':True,'body__emboss':True})
  validate_requested({'body__shadow':True,'body__outline':True})
  with TemporaryDirectory() as d:
   s=WebSession(Path(d));s.set_mode('academic');s.load((S/'conflicting-effects.docx').read_bytes(),'互斥反例.docx');s.start_workbench_draft();s.workbench_policy('keep',acknowledged=True);s.workbench_edit('main','body__shadow','manual',False);s.apply_workbench(acknowledged=True);s.set_types(json.loads((S/'roles.json').read_text()));self.assertTrue(s.confirm_and_check(),s.error)
   f=next(f for f in s.report['findings'] if f['locator'].get('paragraph')==6 and f['locator'].get('run')==1 and f['locator'].get('property')=='text_shadow');self.assertEqual(f['status'],'review');self.assertFalse(f['fixable']);self.assertIn('互斥',f['detail'])
if __name__=='__main__':unittest.main()
