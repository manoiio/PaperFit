"""独立错误清单 + XML/ZIP对象比对 + 选择性修正 + 危险边界。运行不覆盖样本。"""
import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from copy import deepcopy
from pathlib import Path
from lxml import etree as ET
from docx import Document
from core import check_document, fix_document, inspect_document, verify_document
from core.engine import IDS, _ensure
from core.package import DemoError, Package, q, NS

ROOT = Path(__file__).resolve().parents[1]
ERRORS = ROOT / 'samples/errors.docx'
CORRECT = ROOT / 'samples/correct.docx'
BODY = [4, 5, 6, 11, 14, 15]

def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def item(report, rule, **loc):
    return next(f for f in report['findings'] if f['rule_id'] == rule and f['locator'] == loc)

class Acceptance(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='acceptance-')
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def variant(self, mutate, base=CORRECT, styles=False, extra=None):
        p = Package(base)
        mutate(p)
        parts = p.parts.copy()
        parts['word/document.xml'] = ET.tostring(p.root, xml_declaration=True, encoding='UTF-8')
        if styles:
            parts['word/styles.xml'] = ET.tostring(p.styles, xml_declaration=True, encoding='UTF-8')
        parts.update(extra or {})
        path = self.dir / 'variant.docx'
        with zipfile.ZipFile(path, 'w') as z:
            for name, data in parts.items():
                z.writestr(name, data)
        return path

    def test_exact_independent_error_locations_and_expected_values(self):
        report = check_document(ERRORS, ROOT / 'samples/errors.body.json')
        expected = json.loads((ROOT / 'samples/expected_errors.json').read_text())['errors']
        actual = [f for f in report['findings'] if f['status'] == 'problem']
        project = lambda fs: sorted(json.dumps({k: f[k] for k in ('rule_id', 'locator', 'expected')}, sort_keys=True) for f in fs)
        self.assertEqual(project(actual), project(expected))
        self.assertEqual(report['summary']['rule_count'], 7)
        self.assertEqual(report['summary']['problem_count'], 11)
        self.assertEqual(report['summary']['review_count'], 0)
        self.assertEqual(item(report, 'R06', paragraph=4)['actual']['multiple'], 1.2)
        self.assertEqual(item(report, 'R06', paragraph=5)['actual']['line_rule'], 'exact')
        self.assertEqual(item(report, 'R05', paragraph=4, run=2)['actual']['pt'], 10.5)
        self.assertEqual(item(report, 'R03', paragraph=5, run=1)['actual']['eastAsia'], '黑体')
        self.assertEqual(item(report, 'R04', paragraph=14, run=2)['actual']['ascii'], 'Arial')
        self.assertEqual(item(report, 'R05', paragraph=14, run=2)['actual']['pt'], 14)

    def test_correct_character_and_equivalent_point_indents_no_false_positive(self):
        report = check_document(CORRECT, BODY)
        self.assertEqual(report['summary']['problem_count'], 0)
        self.assertEqual(report['summary']['review_count'], 0)
        self.assertEqual(item(report, 'R07', paragraph=6)['status'], 'pass')
        self.assertEqual(item(report, 'R07', paragraph=11)['actual']['first_line_pt'], 24)
        self.assertEqual(item(report, 'R07', paragraph=11)['status'], 'pass')
        self.assertEqual(report['summary']['conclusion'], '已覆盖检查未发现问题')
        self.assertEqual(report['coverage']['body_paragraphs'], BODY)
        self.assertEqual(inspect_document(CORRECT)['section_count'], 2)

    def test_all_four_margins_and_second_section_paper(self):
        def mutate(pkg):
            e = pkg.sections[1].find(q('pgSz'))
            e.set(q('h'), '15000')
            e = pkg.sections[1].find(q('pgMar'))
            for a in ('top', 'bottom', 'left', 'right'):
                e.set(q(a), '0')
        path = self.variant(mutate)
        report = check_document(path, BODY)
        self.assertEqual(item(report, 'R01', section=2)['status'], 'problem')
        self.assertEqual(item(report, 'R02', section=2)['status'], 'problem')
        out = fix_document(path, self.dir / 'fixed.docx', BODY, ['R01', 'R02'])
        self.assertEqual(item(out['after'], 'R01', section=2)['status'], 'pass')
        self.assertEqual(item(out['after'], 'R02', section=2)['status'], 'pass')

    def test_all_selected_fix_reread_original_hash_and_idempotence(self):
        original_hash = digest(ERRORS)
        path = self.dir / 'fixed.docx'
        result = fix_document(ERRORS, path, BODY, IDS)
        self.assertEqual(len(result['changes']), 11)
        self.assertTrue(all(c['verification'] == 'pass' for c in result['changes']))
        self.assertEqual(verify_document(path, BODY)['summary']['problem_count'], 0)
        self.assertEqual(digest(ERRORS), original_hash)
        manifest = result['output_body_manifest']
        self.assertEqual(check_document(path, manifest)['summary']['problem_count'], 0)
        path2 = self.dir / 'twice.docx'
        second = fix_document(path, path2, manifest, IDS)
        self.assertEqual(second['changes'], [])
        self.assertEqual(path2.read_bytes(), path.read_bytes())

    def test_partial_selection_keeps_all_unselected_format_attributes(self):
        out = self.dir / 'only_size.docx'
        result = fix_document(ERRORS, out, BODY, ['R05'])
        self.assertEqual(len(result['changes']), 2)
        self.assertEqual(result['after']['summary']['problem_count'], 9)
        before, after = Package(ERRORS), Package(out)
        for pkg in (before, after):
            for p in pkg.paragraphs:
                for e in p.findall('w:r/w:rPr/w:sz', NS):
                    e.getparent().remove(e)
            for e in pkg.root.findall('.//w:rPr', NS):
                if not len(e) and not e.attrib:
                    e.getparent().remove(e)
        self.assertEqual(ET.tostring(before.root, method='c14n'), ET.tostring(after.root, method='c14n'))
        for rule in ('R01', 'R02', 'R03', 'R04', 'R06', 'R07'):
            def formats(report):
                fs = deepcopy([f for f in report['findings'] if f['rule_id'] == rule])
                if rule == 'R07':
                    for f in fs:
                        # Derived diagnostic changes after fixing R05; indent XML itself is unchanged.
                        f['actual'].pop('uniform_size_pt', None)
                return fs
            self.assertEqual(formats(result['before']), formats(result['after']))

    def test_individual_problem_key_does_not_fix_other_run(self):
        key = item(check_document(ERRORS, BODY), 'R05', paragraph=4, run=2)['key']
        result = fix_document(ERRORS, self.dir / 'one.docx', BODY, selected_keys=[key])
        self.assertEqual(len(result['changes']), 1)
        self.assertEqual(item(result['after'], 'R05', paragraph=14, run=2)['status'], 'problem')

    def test_text_table_image_numbering_and_sections_preserved_independently(self):
        out = self.dir / 'fixed.docx'
        fix_document(ERRORS, out, BODY, IDS)
        a, b = Package(ERRORS), Package(out)
        self.assertEqual(a.root.xpath('//w:t/text()', namespaces=NS), b.root.xpath('//w:t/text()', namespaces=NS))
        for tag in ('tbl', 'drawing', 'numPr'):
            get = lambda pkg: [ET.tostring(e, method='c14n') for e in pkg.root.iter(q(tag))]
            self.assertEqual(get(a), get(b))
        self.assertEqual(len(a.root.findall('.//w:drawing', NS)), 1)
        self.assertEqual(len(a.root.findall('.//w:tbl', NS)), 1)
        media = [n for n in a.parts if n.startswith('word/media/')]
        self.assertEqual(len(media), 1)
        for n in a.parts:
            if n != 'word/document.xml':
                self.assertEqual(a.parts[n], b.parts[n], n)
        # Section properties outside paper/margins must be identical, with same break locations.
        for pkg in (a, b):
            for s in pkg.sections:
                for tag in ('pgSz', 'pgMar'):
                    s.remove(s.find(q(tag)))
        self.assertEqual([ET.tostring(s, method='c14n') for s in a.sections], [ET.tostring(s, method='c14n') for s in b.sections])
        d = Document(out)
        self.assertEqual([p.text for p in d.paragraphs], [p.text for p in Document(ERRORS).paragraphs])
        self.assertEqual([[c.text for row in t.rows for c in row.cells] for t in d.tables],
                         [[c.text for row in t.rows for c in row.cells] for t in Document(ERRORS).tables])

    def test_unconfirmed_paragraph_and_table_fonts_never_fixed(self):
        out = self.dir / 'scoped.docx'
        result = fix_document(ERRORS, out, [4], IDS)
        a, b = Package(ERRORS), Package(out)
        for index in (1, 2, 3, 5, 6, 7, 8, 9, 10, 11, 13, 14, 15):
            self.assertEqual(ET.tostring(a.paragraphs[index - 1]), ET.tostring(b.paragraphs[index - 1]))
        self.assertIn(5, result['after']['coverage']['unconfirmed_paragraphs'])

    def test_direct_theme_unknown_never_passed_or_force_fixed(self):
        def mutate(pkg):
            rp = _ensure(pkg.paragraphs[3].find(q('r')), 'rPr')
            rf = _ensure(rp, 'rFonts')
            rf.set(q('eastAsia'), '宋体')
            rf.set(q('eastAsiaTheme'), 'minorEastAsia')
        path = self.variant(mutate)
        report = check_document(path, [4])
        finding = item(report, 'R03', paragraph=4, run=1)
        self.assertEqual(finding['status'], 'review')
        self.assertFalse(finding['fixable'])
        result = fix_document(path, self.dir / 'fixed.docx', [4], ['R03'])
        self.assertEqual(result['changes'], [])
        self.assertEqual(path.read_bytes(), (self.dir / 'fixed.docx').read_bytes())
        self.assertIn(finding['key'], result['skipped_review_keys'])

    def test_docdefaults_and_direct_override_of_inherited_theme(self):
        def mutate(pkg):
            normal = pkg.styles.find("w:style[@w:styleId='Normal']", NS)
            rp = normal.find(q('rPr'))
            normal.remove(rp)
            defaults = pkg.styles.find('w:docDefaults/w:rPrDefault/w:rPr', NS)
            rf = _ensure(defaults, 'rFonts')
            rf.set(q('eastAsiaTheme'), 'minorEastAsia')
            _ensure(defaults, 'sz').set(q('val'), '21')
            rp = _ensure(pkg.paragraphs[3].find(q('r')), 'rPr')
            _ensure(rp, 'rFonts').set(q('eastAsia'), '宋体')
        path = self.variant(mutate, styles=True)
        report = check_document(path, [4])
        self.assertEqual(item(report, 'R03', paragraph=4, run=1)['status'], 'pass')
        self.assertEqual(item(report, 'R03', paragraph=4, run=2)['status'], 'review')
        self.assertEqual(item(report, 'R05', paragraph=4, run=1)['actual']['pt'], 10.5)

    def test_ascii_and_hansi_are_independent(self):
        def mutate(pkg):
            r = pkg.paragraphs[13].findall(q('r'))[1]
            r.find(q('t')).text = 'Latin café 2026'
            rf = _ensure(_ensure(r, 'rPr'), 'rFonts')
            rf.set(q('ascii'), 'Times New Roman')
            rf.set(q('hAnsi'), 'Arial')
        path = self.variant(mutate)
        f = item(check_document(path, [14]), 'R04', paragraph=14, run=2)
        self.assertEqual(f['actual'], {'ascii': 'Times New Roman', 'hAnsi': 'Arial'})
        result = fix_document(path, self.dir / 'fixed.docx', [14], ['R04'])
        self.assertEqual(item(result['after'], 'R04', paragraph=14, run=2)['status'], 'pass')
        rf = Package(self.dir / 'fixed.docx').paragraphs[13].findall(q('r'))[1].find('w:rPr/w:rFonts', NS)
        self.assertEqual(rf.get(q('hAnsi')), 'Times New Roman')

    def test_character_and_point_conflict_fixed_without_hanging_properties(self):
        def mutate(pkg):
            ind = _ensure(_ensure(pkg.paragraphs[5], 'pPr'), 'ind')
            ind.set(q('firstLineChars'), '200')
            ind.set(q('firstLine'), '960')
            ind.set(q('hanging'), '0')
        path = self.variant(mutate)
        self.assertEqual(item(check_document(path, [6]), 'R07', paragraph=6)['status'], 'problem')
        result = fix_document(path, self.dir / 'fixed.docx', [6], ['R07'])
        self.assertEqual(item(result['after'], 'R07', paragraph=6)['status'], 'pass')
        ind = Package(self.dir / 'fixed.docx').paragraphs[5].find('w:pPr/w:ind', NS)
        self.assertEqual(ind.get(q('firstLineChars')), '200')
        self.assertIsNone(ind.get(q('hanging')))
        self.assertIsNone(ind.get(q('hangingChars')))

    def test_two_chars_not_hardcoded_48pt(self):
        def mutate(pkg):
            p = pkg.paragraphs[10]
            _ensure(_ensure(p.find(q('r')), 'rPr'), 'sz').set(q('val'), '28')
            _ensure(_ensure(p, 'pPr'), 'ind').set(q('firstLine'), '560')
        path = self.variant(mutate)
        self.assertEqual(item(check_document(path, [11]), 'R07', paragraph=11)['status'], 'pass')

    def test_inherited_hanging_and_mixed_point_sizes_are_review(self):
        def mutate(pkg):
            normal = pkg.styles.find("w:style[@w:styleId='Normal']", NS)
            _ensure(_ensure(normal, 'pPr'), 'ind').set(q('hanging'), '240')
        path = self.variant(mutate, styles=True)
        f = item(check_document(path, [6]), 'R07', paragraph=6)
        self.assertEqual(f['status'], 'review')
        self.assertFalse(f['fixable'])

    def test_point_indent_with_mixed_sizes_or_width_changes_is_review(self):
        for mode in ('mixed', 'width'):
            with self.subTest(mode=mode):
                def mutate(pkg):
                    p = pkg.paragraphs[10]
                    r = p.find(q('r'))
                    if mode == 'mixed':
                        second = deepcopy(r)
                        _ensure(_ensure(second, 'rPr'), 'sz').set(q('val'), '28')
                        p.append(second)
                    else:
                        _ensure(_ensure(r, 'rPr'), 'w').set(q('val'), '80')
                path = self.variant(mutate)
                self.assertEqual(item(check_document(path, [11]), 'R07', paragraph=11)['status'], 'review')

    def test_special_paragraph_structures_skipped_and_xml_retained(self):
        for tag in ('hyperlink', 'ins', 'fldSimple', 'sdt'):
            with self.subTest(tag=tag):
                def mutate(pkg):
                    p = pkg.paragraphs[3]
                    wrapper = ET.Element(q(tag))
                    r = p.findall(q('r'))[-1]
                    p.remove(r)
                    wrapper.append(r)
                    p.append(wrapper)
                path = self.variant(mutate, base=ERRORS)
                report = check_document(path, [4])
                self.assertTrue(all(f['status'] == 'review' and not f['fixable'] for f in report['findings'] if f['locator'] == {'paragraph': 4}))
                out = self.dir / (tag + '.docx')
                result = fix_document(path, out, [4], IDS)
                self.assertEqual(ET.tostring(Package(path).paragraphs[3]), ET.tostring(Package(out).paragraphs[3]))
                self.assertTrue(result['skipped_review_keys'])

    def test_equation_and_field_paragraphs_are_not_force_formatted(self):
        for kind in ('equation', 'field'):
            with self.subTest(kind=kind):
                def mutate(pkg):
                    p = pkg.paragraphs[3]
                    if kind == 'equation':
                        p.append(ET.Element('{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath'))
                    else:
                        _ensure(p.find(q('r')), 'fldChar').set(q('fldCharType'), 'begin')
                path = self.variant(mutate)
                report = check_document(path, [4])
                self.assertEqual(item(report, 'R05', paragraph=4)['status'], 'review')
                self.assertEqual(fix_document(path, self.dir / (kind + '.docx'), [4], IDS)['changes'], [])

    def test_superscript_skipped_and_paragraph_format_changes_blocked(self):
        def mutate(pkg):
            r = pkg.paragraphs[3].findall(q('r'))[1]
            _ensure(_ensure(r, 'rPr'), 'vertAlign').set(q('val'), 'superscript')
        path = self.variant(mutate, base=ERRORS)
        before = Package(path).paragraphs[3].findall(q('r'))[1]
        report = check_document(path, [4])
        self.assertFalse(any(f['locator'] == {'paragraph': 4, 'run': 2} for f in report['findings']))
        self.assertEqual(item(report, 'R06', paragraph=4)['status'], 'review')
        self.assertGreater(report['summary']['review_count'], 0)
        out = self.dir / 'fixed.docx'
        fix_document(path, out, [4], IDS)
        after = Package(out).paragraphs[3].findall(q('r'))[1]
        self.assertEqual(ET.tostring(before), ET.tostring(after))

    def test_numbering_indent_and_document_grid_review(self):
        def mutate(pkg):
            _ensure(_ensure(pkg.paragraphs[3], 'pPr'), 'numPr')
            _ensure(pkg.sections[0], 'docGrid').set(q('type'), 'lines')
        path = self.variant(mutate)
        report = check_document(path, [4])
        self.assertEqual(item(report, 'R06', paragraph=4)['status'], 'review')
        self.assertEqual(item(report, 'R07', paragraph=4)['status'], 'review')

    def test_style_cycle_is_review_not_infinite_loop(self):
        def mutate(pkg):
            s = pkg.styles.find("w:style[@w:styleId='DemoBodyBase']", NS)
            s.find(q('basedOn')).set(q('val'), 'DemoBody')
        path = self.variant(mutate, styles=True)
        self.assertEqual(item(check_document(path, [5]), 'R05', paragraph=5)['status'], 'review')

    def test_missing_page_properties_is_review(self):
        def mutate(pkg):
            pkg.sections[0].remove(pkg.sections[0].find(q('pgSz')))
        path = self.variant(mutate)
        f = item(check_document(path, BODY), 'R01', section=1)
        self.assertEqual(f['status'], 'review')
        self.assertFalse(f['fixable'])

    def test_ambiguous_nested_section_reported_and_fix_blocked(self):
        def mutate(pkg):
            change = _ensure(pkg.sections[0], 'sectPrChange')
            change.append(deepcopy(pkg.sections[1]))
        path = self.variant(mutate)
        self.assertEqual(item(check_document(path, BODY), 'R01', document=1)['status'], 'review')
        with self.assertRaises(DemoError):
            fix_document(path, self.dir / 'fixed.docx', BODY, IDS)

    def test_original_existing_output_and_symlink_blocked(self):
        original = digest(ERRORS)
        with self.assertRaises(DemoError):
            fix_document(ERRORS, ERRORS, BODY, IDS)
        existing = self.dir / 'existing.docx'
        existing.write_bytes(b'keep existing')
        with self.assertRaises(DemoError):
            fix_document(ERRORS, existing, BODY, IDS)
        self.assertEqual(existing.read_bytes(), b'keep existing')
        alias = self.dir / 'alias.docx'
        alias.symlink_to(ERRORS)
        with self.assertRaises(DemoError):
            fix_document(ERRORS, alias, BODY, IDS)
        self.assertEqual(digest(ERRORS), original)

    def test_manifest_mismatch_duplicate_and_out_of_range_rejected(self):
        with self.assertRaises(DemoError):
            check_document(CORRECT, ROOT / 'samples/errors.body.json')
        for body in ([4, 4], [0], [999], ['4'], None):
            with self.subTest(body=body), self.assertRaises(DemoError):
                check_document(CORRECT, body)

    def test_bad_docx_non_docx_and_input_limit_rejected(self):
        for name in ('broken.docx', 'wrong.txt'):
            p = self.dir / name
            p.write_bytes(b'not a docx')
            with self.assertRaises(DemoError):
                check_document(p, BODY)
        p = self.dir / 'large.docx'
        with p.open('wb') as f:
            f.truncate(10 * 1024 * 1024 + 1)
        with self.assertRaises(DemoError):
            check_document(p, BODY)

    def test_signed_file_fix_rejected(self):
        path = self.variant(lambda pkg: None, extra={'_xmlsignatures/sig1.xml': b'<signature/>'})
        with self.assertRaises(DemoError):
            fix_document(path, self.dir / 'fixed.docx', BODY, IDS)
        self.assertFalse((self.dir / 'fixed.docx').exists())

    def test_cli_two_full_fresh_process_runs_and_exit_codes(self):
        for n in (1, 2):
            directory = self.dir / str(n)
            args = [sys.executable, str(ROOT / 'cli.py')]
            checked = subprocess.run(args + ['check', str(ERRORS), '--body-manifest', str(ROOT / 'samples/errors.body.json'),
                                     '--report', str(directory / 'check.json')], capture_output=True, text=True)
            self.assertEqual(checked.returncode, 1, checked.stderr)
            fixed = subprocess.run(args + ['fix', str(ERRORS), '--body-ranges', '4-6,11,14-15', '--select', ','.join(IDS),
                                   '--output', str(directory / 'fixed.docx'), '--report', str(directory / 'changes.json'),
                                   '--recheck-report', str(directory / 'recheck.json')], capture_output=True, text=True)
            self.assertEqual(fixed.returncode, 0, fixed.stderr)
            self.assertEqual(json.loads((directory / 'recheck.json').read_text())['summary']['problem_count'], 0)
            # Repeating the same output paths must fail before any writes.
            again = subprocess.run(args + ['fix', str(ERRORS), '--body-ranges', '4-6,11,14-15', '--select', 'R05',
                                    '--output', str(directory / 'fixed.docx'), '--report', str(directory / 'new.json'),
                                    '--recheck-report', str(directory / 'new-recheck.json')], capture_output=True, text=True)
            self.assertEqual(again.returncode, 3)
            self.assertFalse((directory / 'new.json').exists())

    def test_cli_reports_cannot_overwrite_original_or_collide(self):
        before = digest(ERRORS)
        cmd = [sys.executable, str(ROOT / 'cli.py'), 'fix', str(ERRORS), '--body-ranges', '4', '--select', 'R05',
               '--output', str(self.dir / 'fixed.docx'), '--report', str(self.dir / 'same.json'),
               '--recheck-report', str(self.dir / 'same.json')]
        result = subprocess.run(cmd, capture_output=True, text=True)
        self.assertEqual(result.returncode, 3)
        self.assertFalse((self.dir / 'fixed.docx').exists())
        self.assertEqual(digest(ERRORS), before)

if __name__ == '__main__':
    unittest.main(verbosity=2)
