"""Actual transactions plus independent XML/ZIP oracles; entirely portable/offline."""
from pathlib import Path
from copy import deepcopy
from unittest import TestCase,mock
import tempfile,json,zipfile
from io import BytesIO
from lxml import etree as E
from core.package import Package,q,sha256,DemoError,NS
from core.rehearsal import reconcile
from scripts.reproduce_scope_demo import session,old_check,ROOT
from web.session import WebSession
from core import fix_document

KEY='R05:paragraph=8:run=1'
class RehearsalTests(TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.d=Path(self.tmp.name);self.s=session(self.d/'a');old_check(self.s);self.s.set_selection([KEY])
 def preview(self,s=None):
  s=s or self.s;self.assertTrue(s.rehearse(),s.rehearsal_error);return s.rehearsal['id']
 def adopt(self,s=None):
  s=s or self.s;self.assertTrue(s.adopt_rehearsal(s.rehearsal['id'],acknowledged=True),s.rehearsal_error)
 def test_actual_rehearsal_does_not_publish_or_switch(self):
  before=self.s.input_path;raw=before.read_bytes();self.preview();self.assertEqual(self.s.current_round,0);self.assertEqual(self.s.generation_count,0);self.assertEqual(self.s.history,[]);self.assertIsNone(self.s.result);self.assertEqual(self.s.input_path,before);self.assertEqual(before.read_bytes(),raw)
 def test_natural_R07_and_distinct_count_meanings(self):
  self.preview();v=self.s.rehearsal_view();i=v['impact'];self.assertEqual(i['change_record_count'],1);self.assertEqual(i['xml_attribute_change_count'],1);self.assertEqual(i['counts']['selected_resolved'],1);self.assertEqual(i['counts']['natural_resolved'],1);self.assertEqual(i['counts']['remaining_unselected'],49);self.assertIn('未修改缩进',next(x for x in i['records'] if x['kind']=='natural_resolved')['reason'])
 def test_independent_whitelist_exact_xml_and_zip(self):
  self.preview();a=Package(self.s.input_path);b=Package(self.s.rehearsal['result']['output']);old=a.paragraphs[7].findall(q('r'))[0].find('w:rPr/w:sz',NS);new=b.paragraphs[7].findall(q('r'))[0].find('w:rPr/w:sz',NS)
  self.assertEqual((old.get(q('val')),new.get(q('val'))),('26','24'));old.set(q('val'),'24');self.assertEqual(E.tostring(a.root,method='c14n'),E.tostring(b.root,method='c14n'))
  for n,v in a.parts.items():
   if n!='word/document.xml':self.assertEqual(v,b.parts[n])
 def test_same_bytes_as_direct_legal_path_and_existing_visual_fixture(self):
  self.preview();r=fix_document(self.s.input_path,self.d/'direct.docx',self.s.confirmed,selected_keys=[KEY],**self.s._rules_arguments());preview=Path(self.s.rehearsal['result']['output']).read_bytes();self.assertEqual(preview,Path(r['output']).read_bytes());self.assertEqual(preview,(ROOT/'samples/phase26/fixed.docx').read_bytes());self.adopt();self.assertEqual(preview,self.s.downloads()['fixed.docx'])
 def test_adoption_does_not_fix_again_and_duplicate_only_one_round(self):
  i=self.preview();p=Path(self.s.rehearsal['directory'])
  with mock.patch('web.session.fix_document',side_effect=AssertionError('second fix')):
   self.adopt();self.assertTrue(self.s.adopt_rehearsal(i,acknowledged=True))
  self.assertEqual(self.s.generation_count,1);self.assertEqual(len(self.s.history),1);self.assertFalse(p.exists());self.assertEqual(self.s.current_round,0);self.assertEqual(self.s.result['rehearsal']['fix_calls_at_adoption'],0)
 def test_repeated_preview_and_download_no_fix_or_extra_files(self):
  self.preview();files=set(self.s.task_dir.rglob('*'))
  with mock.patch('web.session.fix_document',side_effect=AssertionError()):
   self.assertTrue(self.s.rehearse());self.s.rehearsal_download();self.s.rehearsal_view()
  self.assertEqual(files,set(self.s.task_dir.rglob('*')))
 def test_ack_required_keeps_pending(self):
  i=self.preview();self.assertFalse(self.s.adopt_rehearsal(i));self.assertIsNotNone(self.s.rehearsal_view());self.assertEqual(len(self.s.history),0)
 def test_cancel_cleans_only_owned_files_and_keeps_input(self):
  other=self.s.task_dir/'foreign.txt';other.write_text('keep');i=self.preview();p=Path(self.s.rehearsal['directory']);self.s.cancel_rehearsal();self.assertFalse(p.exists());self.assertEqual(other.read_text(),'keep');self.assertFalse(self.s.adopt_rehearsal(i,acknowledged=True));self.assertIsNone(self.s.result)
 def test_selection_change_and_return_to_original_still_requires_new_preview(self):
  i=self.preview();self.s.set_selection([]);self.s.set_selection([KEY]);self.assertIsNone(self.s.rehearsal);self.assertFalse(self.s.adopt_rehearsal(i,acknowledged=True))
 def test_scope_draft_and_confirmation_changes_rejected(self):
  i=self.preview();self.s.note_draft['objects']=[];self.assertFalse(self.s.adopt_rehearsal(i,acknowledged=True));self.assertEqual(len(self.s.history),0)
  self.s=session(self.d/'b');old_check(self.s);self.s.set_selection([KEY]);i=self.preview();self.s.classification_confirmed['confirmation']='changed';self.assertFalse(self.s.adopt_rehearsal(i,acknowledged=True))
 def test_rules_and_dependency_bytes_change_rejected(self):
  rp=self.d/'rules.json';rp.write_bytes(self.s.rules_path.read_bytes());self.s.rules_path=rp;i=self.preview();rp.write_bytes(rp.read_bytes()+b' ');self.assertFalse(self.s.adopt_rehearsal(i,acknowledged=True));self.assertFalse(self.s.history)
 def test_execution_version_change_rejected(self):
  i=self.preview()
  with mock.patch('web.rehearsal_session.execution_hash',return_value='new-version'):self.assertFalse(self.s.adopt_rehearsal(i,acknowledged=True))
 def test_external_input_change_rejected(self):
  i=self.preview();self.s.input_path.write_bytes((ROOT/'samples/correct.docx').read_bytes());self.assertFalse(self.s.adopt_rehearsal(i,acknowledged=True));self.assertFalse(self.s.history)
 def test_output_missing_or_tampered_rejected(self):
  for name in ('remove','change'):
   s=session(self.d/name);old_check(s);s.set_selection([KEY]);i=self.preview(s);p=Path(s.rehearsal['result']['output']);p.unlink() if name=='remove' else p.write_bytes(b'changed');self.assertFalse(s.adopt_rehearsal(i,acknowledged=True));self.assertFalse(s.history)
 def test_transaction_report_and_memory_impact_tamper_rejected(self):
  i=self.preview();(Path(self.s.rehearsal['directory'])/'transaction.json').write_text('{}');self.assertFalse(self.s.adopt_rehearsal(i,acknowledged=True));self.preview();i=self.s.rehearsal['id'];self.s.rehearsal['view']['impact']['counts']['natural_resolved']=999;self.assertFalse(self.s.adopt_rehearsal(i,acknowledged=True))
 def test_archive_storage_failure_no_half_round_or_orphan(self):
  i=self.preview();before=set(self.s.task_dir.iterdir())
  import web.session as api
  write=api.write_new;calls=[]
  def fail(*a,**kw):
   calls.append(a[0]);
   if len(calls)==3:raise OSError('disk full')
   return write(*a,**kw)
  with mock.patch('web.session.write_new',side_effect=fail):self.assertFalse(self.s.adopt_rehearsal(i,acknowledged=True))
  self.assertEqual(self.s.generation_count,0);self.assertFalse(self.s.history);self.assertIsNone(self.s.result);self.assertEqual(before,set(self.s.task_dir.iterdir()));self.adopt()
 def test_final_archive_binding_race_refuses_publication(self):
  i=self.preview();original=self.s._archive_round
  def race(staged,result):
   original_fn=WebSession._archive_round;original_fn(staged,result);self.s.selected=[]
  with mock.patch.object(WebSession,'_archive_round',race):self.assertFalse(self.s.adopt_rehearsal(i,acknowledged=True))
  self.assertEqual(self.s.generation_count,0);self.assertFalse(self.s.history);self.assertFalse(list(self.s.task_dir.glob('fixed-*')))
 def test_project_restore_keeps_selection_but_never_preview_authorization(self):
  self.preview();raw=self.s.prepare_project();z=zipfile.ZipFile(BytesIO(raw));self.assertFalse(any('rehearsal-' in n for n in z.namelist()));t=WebSession(self.d/'restore');t.stage_project(raw);r=t.restore_project(acknowledged=True);self.assertIsNone(r.rehearsal);self.assertTrue(r.report);self.assertEqual(r.selected,[KEY]);self.assertFalse(r.adopt_rehearsal(self.s.rehearsal['id'],acknowledged=True));self.assertTrue(r.rehearse(),r.rehearsal_error)
 def test_formal_adopt_archive_project_and_next_round(self):
  self.preview();self.adopt();raw=self.s.prepare_project();t=WebSession(self.d/'restore');t.stage_project(raw);r=t.restore_project(acknowledged=True);self.assertEqual(len(r.history),1);self.assertTrue(r.continue_history(r.history[0]['stem'],acknowledged=True),r.error);self.assertEqual(r.current_round,1);self.assertIsNone(r.rehearsal)
 def test_two_sessions_do_not_accept_each_others_preview(self):
  i=self.preview();t=session(self.d/'b');old_check(t);t.set_selection([KEY]);self.preview(t);self.assertFalse(t.adopt_rehearsal(i,acknowledged=True));self.assertTrue(self.s.rehearsal_view());self.assertEqual(t.generation_count,0)
 def test_direct_compatibility_path_invalidates_pending_preview(self):
  i=self.preview();self.assertTrue(self.s.generate(),self.s.error);self.assertIsNone(self.s.rehearsal);self.assertEqual(self.s.generation_count,1);self.assertFalse(self.s.adopt_rehearsal(i,acknowledged=True))
 def test_report_html_md_json_zip_keep_impact_and_escape(self):
  self.preview();self.adopt();files=self.s.prepare_reading_export();d=json.loads(files['report.json']);self.assertEqual(d['impact'],self.s.result['impact']);self.assertIn('预演与正式采用',files['report.html'].decode());self.assertIn('预演与正式采用',files['report.md'].decode());self.assertNotIn(str(self.s.task_dir),files['report.html'].decode());z=zipfile.ZipFile(BytesIO(self.s.result_package()));self.assertEqual(json.loads(z.read('rehearsal-impact.json'))['impact'],d['impact']);self.assertIn('预演及采用影响',z.read('changes.md').decode())
 def test_problem_to_review_does_not_count_as_resolved(self):
  self.preview();r=deepcopy(self.s.rehearsal['result']);f=next(f for f in r['after']['findings'] if f['rule_id']=='R07');f['status']='review';out=reconcile(r,Package(self.s.input_path),Package(r['output']));self.assertEqual(out['counts']['natural_resolved'],0);self.assertEqual(out['counts']['problem_to_review'],1)
 def test_key_change_same_stable_object_is_not_treated_as_disappearance(self):
  self.preview();r=deepcopy(self.s.rehearsal['result']);next(f for f in r['after']['findings'] if f['rule_id']=='R07')['key']='changed-key';i=reconcile(r,Package(self.s.input_path),Package(r['output']));self.assertEqual(i['counts']['natural_resolved'],1)
 def test_R05_creates_unselected_R07_without_changing_indent(self):
  from docx import Document
  from docx.shared import Pt
  # Build by independently changing the demo body firstLine to26pt: 13pt×2.
  p=Package(ROOT/'samples/phase26/composite.docx');p.paragraphs[7].find('w:pPr/w:ind',NS).set(q('firstLine'),'520');src=self.d/'new.docx'
  with zipfile.ZipFile(src,'w') as z:
   for n,v in p.parts.items():z.writestr(n,E.tostring(p.root,encoding='utf-8',xml_declaration=True) if n=='word/document.xml' else v)
  s=WebSession(self.d/'new-web');s.load(src.read_bytes(),'合成缩进案例.docx');s.set_body([8]);s.confirm();self.assertTrue(s.check());s.set_selection([KEY]);self.preview(s);r=s.rehearsal['result'];imp=s.rehearsal_view()['impact'];self.assertEqual(imp['counts']['new_problem'],1);self.assertEqual(imp['counts']['natural_resolved'],0);self.assertIn('26pt',next(x['reason'] for x in imp['records'] if x['kind']=='new_problem'));self.assertEqual(Package(r['output']).paragraphs[7].find('w:pPr/w:ind',NS).get(q('firstLine')),'520')
 def test_R33_existing_split_protection_and_coverage_change(self):
  from tests.test_citations import CitationTests,manifest
  p=ROOT/'samples/phase19/errors.docx';m=manifest(p);s=WebSession(self.d/'c');s.set_mode('academic');s.load(p.read_bytes(),'引文演示.docx');s.set_types(m['paragraph_types']);s.enable_citations(True);s.edit_citations(m['citation_scope']);s.confirm_citations();self.assertTrue(s.confirm_and_check());keys=[f['key'] for f in s.eligible() if f['rule_id']=='R33'];s.set_selection(keys[1:2]);self.preview(s);self.assertIsNone(s.rehearsal_view()['impact']['xml_attribute_change_count']);self.adopt(s);self.assertTrue(s.result['rehearsal']['validation']['white_list_revalidated'])
 def test_R33_ordinary_run_combination_still_rejected(self):
  from tests.test_citations import manifest,rewrite
  p=ROOT/'samples/phase19/errors.docx';m=manifest(p)
  def bad(pkg):
   from core.engine import _ensure
   _ensure(_ensure(pkg.paragraphs[2].findall(q('r'))[0],'rPr'),'sz').set(q('val'),'30')
  src=rewrite(p,self.d/'bad.docx',bad);m=manifest(src);s=WebSession(self.d/'c');s.set_mode('academic');s.load(src.read_bytes(),'引文演示.docx');s.set_types(m['paragraph_types']);s.enable_citations(True);s.edit_citations(m['citation_scope']);s.confirm_citations();self.assertTrue(s.confirm_and_check());ks=[f['key'] for f in s.eligible() if f['locator'].get('paragraph')==3 and f['rule_id'] in ('R05','R33')];self.assertTrue(any(k.startswith('R05') for k in ks));s.set_selection(ks);self.assertFalse(s.rehearse());self.assertIn('分轮',s.rehearsal_error);self.assertFalse(s.history);self.assertFalse(list(s.task_dir.glob('rehearsal-*')))
 def test_R11_preserves_unselected_effective_format_and_styles(self):
  from tests.test_extended import ERROR,MANIFEST
  from core.extended import formatting_snapshot,snapshots_equal
  s=WebSession(self.d/'h');s.set_mode('extended');s.load(ERROR.read_bytes(),'标题演示.docx');s.set_types(MANIFEST['paragraph_types']);self.assertTrue(s.confirm_and_check());s.set_selection([f['key'] for f in s.eligible() if f['rule_id']=='R11']);self.preview(s);self.adopt(s);a,b=Package(s.input_path),Package(s.result['output']);self.assertEqual(a.parts['word/styles.xml'],b.parts['word/styles.xml']);self.assertTrue(snapshots_equal(formatting_snapshot(a,a.paragraphs[12]),formatting_snapshot(b,b.paragraphs[12])))
 def test_notes_story_and_table_paths_use_same_guard_and_keep_unselected_parts(self):
  for rid in ('R23','R24','R26'):
   with self.subTest(rule=rid):
    s=session(self.d/rid);old_check(s);items=[f for f in s.eligible() if f['rule_id']==rid]
    self.assertTrue(items,rid)
    s.set_selection([items[0]['key']]);self.preview(s);self.adopt(s);self.assertTrue(s.result['rehearsal']['validation']['white_list_revalidated']);self.assertEqual(len(s.result['changes']),1)
    a,b=Package(s.input_path),Package(s.result['output'])
    if items[0].get('story_part'):self.assertEqual(a.parts['word/document.xml'],b.parts['word/document.xml'])
 def test_mode_change_cancels_and_isolation_of_cleanup(self):
  i=self.preview();p=Path(self.s.rehearsal['directory']);t=session(self.d/'t');old_check(t);t.set_selection([KEY]);self.preview(t);other=Path(t.rehearsal['result']['output']);self.s.set_mode('basic');self.assertFalse(p.exists());self.assertTrue(other.exists());self.assertFalse(self.s.adopt_rehearsal(i,acknowledged=True));self.assertTrue(t.rehearsal_view())
 def test_recheck_invalidates_prior_preview(self):
  i=self.preview();self.assertTrue(self.s.check());self.assertIsNone(self.s.rehearsal);self.assertFalse(self.s.adopt_rehearsal(i,acknowledged=True))
 def test_caller_cannot_mutate_view_or_download_to_authorize_new_values(self):
  self.preview();a=self.s.rehearsal_view();a['impact']['counts']['new_problem']=999;self.assertNotEqual(a,self.s.rehearsal_view());self.assertEqual(self.s.generation_count,0)
 def test_storage_failure_during_preview_cleans_owned_directory(self):
  with mock.patch('web.rehearsal_session.write_new',side_effect=OSError('full')):self.assertFalse(self.s.rehearse())
  self.assertIsNone(self.s.rehearsal);self.assertFalse(list(self.s.task_dir.glob('rehearsal-*')));self.assertFalse(self.s.history)
 def test_html_escape_of_impact_uses_original_report_protection(self):
  self.preview();self.adopt();from web.reading_report import portable,html_report
  b=self.s.current_bundle();b['impact']['records'][0]['reason']='<script>bad()</script>';html=html_report(portable(b));self.assertNotIn('<script>bad()',html);self.assertIn('&lt;script&gt;',html)
 def test_actual_streamlit_preview_and_ack_adoption(self):
  from streamlit.testing.v1 import AppTest
  app=AppTest.from_file(str(ROOT/'app.py'),default_timeout=30);app.session_state['workflow']=self.s;app.run();self.assertFalse(app.exception)
  next(b for b in app.button if b.label=='预演所选修正').click();app.run();self.assertFalse(app.exception);self.assertEqual(self.s.generation_count,0);self.assertTrue(self.s.rehearsal)
  next(c for c in app.checkbox if c.label=='我已核对全部所选修改、关联影响及剩余事项，确认采用这一副本').check();app.run();next(b for b in app.button if b.label=='确认采用此副本').click();app.run();self.assertFalse(app.exception);self.assertEqual(self.s.generation_count,1);self.assertEqual(len(self.s.history),1)
 def test_selected_note_story_only_transaction_is_revalidated(self):
  from tests.test_notes import manifest,S
  src=S/'errors.docx';m=manifest(src);s=WebSession(self.d/'n');s.set_mode('academic');s.load(src.read_bytes(),'注释合成.docx');s.set_types(m['paragraph_types']);s.enable_notes(True);s.edit_notes(m['note_scope']);s.confirm_notes();self.assertTrue(s.confirm_and_check())
  for rid in ('R36','R37'):
   keys=[f['key'] for f in s.eligible() if f['rule_id']==rid];self.assertTrue(keys);s.set_selection(keys[:1]);self.preview(s);self.adopt(s);a,b=Package(s.input_path),Package(s.result['output']);self.assertEqual(a.parts['word/document.xml'],b.parts['word/document.xml']);self.assertTrue(s.next_round())
 def test_ordinary_header_story_reuse_and_shared_unsafe_remain_ineligible(self):
  from tests.test_front import M,BAD
  m=deepcopy(M);m['input_sha256']=m['front_scope']['input_sha256']=sha256(BAD.read_bytes());s=WebSession(self.d/'f');s.set_mode('academic');s.load(BAD.read_bytes(),'分节合成.docx');s.set_types(m['paragraph_types']);s.enable_front(True);s.edit_front(m['front_scope']);s.confirm_front();self.assertTrue(s.confirm_and_check());fs=[f for f in s.eligible() if f['rule_id']=='R27'];self.assertTrue(fs);s.set_selection([fs[0]['key']]);self.preview(s);self.adopt(s);a,b=Package(s.input_path),Package(s.result['output']);self.assertEqual(a.parts['word/document.xml'],b.parts['word/document.xml']);self.assertEqual(a.parts['word/styles.xml'],b.parts['word/styles.xml'])
 def test_final_binding_validation_failure_before_publication(self):
  i=self.preview();binding=self.s.rehearsal_binding
  def fail(**kwargs):
   if kwargs:raise OSError('final read failed')
   return binding()
  with mock.patch.object(self.s,'rehearsal_binding',side_effect=fail):self.assertFalse(self.s.adopt_rehearsal(i,acknowledged=True))
  self.assertFalse(self.s.history);self.assertEqual(self.s.generation_count,0);self.assertIsNone(self.s.result);self.assertFalse(list(self.s.task_dir.glob('fixed-*')));self.assertTrue(self.s.rehearsal_view())
 def test_remaining_reviews_and_coverage_preserved_in_full_impact(self):
  self.preview();i=self.s.rehearsal_view()['impact'];self.assertEqual(i['counts']['remaining_review'],self.s.report['summary']['review_findings']);self.assertEqual(i['coverage_notes'],self.s.rehearsal['result']['after']['coverage_notes']);self.assertEqual(i['coverage_changes']['after_count'],len(i['coverage_notes']))
 def test_same_table_run_different_properties_align_independently(self):
  self.s.set_selection([f['key'] for f in self.s.eligible() if f['rule_id']=='R23'][:1]);self.preview();im=self.s.rehearsal_view()['impact'];self.assertEqual(im['counts']['selected_resolved'],1);self.assertEqual(im['counts']['unaligned'],0)
 def test_cleanup_failure_does_not_delete_adopted_round(self):
  i=self.preview()
  with mock.patch('web.rehearsal_session.shutil.rmtree',side_effect=OSError('busy')):self.assertTrue(self.s.adopt_rehearsal(i,acknowledged=True))
  self.assertEqual(len(self.s.history),1);self.assertEqual(self.s.generation_count,1);self.assertTrue(self.s.downloads());self.assertTrue(self.s.adopt_rehearsal(i,acknowledged=True))
 def test_formal_history_and_reading_summary_use_stable_impact_not_legacy_delta(self):
  self.preview();self.adopt();self.s.result['problem_delta']['resolved_keys']=['fake']*50;self.s.history[0]['result']['problem_delta']['resolved_keys']=['fake']*50
  self.assertEqual(self.s.history_rows()[0]['counts']['格式已解决'],2)
  from web.reading_report import portable,html_report
  html=html_report(portable(self.s.current_bundle()));self.assertNotIn('已解决 50',html);self.assertIn('未选问题自然符合',html)
 def test_preview_adopt_download_does_not_call_AI_renderer_or_cancel_background_jobs(self):
  with mock.patch('urllib.request.urlopen',side_effect=AssertionError('network')) as network,mock.patch('core.preview.render_preview',side_effect=AssertionError('render')) as render,mock.patch.object(self.s,'invalidate_preview_jobs',side_effect=AssertionError('cancel worker')):
   self.preview();self.adopt();self.s.downloads();self.s.result_package()
  self.assertEqual(network.call_count,0);self.assertEqual(render.call_count,0)
 def test_cancel_cleanup_failure_revokes_authorization_without_crashing_workflow(self):
  i=self.preview()
  with mock.patch('web.rehearsal_session.shutil.rmtree',side_effect=OSError('busy')):self.s.cancel_rehearsal()
  self.assertIsNone(self.s.rehearsal);self.assertIn('已取消',self.s.rehearsal_error);self.assertFalse(self.s.adopt_rehearsal(i,acknowledged=True));self.assertEqual(self.s.generation_count,0);self.assertTrue(self.s.report)
