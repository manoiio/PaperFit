from pathlib import Path
from tempfile import TemporaryDirectory
from copy import deepcopy
from unittest.mock import patch
from lxml import etree as E
import unittest,json
from core.package import Package,NS,q,DemoError
from core import word_fields as wf
ROOT=Path(__file__).resolve().parents[1]
class WordFieldsTests(unittest.TestCase):
 def setUp(self):self.tmp=TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.root=Path(self.tmp.name);self.source=ROOT/'samples/phase50/errors.docx';pkg=Package(self.source);self.m={'input_sha256':pkg.hash,'paragraph_types':json.loads((ROOT/'samples/phase50/roles.json').read_text()),'confirmed':True,'confirmation':'独立合成角色'}
 def test_inventory_explicit_safe_and_kept(self):
  rows=wf.inventory(self.source)['objects'];self.assertEqual([r['kind'] for r in rows if r['safe']],['TOC','REF','PAGEREF']);self.assertTrue(any('锁定' in (r['reason'] or '') for r in rows));self.assertTrue(any('未找到' in (r['reason'] or '') for r in rows))
 def mock(self,raw,codes,targets,update=True):
  path=self.root/'current.docx';path.write_bytes(raw);p=Package(path)
  pages={n:{'adjusted':7 if n=='ChapterTwo' or n.endswith('_11') else 5,'physical':7 if n=='ChapterTwo' or n.endswith('_11') else 5} for n in targets}
  if update:
   for f in reversed(wf.fields(p)[0]):
    if f['instruction'] in codes and f['instruction'].strip().startswith('REF '):wf.set_cache(p,f,'图1-2 合成曲线')
    if f['instruction'] in codes and f['instruction'].strip().startswith('PAGEREF '):wf.set_cache(p,f,'7')
   for f in p.root.iter(q('fldSimple')):
    if f.get(q('instr'),'').strip().startswith('PAGEREF PFToc_'):
     n=f.get(q('instr')).split()[1];f.find('w:r/w:t',NS).text=str(pages[n]['adjusted'])
  return {'raw':p.bytes(True),'pages':pages,'seconds':0,'backend':'mock'}
 def test_full_backfill_exact_parts_and_preserved_targets(self):
  r=wf.update(self.source,self.root/'fixed.docx',self.m,wf.plan(self.source,self.m),calculator=self.mock,artifact_dir=self.root/'art');b=Package(r['output']);a=Package(self.source)
  self.assertEqual(r['proof']['passes'][-1]['pass'],2);self.assertTrue(r['proof']['passes'][-1]['final_target_pages_equal']);self.assertFalse(r['proof']['calculator_package_published'])
  for n in a.parts:
   if n!='word/document.xml':self.assertEqual(a.parts[n],b.parts[n])
  fields=wf.fields(b)[0];self.assertEqual([wf.cache_text(b,f) for f in fields if f['instruction'].strip().startswith('REF Missing')],['保持缺失目标缓存']);self.assertEqual([wf.cache_text(b,f) for f in fields if f['instruction'].strip().startswith('SEQ ')],['66'])
 def test_binding_and_unsupported_rejected_before_word(self):
  t=wf.plan(self.source,self.m);t['input_sha256']='0'*64
  with self.assertRaises(DemoError):wf.update(self.source,self.root/'bad.docx',self.m,t,calculator=lambda *a,**k:self.fail('must not call'))
  with self.assertRaises(DemoError):wf.plan(self.source,self.m,['field:5'])
 def test_read_only_does_not_change_file(self):
  before=self.source.read_bytes();wf.inventory(self.source);wf.plan(self.source,self.m);self.assertEqual(before,self.source.read_bytes())
 def test_failure_does_not_create_final(self):
  with self.assertRaises(DemoError):wf.update(self.source,self.root/'bad.docx',self.m,wf.plan(self.source,self.m),calculator=lambda *a,**k:(_ for _ in ()).throw(DemoError('Word unavailable')))
  self.assertFalse((self.root/'bad.docx').exists())
 def test_creation_requires_confirmed_ordinary_region(self):
  with self.assertRaises(DemoError):wf.plan(self.source,self.m,[],{'start':4,'end':5,'levels':[1,2,3],'confirmed':True})
 def test_native_never_called_by_regular_check(self):
  from core import check_document
  with patch('core.word_native.calculate',side_effect=AssertionError('no Word in ordinary check')):check_document(self.source,mode='academic',classification_manifest=self.m)
 def test_cancellation_ownership_and_queued_worker(self):
  from core.preview_jobs import PreviewJobs
  from threading import Event
  from core.package import sha256
  gate=Event();j=PreviewJobs();root=self.root/'queue'
  def worker(source,out,context):gate.wait(2);return {'document_sha256':sha256(source.read_bytes()),'render_config_sha256':'seal'}
  identity=j.submit('one',root,self.source.read_bytes(),('a','seal'),{},worker)
  with self.assertRaises(DemoError):j.status('two',identity)
  j.cancel('one',identity);self.assertEqual(j.status('one',identity)['state'],'cancelled');gate.set()
 def test_same_code_partial_selection_is_rejected(self):
  p=Package(self.source);f=p.root.xpath('.//w:fldSimple',namespaces=NS)[1];f.addnext(deepcopy(f));path=self.root/'dupe.docx';path.write_bytes(p.bytes(True));m=deepcopy(self.m);m['input_sha256']=Package(path).hash
  with self.assertRaises(DemoError):wf.plan(path,m,['field:2'])
 def test_same_paragraph_same_instruction_ref_keeps_each_occurrence(self):
  p=Package(self.source);f=next(f for f in wf.fields(p)[0] if f['simple'] and f['instruction'].strip().startswith('REF CaptionOne'))
  f['node'].addnext(deepcopy(f['node']));path=self.root/'same-paragraph.docx';path.write_bytes(p.bytes(True));m=deepcopy(self.m);m['input_sha256']=Package(path).hash
  r=wf.update(path,self.root/'both.docx',m,wf.plan(path,m),calculator=self.mock,artifact_dir=self.root/'both-art')
  out=Package(r['output']);values=[wf.cache_text(out,x) for x in wf.fields(out)[0] if x['instruction'].strip().startswith('REF CaptionOne')]
  self.assertEqual(values,['图1-2 合成曲线','图1-2 合成曲线'])
 def test_locked_nested_toc_and_unsupported_switches_remain_protected(self):
  p=Package(self.source);f=p.root.xpath('.//w:fldSimple',namespaces=NS)[0];f.set(q('instr'),' TOC \\t "Custom,1" ');path=self.root/'custom.docx';path.write_bytes(p.bytes(True));self.assertFalse(wf.inventory(path)['objects'][0]['safe'])
 def test_no_stable_page_signature_never_publishes(self):
  calls=[]
  def moving(raw,codes,targets,update=True):
   r=self.mock(raw,codes,targets,update);calls.append(update)
   if not update:
    for v in r['pages'].values():v['physical']+=1
   return r
  with self.assertRaisesRegex(DemoError,'未稳定'):wf.update(self.source,self.root/'unstable.docx',self.m,wf.plan(self.source,self.m),calculator=moving,artifact_dir=self.root/'unstable')
  self.assertEqual(len(calls),8);self.assertFalse((self.root/'unstable.docx').exists())
 def test_creation_works_only_with_hand_confirmed_toc_and_levels(self):
  from core.toc import make_scope
  p=Package(self.source);f=p.root.xpath('.//w:fldSimple',namespaces=NS)[0];f.getparent().remove(f);path=self.root/'manual.docx';path.write_bytes(p.bytes(True));m=deepcopy(self.m);m['input_sha256']=Package(path).hash;scope=make_scope(path);scope.update(confirmed=True,objects={'d:4':'title','d:5':'level1'});m['toc_scope']=scope
  r=wf.update(path,self.root/'created.docx',m,wf.plan(path,m,[],{'start':5,'end':5,'levels':[1,2,3],'confirmed':True}),calculator=self.mock,artifact_dir=self.root/'creation');self.assertTrue(r['changes']);self.assertEqual(wf.inventory(r['output'])['objects'][0]['kind'],'TOC')
 def test_async_adoption_restore_and_tampered_preview_are_bound(self):
  import time
  from web.session import WebSession
  from web import field_session as session
  from core.preview_jobs import PreviewJobs
  s=WebSession(self.root/'web');s.set_mode('academic');s.load(self.source.read_bytes(),'公开合成.docx');s.set_types(self.m['paragraph_types']);s.confirm_classification()
  real=wf.update
  def controlled(*a,**kw):return real(*a,**kw,calculator=self.mock)
  state={}
  with patch('core.word_fields.update',side_effect=controlled):
   session.submit(s,state,scheduler=PreviewJobs())
   for _ in range(200):
    status=session.receive(s,state)
    if status['state'] not in ('queued','running'):break
    time.sleep(.02)
   self.assertEqual(status['state'],'succeeded',status);self.assertEqual(s.generation_count,0)
   with patch('web.session.WebSession._archive_round',side_effect=OSError('storage failure')):
    with self.assertRaises(OSError):session.adopt(s,state)
   self.assertEqual(s.generation_count,0);self.assertEqual(len(s.history),0)
   with patch('core.word_native.calculate',side_effect=AssertionError('adoption must not recalculate')):session.adopt(s,state);session.adopt(s,state)
  self.assertEqual(s.generation_count,1);self.assertEqual(len(s.history),1);self.assertEqual(s.current_round,0)
  raw=s.prepare_project();other=WebSession(self.root/'restore');other.stage_project(raw)
  with patch('core.word_native.calculate',side_effect=AssertionError('restore must not update Word')):other=other.restore_project(acknowledged=True)
  self.assertEqual(other.history[-1]['result']['kind'],'field_update');self.assertTrue(other.history[-1]['result']['restored_history'])
  Path(state['result']['output']).write_bytes(b'changed')
  with self.assertRaises(DemoError):s.downloads()
 def test_outline_filter_reads_existing_resolver_inheritance(self):
  from docx import Document
  from docx.enum.style import WD_STYLE_TYPE
  from docx.oxml import OxmlElement
  from docx.oxml.ns import qn
  d=Document();style=d.styles.add_style('CustomOutline',WD_STYLE_TYPE.PARAGRAPH);ppr=style.element.get_or_add_pPr();x=OxmlElement('w:outlineLvl');x.set(qn('w:val'),'1');ppr.append(x);d.add_paragraph('自定义层级标题','CustomOutline');path=self.root/'outline.docx';d.save(path);pkg=Package(path)
  self.assertIsNone(wf.paragraph_level(pkg,pkg.paragraphs[0],False));self.assertEqual(wf.paragraph_level(pkg,pkg.paragraphs[0],True),2)
 def test_hidden_target_keeps_reference_but_leaves_independent_toc(self):
  p=Package(self.source);bm=p.root.xpath('.//w:bookmarkStart[@w:name="CaptionOne"]',namespaces=NS)[0];r=bm.getparent().find(q('r'));pr=r.find(q('rPr'))
  if pr is None:pr=E.Element(q('rPr'));r.insert(0,pr)
  E.SubElement(pr,q('vanish'));path=self.root/'hidden.docx';path.write_bytes(p.bytes(True));rows=wf.inventory(path)['objects'];self.assertTrue(rows[0]['safe']);self.assertFalse(rows[1]['safe']);self.assertIn('隐藏',rows[1]['reason'])
 def test_onepass_word_adoption_counts_every_actual_round_and_parent(self):
  import time
  from web.session import WebSession
  from web.onepass_session import start,poll_fields
  from core.page_policy import DEFAULT
  s=WebSession(self.root/'smart');s.set_mode('academic');s.load(self.source.read_bytes(),'公开合成.docx');s.page_layout_draft={'enabled':True,'policy':{**DEFAULT,'pages':'front-body','footer':'chapter','body_start':5},'regions':[],'input_sha256':s.inventory['input_sha256']};s.field_updates_enabled=True;state={}
  with patch('core.word_native.calculate',side_effect=self.mock):
   start(s,state,structure_runner=lambda *_:(_ for _ in ()).throw(DemoError('no API')))
   self.assertEqual(state['status'],'updating',state.get('error'))
   for _ in range(300):
    poll_fields(s,state)
    if state['status']!='updating':break
    time.sleep(.02)
  self.assertEqual(state['status'],'done',state.get('error'));self.assertEqual(state['audit']['field_update']['status'],'updated');self.assertEqual(state['outcome']['rounds'],3);self.assertEqual(s.generation_count,3);steps=state['audit']['assembly_steps'];self.assertEqual(steps[-1]['kind'],'field_update');self.assertEqual(steps[-2]['output_sha256'],steps[-1]['input_sha256']);self.assertEqual(state['outcome']['changes'],sum(x['changes'] for x in steps));self.assertEqual(state['outcome']['output_sha256'],s.result['after']['input_sha256']);self.assertEqual(state['outcome']['summary'],s.result['after']['summary'])
if __name__=='__main__':unittest.main()
