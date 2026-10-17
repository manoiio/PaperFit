"""Native file extraction, atomic overlay, reusable package. All HTTP mocked."""
from unittest import TestCase
from tempfile import TemporaryDirectory
from pathlib import Path
from copy import deepcopy
import json
from web.session import WebSession
from web.requirements_input import file_input,saved_input
from web.onepass_session import start,continue_command,current
from ai.file_spec import request_for,adapt,parse_requirements,VERSION
from ai.short_spec import Settings
from core.package import DemoError
ROOT=Path(__file__).resolve().parents[1]
EXPECTED=json.loads((ROOT/'samples/phase44/expected-overrides.json').read_text())['fields']
def answer(extraction):
 r=request_for(extraction)
 return {'version':VERSION,'items':[{'field':fid,'value':EXPECTED[f['name']]['value'],'source':next(s['id'] for s in r['sources'] if EXPECTED[f['name']]['quote'] in s['text'])} for fid,f in r['fields'].items() if f['name'] in EXPECTED],'background':[],'issues':[]}
def parser(extraction):return parse_requirements(extraction,config=Settings('synthetic'),transport=lambda *_:json.dumps({'choices':[{'finish_reason':'stop','message':{'content':json.dumps(answer(extraction),ensure_ascii=False)}}]}))
class FileRequirementsAcceptance(TestCase):
 def setUp(self):self.tmp=TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.root=Path(self.tmp.name)
 def model(self,n='a'):
  s=WebSession(self.root/n);s.set_mode('academic');s.load((ROOT/'samples/phase44/supplement.docx').read_bytes(),'论文.docx');return s
 def source(self,s,kind='txt',extra=''):
  data=(ROOT/f'samples/phase46/requirements.{kind}').read_bytes()
  if extra:data+=extra.encode()
  return file_input(s,data,'老师要求.'+kind)
 def test_txt_docx_same_16_values_real_positions_and_original_input_unchanged(self):
  for kind in ('txt','docx'):
   s=self.model(kind);src=self.source(s,kind);e=src['extraction'];r=parser(e)
   self.assertEqual(len(r['draft']['fields']),16);self.assertFalse(r['notices'])
   for n,f in r['draft']['fields'].items():self.assertEqual(f['value'],EXPECTED[n]['value']);self.assertEqual(f['location'],'TXT:L1' if kind=='txt' else 'DOCX:P1')
   old=s.input_path.read_bytes();state={};start(s,state,requirements=src,file_parser=parser)
   self.assertEqual(state['status'],'done',state.get('error'));self.assertEqual(s.input_path.read_bytes(),old)
   self.assertEqual(state['audit']['API']['requirements']['source_sha256'],src['extraction']['file_sha256']);self.assertTrue(current(s,state,'',src));self.assertFalse(current(s,state,'',self.source(s,'txt','\n本说明用于测试。')))
   self.assertEqual(s.downloads()['fixed.docx'],(ROOT/'samples/phase45/supplement.docx').read_bytes())
 def test_partial_background_unsupported_and_conflict_require_explicit_decision(self):
  s=self.model();src=self.source(s,extra='\n本说明用于测试。\n页码从1开始。');e=src['extraction'];d=answer(e);sources=request_for(e)['sources'];d['background']=[sources[-2]['id']];d['issues']=[{'source':sources[-1]['id'],'reason':'页码生成不支持'}]
  r=adapt(d,e);self.assertTrue(r['notices']);self.assertEqual(len(r['draft']['fields']),16)
  state={};before=s.rules_hash;start(s,state,requirements=src,file_parser=lambda _:dict(r,receipt={'kind':'mock_api','model':'deepseek-flash','timestamp_utc':'2026-10-17T00:00:00+00:00'}))
  self.assertEqual(state['status'],'waiting',state.get('error'));self.assertEqual(s.rules_hash,before);self.assertFalse(s.history)
  continue_command(s,state,'',{'requirements_partial':'partial'},requirements=src);self.assertEqual(state['status'],'done',state.get('error'))
  t=self.model('template');st={};src=self.source(t,extra='\n页码从1开始。');e=src['extraction'];d=answer(e);d['issues']=[{'source':request_for(e)['sources'][-1]['id'],'reason':'不支持'}];r=adapt(d,e)
  start(t,st,requirements=src,file_parser=lambda _:dict(r,receipt={'kind':'mock_api','model':'deepseek-flash','timestamp_utc':'2026-10-17T00:00:00+00:00'}));continue_command(t,st,'',{'requirements_partial':'template'},requirements=src);self.assertEqual(st['status'],'done',st.get('error'));self.assertTrue(st['audit']['file_requirements_declined'])
  e=self.source(self.model('conflict'),extra='\n正文三号。')['extraction'];r=adapt(answer(e),e);self.assertNotIn('font_size',r['draft']['fields']);self.assertTrue(r['notices'])
 def test_file_rejection_protocol_missing_and_cache_and_reuse_zero_parser(self):
  s=self.model()
  for raw in (b'\xffbad',('正文小四。'*3000).encode()):
   with self.assertRaises(DemoError):file_input(s,raw,'规范.txt')
  src=self.source(s);d=answer(src['extraction']);d['items'].append(d['items'][0])
  with self.assertRaises(DemoError):adapt(d,src['extraction'])
  d=answer(src['extraction']);d['items'].pop();self.assertTrue(adapt(d,src['extraction'])['notices'])
  st={};cache={};calls=[]
  def p(e):calls.append(1);return parser(e)
  start(s,st,requirements=src,requirements_cache=cache,file_parser=p);start(s,st,requirements=src,requirements_cache=cache,file_parser=p);self.assertEqual(calls,[1]);self.assertEqual(len(s.history),1)
  other=self.model('reuse');bundle=s.rule_bundle_download();saved=saved_input(other,bundle,'排版要求.json');state={};start(other,state,requirements=saved,file_parser=lambda _:self.fail('saved rules must not parse'))
  self.assertEqual(state['status'],'done',state.get('error'));self.assertEqual(other.downloads()['fixed.docx'],s.downloads()['fixed.docx'])
  local=self.model('local');local.import_rule_bundle(bundle);local.apply_workbench(acknowledged=True);self.assertEqual(local.current_rule_bundle()['provenance'],s.current_rule_bundle()['provenance'])
 def test_cancelled_waiting_parse_cache_and_cross_session_no_authorization(self):
  from web.onepass_session import cancel
  s=self.model();src=self.source(s,extra='\n页码从1开始。');e=src['extraction'];data=answer(e);data['issues']=[{'source':request_for(e)['sources'][-1]['id'],'reason':'未支持'}]
  calls=[]
  def p(_):calls.append(1);return dict(adapt(data,e),receipt={'kind':'mock_api','model':'deepseek-flash','timestamp_utc':'2026-10-17T00:00:00+00:00','request_count':1})
  cache={};state={};start(s,state,requirements=src,file_parser=p,requirements_cache=cache);self.assertEqual(state['status'],'waiting');cancel(s,state)
  start(s,state,requirements=src,file_parser=p,requirements_cache=cache);self.assertEqual(calls,[1]);self.assertTrue(state['audit']['API']['requirements']['receipt']['cache_hit']);self.assertFalse(s.history)
  other=self.model('other');self.assertFalse(current(other,state,'',src))
  changed=deepcopy(src);changed['extraction']['segments'][0]['text']='正文三号'
  with self.assertRaises(DemoError):start(other,{},requirements=changed,file_parser=p)

 def test_background_laundering_wrong_source_types_and_complex_extraction(self):
  s=self.model();e=self.source(s,extra='\n页码从1开始。')['extraction'];d=answer(e);d['background']=[request_for(e)['sources'][-1]['id']];self.assertTrue(adapt(d,e)['notices'])
  for mutation in ('unknown','value','extra'):
   d=answer(e)
   if mutation=='unknown':d['items'][0]['source']='invented'
   if mutation=='value':d['items'][0]['value']=True
   if mutation=='extra':d['confirmed']=True
   if mutation=='value':self.assertTrue(adapt(d,e)['notices'])
   else:
    with self.assertRaises(DemoError):adapt(d,e)
  from docx import Document
  from io import BytesIO
  from docx.oxml import OxmlElement
  doc=Document();doc.add_paragraph('正文小四。');p=doc.add_paragraph();p._p.append(OxmlElement('w:fldSimple'));b=BytesIO();doc.save(b)
  source=file_input(s,b.getvalue(),'复杂.docx');self.assertTrue(source['extraction']['unextracted'])
  doc=Document();doc.add_paragraph('正文小四。');doc.sections[0].header.paragraphs[0].text='页眉中的规范';b=BytesIO();doc.save(b);self.assertTrue(file_input(s,b.getvalue(),'页眉.docx')['extraction']['unextracted'])
