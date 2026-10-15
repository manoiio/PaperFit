"""第五阶段实际行为验收。所有自动API使用模拟响应，绝不消耗余额。"""
from pathlib import Path
from copy import deepcopy
from tempfile import TemporaryDirectory
from unittest.mock import patch
from urllib.error import HTTPError
from io import BytesIO
import hashlib,json,socket,unittest
from docx import Document
from streamlit.testing.v1 import AppTest
from core import load_rules,load_custom_rules,check_document,fix_document
from core.engine import _ensure,DEFAULT_RULES
from core.package import DemoError,Package,q,NS,write_new
from core.norms import extract_norm,selected_segments,marked_text
from core.templates import FIELDS,validate_draft,validate_template,make_template,defaults
from core.rule_config import strict_json
from ai.short_spec import Settings
from ai.template_spec import parse_template_requirements
from web.session import WebSession

ROOT=Path(__file__).resolve().parents[1]
CORRECT=ROOT/'samples/extended-correct.docx'
MANIFEST=json.loads(CORRECT.with_suffix('.types.json').read_text())
EXPECTED=json.loads((ROOT/'samples/norm-expected.json').read_text())
BODY=[int(n) for n,t in MANIFEST['paragraph_types'].items() if t=='body']
QUOTES_A={'paper':'纸张为A4竖向。','margin_top':'上下左右边距均为2.5cm。','margin_bottom':'上下左右边距均为2.5cm。','margin_left':'上下左右边距均为2.5cm。','margin_right':'上下左右边距均为2.5cm。','font_cn':'正文中文字体为宋体。','font_western':'正文西文字体为Times New Roman。','font_size':'正文普通文字字号为小四号。','line_spacing':'正文采用1.5倍行距。','first_line_indent':'正文首行缩进2字。'}
QUOTES_B={**QUOTES_A,'margin_top':'上边距3cm。','margin_bottom':'下边距2cm。','margin_left':'左边距3.5cm。','margin_right':'右边距3cm。','font_cn':'正文中文字体为黑体。','font_western':'正文西文字体为Arial。','font_size':'正文普通文字字号为四号。','line_spacing':'正文采用2倍行距。','first_line_indent':'正文首行缩进1字符。'}

def synthetic_draft(segments,kind='B'):
    quotes=QUOTES_A if kind=='A' else QUOTES_B
    draft={n:{'value':EXPECTED[kind][n],'unit':FIELDS[n][2],'quote':quotes[n],'location':next(s['location'] for s in segments if quotes[n] in s['text']),'support':'supported'} for n in FIELDS}
    draft['unsupported']=[{'quote':'请双面打印。','location':next(s['location'] for s in segments if '请双面打印。' in s['text']),'reason':'printing'}]
    if kind=='B':draft['unsupported'].append({'quote':'标题采用固定20磅行距。','location':next(s['location'] for s in segments if '标题采用固定20磅行距。' in s['text']),'reason':'title'})
    return draft

def fake_parse(segments,kind='B'):
    draft=synthetic_draft(segments,kind)
    return parse_template_requirements(segments,config=Settings('mock-secret'),transport=lambda request,timeout:json.dumps({'choices':[{'finish_reason':'stop','message':{'content':json.dumps(draft,ensure_ascii=False)}}]}).encode())

class TemplateAcceptance(unittest.TestCase):
    def setUp(self):
        self.tmp=TemporaryDirectory(dir=ROOT/'work/phase5');self.dir=Path(self.tmp.name)
        self.base=load_rules();self.model=WebSession(self.dir)
    def tearDown(self):self.tmp.cleanup()
    def norm(self,kind='B'):
        path=ROOT/'samples'/('norm-A.txt' if kind=='A' else 'norm-B.docx')
        self.assertTrue(self.model.load_norm(path.read_bytes(),path.name),self.model.template_error)
        self.model.select_norm_segments([s['location'] for s in self.model.norm_inventory['segments']]);self.model.confirm_norm_text()
        return self.model.norm_segments
    def applied(self,kind='B',paper=True):
        self.norm(kind);m=self.model
        self.assertTrue(m.parse_template(parser=lambda s:fake_parse(s,kind)),m.template_error)
        self.assertTrue(m.apply_template(),m.template_error)
        if paper:
            m.set_mode('extended');self.assertTrue(m.load(CORRECT.read_bytes(),CORRECT.name),m.error)
            m.set_types(MANIFEST['paragraph_types']);m.confirm();self.assertTrue(m.check(),m.error)
        return m
    def template(self,kind='B'):
        return self.applied(kind,paper=False).active_template

    def test_txt_encodings_are_strict_and_no_silent_byte_loss(self):
        text='正文中文字体为宋体。\n正文小四号。'
        for data,codec in [(text.encode(),'auto'),(text.encode('utf-16'),'auto'),(text.encode('gb18030'),'gb18030')]:
            result=extract_norm(data,'txt',encoding=codec)
            self.assertEqual([s['text'] for s in result['segments']],text.splitlines())
        for data in (b'\xffbroken',b'hello\x00world',text.encode('gb18030')):
            with self.assertRaises(DemoError):extract_norm(data,'txt')
        with self.assertRaises(DemoError):extract_norm('正文\ufffd'.encode(),'txt')

    def test_docx_order_and_exact_paragraph_table_source_locations(self):
        raw=(ROOT/'samples/norm-B.docx').read_bytes()
        p=self.dir/'norm.docx';p.write_bytes(raw)
        inv=extract_norm(raw,'docx',path=p);segments=inv['segments']
        self.assertEqual(segments[2],{'location':'DOCX:P3','text':'纸张为A4竖向。'})
        self.assertEqual(segments[4],{'location':'DOCX:T1:R1:C2:P1','text':'上边距3cm。'})
        self.assertEqual(segments[-2]['text'],'标题采用固定20磅行距。')
        self.assertEqual(segments[-1]['location'],'DOCX:P5')
        self.assertEqual(inv['file_sha256'],hashlib.sha256(raw).hexdigest())
        with self.assertRaises(DemoError):extract_norm(raw+b'changed','docx',path=p)

    def test_select_subset_edit_rebind_and_length_never_truncated(self):
        self.norm();m=self.model
        loc=next(s['location'] for s in m.norm_segments if '上边距3cm' in s['text'])
        m.select_norm_segments([loc]);self.assertEqual(len(m.norm_segments),1)
        self.assertIsNone(m.norm_confirmed_hash)
        m.set_norm_text('正文四号。');m.confirm_norm_text()
        self.assertEqual(m.norm_segments,[{'location':'摘录:L1','text':'正文四号。'}])
        with self.assertRaises(DemoError):selected_segments('文'*12001)
        with self.assertRaises(DemoError):m.select_norm_segments(['DOCX:P999'])
        self.assertEqual(m.norm_inventory['file_sha256'],hashlib.sha256(m.norm_bytes).hexdigest())

    def test_cm_and_chinese_size_conversion_are_evidence_checked(self):
        segments=self.norm();draft=synthetic_draft(segments)
        validate_draft(draft,segments)
        self.assertEqual(draft['margin_left']['value'],35);self.assertEqual(draft['font_size']['value'],14)
        for name,value in [('margin_left',3.5),('font_size',12)]:
            bad=deepcopy(draft);bad[name]['value']=value
            with self.assertRaises(DemoError):validate_draft(bad,segments)

    def test_independent_margins_missing_fields_require_explicit_decisions(self):
        s=selected_segments('左右边距3cm。');d={n:None for n in FIELDS};d['unsupported']=[]
        for n in ('margin_left','margin_right'):d[n]={'value':30,'unit':'mm','quote':'左右边距3cm。','location':'摘录:L1','support':'supported'}
        validate_draft(d,s)
        m=self.model;m.set_norm_text(s[0]['text']);m.confirm_norm_text()
        self.assertTrue(m.parse_template(parser=lambda _: {'draft':d,'receipt':{'kind':'mock_api','model':'deepseek-flash','timestamp_utc':'2026-10-14T00:00:00+00:00'}}))
        self.assertIsNone(m.template_draft['margin_top']);self.assertEqual(m.template_modes['margin_top'],'choose')
        self.assertFalse(m.apply_template())
        for n in FIELDS:
            if d[n] is None:m.edit_template(n,'default')
        self.assertTrue(m.apply_template(),m.template_error)
        values=m.active_template['final_values'];self.assertEqual((values['margin_top'],values['margin_bottom'],values['margin_left'],values['margin_right']),(25,25,30,30))
        self.assertEqual(m.active_template['field_sources']['margin_top']['kind'],'default_inherited')

    def test_nonbody_fixed_indent_units_and_conflicts_are_not_executable(self):
        cases=[('font_size',14,'标题四号。'),('font_cn','宋体','摘要使用宋体。'),('line_spacing',2,'正文固定20磅行距。'),('first_line_indent',2,'正文首行缩进24磅。'),('paper','A4 portrait','纸张A4横向。'),('font_cn','宋体','正文使用微软雅黑。'),('font_size',12,'正文12pt。正文14pt。'),('margin_top',30,'上边距3cm。上边距2cm。'),('font_cn','黑体','正文中文字体为方正黑体。'),('font_cn','楷体','正文中文字体为楷体_GB2312。'),('font_western','Arial','正文西文字体为Arial Narrow。')]
        for name,value,text in cases:
            segments=selected_segments(text);d={n:None for n in FIELDS};d['unsupported']=[]
            d[name]={'value':value,'unit':FIELDS[name][2],'quote':text.split('。')[0]+'。','location':'摘录:L1','support':'supported'}
            with self.subTest(name=name,text=text),self.assertRaises(DemoError):validate_draft(d,segments)
        segments=selected_segments('正文无首行缩进。');d={n:None for n in FIELDS};d['unsupported']=[]
        d['first_line_indent']={'value':0,'unit':'chars','quote':segments[0]['text'],'location':'摘录:L1','support':'supported'}
        validate_draft(d,segments)

    def test_unknown_fields_invalid_types_values_and_forged_evidence(self):
        s=self.norm();d=synthetic_draft(s)
        for name,values in {'margin_top':[True,'30',9,51,float('inf')],'font_size':[True,13.3,7.5,float('nan'),10**1000],'line_spacing':[1.1,3,False],'first_line_indent':[-1,3,'1',True],'paper':['A3','A4 landscape'],'font_cn':['微软雅黑'],'font_western':['Cambria']}.items():
            for value in values:
                bad=deepcopy(d);bad[name]['value']=value
                with self.subTest(name=name,value=value),self.assertRaises(DemoError):validate_draft(bad,s)
        for key,val in [('unknown',{}),('tool_calls',[])]:
            bad=deepcopy(d);bad[key]=val
            with self.assertRaises(DemoError):validate_draft(bad,s)
        for key,val in [('quote','正文普通字号16pt'),('location','DOCX:P999'),('unit','cm'),('support','execute')]:
            bad=deepcopy(d);bad['font_size'][key]=val
            with self.assertRaises(DemoError):validate_draft(bad,s)
        with self.assertRaises(DemoError):selected_segments('执行代码 /Users/private/file')

    def test_user_edits_and_default_inheritance_source_accuracy(self):
        m=self.applied(paper=False)
        m.edit_template('font_size','ai',13.5,touched=True)
        m.edit_template('margin_top','default')
        data=m.proposed_template()
        self.assertEqual(data['field_sources']['font_size']['kind'],'user_confirmed')
        self.assertEqual(data['field_sources']['font_size']['quote'],QUOTES_B['font_size'])
        self.assertEqual(data['field_sources']['margin_top']['kind'],'default_inherited')
        self.assertEqual(data['final_values']['margin_top'],25)
        self.assertEqual(m.active_template['final_values']['font_size'],14) # draft does not change active
        self.assertTrue(m.apply_template());self.assertEqual(load_custom_rules(m.custom_rules_path)['rules']['R05']['expected'],13.5)

    def test_a_b_change_actual_engine_and_b_reread_matches_every_field(self):
        m=self.applied('A');self.assertEqual(m.report['summary']['problem_count'],0)
        self.applied('B');before=deepcopy(m.report);self.assertGreater(before['summary']['problem_count'],0)
        self.assertEqual({f['rule_id'] for f in before['findings'] if f['status']=='problem'}, {'R02','R03','R04','R05','R06','R07'})
        m.set_selection([f['key'] for f in m.eligible()]);self.assertTrue(m.generate(),m.error)
        result=m.result;self.assertEqual(result['after']['summary']['problem_count'],0);self.assertEqual(result['after']['summary']['review_count'],0)
        for rid,value in [('R03','黑体'),('R04','Arial'),('R05',14),('R06',2),('R07',1)]:
            self.assertTrue(all(f['expected']==value and f['status']=='pass' for f in result['after']['findings'] if f['rule_id']==rid and f['status']!='not_applicable'))
        pkg=Package(result['output']);res=__import__('core.formats',fromlist=['Resolver']).Resolver(pkg)
        for n in BODY:
            self.assertEqual(res.attrs(res.layers(pkg.paragraphs[n-1],kind='pPr'),'ind')['firstLineChars'],'100')
        independent=check_document(result['output'],result['output_body_manifest'],custom_rules_path=m.custom_rules_path,mode='extended',classification_manifest=result['output_classification_manifest'])
        self.assertEqual(independent['summary']['problem_count'],0)
        self.assertEqual(independent['summary']['review_count'],0)
        self.assertEqual(m.original_bytes,CORRECT.read_bytes())
        self.assertEqual(result['round']['rules_sha256'],result['after']['rules']['sha256'])
        old=Package(CORRECT)
        self.assertEqual({k:v for k,v in old.parts.items() if k!='word/document.xml'},{k:v for k,v in pkg.parts.items() if k!='word/document.xml'})
        self.assertEqual(old.root.xpath('.//w:t/text()',namespaces=NS),pkg.root.xpath('.//w:t/text()',namespaces=NS))
        self.assertEqual([__import__('lxml').etree.tostring(t) for t in old.root.findall('.//w:tbl',NS)],[__import__('lxml').etree.tostring(t) for t in pkg.root.findall('.//w:tbl',NS)])
        self.assertEqual(result['preservation']['other_zip_member_bytes_equal'],True)
        self.assertTrue(all(f['status']=='pass' for f in independent['findings'] if f['rule_id'] in ('R08','R09','R10','R11','R12')))

    def test_zero_one_two_character_indent_actual_correct_wrong_and_fix(self):
        for chars in (0,1,2):
            with self.subTest(chars=chars):
                m=self.applied(paper=False);m.edit_template('first_line_indent','manual',chars,touched=True);self.assertTrue(m.apply_template(),m.template_error)
                pkg=Package(CORRECT)
                for n in BODY:
                    ind=_ensure(_ensure(pkg.paragraphs[n-1],'pPr'),'ind');ind.attrib.clear();ind.set(q('firstLineChars'),str(chars*100));ind.set(q('firstLine'),'0')
                p=self.dir/f'indent-good-{chars}.docx';write_new(p,pkg.bytes(True))
                mfest={'input_sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'body_paragraphs':BODY}
                report=check_document(p,mfest,custom_rules_path=m.custom_rules_path)
                self.assertTrue(all(f['status']=='pass' for f in report['findings'] if f['rule_id']=='R07'))
                ind.set(q('firstLineChars'),str(((chars+1)%3)*100))
                bad=self.dir/f'indent-bad-{chars}.docx';write_new(bad,pkg.bytes(True));mfest['input_sha256']=hashlib.sha256(bad.read_bytes()).hexdigest()
                report=check_document(bad,mfest,custom_rules_path=m.custom_rules_path)
                self.assertEqual(sum(f['rule_id']=='R07' and f['status']=='problem' for f in report['findings']),1)
                result=fix_document(bad,self.dir/f'indent-fixed-{chars}.docx',mfest,selected_rules=['R07'],custom_rules_path=m.custom_rules_path)
                self.assertTrue(all(f['status']=='pass' for f in result['after']['findings'] if f['rule_id']=='R07'))

    def test_zero_absent_and_one_point_indent_and_conflict(self):
        for chars in (0,1):
            m=self.applied(paper=False);m.edit_template('first_line_indent','manual',chars,touched=True);m.apply_template()
            pkg=Package(CORRECT)
            for node in pkg.styles.findall('.//w:ind',NS):node.getparent().remove(node)
            pkg.parts['word/styles.xml']=__import__('lxml').etree.tostring(pkg.styles)
            for n in BODY:
                ind=_ensure(_ensure(pkg.paragraphs[n-1],'pPr'),'ind');ind.attrib.clear()
                if chars:ind.set(q('firstLine'),'240') # 12pt = one character before fontsize fix
            path=self.dir/f'point-{chars}.docx';write_new(path,pkg.bytes(True))
            r=check_document(path,BODY,custom_rules_path=m.custom_rules_path)
            self.assertTrue(all(f['status']=='pass' for f in r['findings'] if f['rule_id']=='R07'))
            if chars==0:
                ind.set(q('firstLineChars'),'0');ind.set(q('firstLine'),'240')
                bad=self.dir/'zero-conflict.docx';write_new(bad,pkg.bytes(True))
                r=check_document(bad,BODY,custom_rules_path=m.custom_rules_path)
                self.assertEqual(next(f for f in r['findings'] if f['key']==f'R07:paragraph={BODY[-1]}')['status'],'problem')

    def test_roundtrip_template_same_rules_and_results_without_api(self):
        m=self.applied();before=deepcopy(m.report)
        raw=m.template_download();self.assertNotIn(str(ROOT).encode(),raw);self.assertNotIn(b'API_KEY',raw)
        second=WebSession(self.dir)
        with patch('ai.short_spec.http_post',side_effect=AssertionError('no API allowed')):
            self.assertTrue(second.import_template(raw),second.template_error)
            self.assertIsNone(second.active_template)
            self.assertTrue(second.apply_template(),second.template_error)
            second.set_mode('extended');second.load(CORRECT.read_bytes(),CORRECT.name);second.set_types(MANIFEST['paragraph_types']);second.confirm();self.assertTrue(second.check())
        self.assertEqual(m.active_template['final_values'],second.active_template['final_values'])
        self.assertEqual([(f['key'],f['expected'],f['status']) for f in before['findings']],[(f['key'],f['expected'],f['status']) for f in second.report['findings']])
        self.assertTrue(second.active_template['confirmation']['imported']);self.assertIn('导入模板',second.report['findings'][0]['source']['label'])
        self.assertNotEqual(second.task_dir,m.task_dir);self.assertNotEqual(second.norm_dir,m.norm_dir)

    def test_invalid_templates_reject_keys_version_types_provenance_and_paths(self):
        data=self.template()
        for mutation in ('version','unknown','field','bool','range','name','default','quote','path','confirmation','unicode'):
            bad=deepcopy(data)
            if mutation=='version':bad['structure_version']=True
            if mutation=='unknown':bad['execute']='code'
            if mutation=='field':bad['final_values']['R08']=18
            if mutation=='bool':bad['final_values']['margin_left']=True
            if mutation=='range':bad['final_values']['first_line_indent']=3
            if mutation=='name':bad['name']='学校官方规范'
            if mutation=='default':bad['field_sources']['margin_left']={'kind':'default_inherited','quote':'','location':None}
            if mutation=='quote':bad['field_sources']['font_size']['quote']='伪造正文14pt'
            if mutation=='path':bad['name']='/Users/private/file'
            if mutation=='confirmation':bad['confirmation']['tool_calls']=[]
            if mutation=='unicode':bad['norm_excerpt'][0]['text']='\ud800'
            with self.subTest(mutation=mutation),self.assertRaises(DemoError):validate_template(bad,self.base)
        for raw in (b'{"x":NaN}',b'{"x":1,"x":2}',b'\xff'):
            self.assertFalse(self.model.import_template(raw));self.assertIsNone(self.model.template_import)

    def test_norm_template_change_invalidates_reports_downloads_and_confirmation(self):
        m=self.applied();m.set_selection([f['key'] for f in m.eligible()]);m.generate();self.assertTrue(m.downloads())
        active=deepcopy(m.active_template)
        m.set_norm_text('正文五号。');self.assertIsNone(m.report);self.assertFalse(m.downloads());self.assertIsNone(m.confirmed);self.assertIsNone(m.template_draft)
        self.assertEqual(m.active_template,active)
        self.assertTrue(m.import_template(json.dumps(active,ensure_ascii=False).encode()));self.assertIsNone(m.report);self.assertFalse(m.downloads())
        self.assertTrue(m.apply_template());self.assertIsNone(m.confirmed)
        m.use_default_rules();self.assertIsNone(m.active_template);self.assertIsNone(m.custom_rules_path)
        m.confirm();m.check();self.assertEqual(m.report['summary']['problem_count'],0)

    def test_actual_request_only_selected_norm_and_one_call_no_secret_record(self):
        s=self.norm();s=[next(p for p in s if '正文中文字体为黑体。'==p['text'])]
        d={n:None for n in FIELDS};d['unsupported']=[];d['font_cn']={'value':'黑体','unit':'font','quote':s[0]['text'],'location':s[0]['location'],'support':'supported'}
        calls=[]
        def transport(req,timeout):
            calls.append((req,timeout));return json.dumps({'choices':[{'finish_reason':'stop','message':{'content':json.dumps(d,ensure_ascii=False)}}]}).encode()
        result=parse_template_requirements(s,config=Settings('never-log-secret'),transport=transport)
        self.assertEqual(len(calls),1);self.assertEqual(calls[0][1],30)
        payload=json.loads(calls[0][0].data);self.assertEqual(payload['messages'][1]['content'],marked_text(s))
        self.assertEqual(payload['response_format'],{'type':'json_object'})
        serialized=json.dumps(result,ensure_ascii=False);self.assertNotIn('never-log-secret',serialized);self.assertNotIn(str(CORRECT),serialized)
        self.assertNotIn('软件演示正文',calls[0][0].data.decode())

    def test_api_failure_preserves_active_and_fixed_flows_and_no_retries(self):
        m=self.applied();active=deepcopy(m.active_template)
        m.set_norm_text('正文五号。');m.confirm_norm_text()
        for failure in (socket.timeout(),HTTPError('https://api.deepseek.com',402,'private',{},BytesIO(b''))):
            calls=[]
            def transport(req,timeout):calls.append(1);raise failure
            self.assertFalse(m.parse_template(parser=lambda s:parse_template_requirements(s,transport=transport,config=Settings('secret'))))
            self.assertEqual(len(calls),1);self.assertEqual(m.active_template,active);self.assertIsNone(m.template_draft)
            m.confirm();self.assertTrue(m.check(),m.error)
        m.use_default_rules();m.confirm();m.check();self.assertEqual(m.report['summary']['problem_count'],0)

    def test_sessions_rounds_unselected_and_files_isolated(self):
        m=self.applied();m.set_selection([f['key'] for f in m.eligible() if f['rule_id']=='R05']);m.generate()
        first=Path(m.result['output']);first_bytes=first.read_bytes();first_cfg=m.result['after']['rules']['sha256']
        before=Package(CORRECT);after=Package(first)
        for a,b in zip(before.sections,after.sections):self.assertEqual(__import__('lxml').etree.tostring(a),__import__('lxml').etree.tostring(b))
        for n in BODY:
            self.assertEqual([__import__('lxml').etree.tostring(p) for p in before.paragraphs[n-1].findall(q('pPr'))],[__import__('lxml').etree.tostring(p) for p in after.paragraphs[n-1].findall(q('pPr'))])
        self.assertGreater(m.result['after']['summary']['problem_count'],0)
        self.assertTrue(m.next_round());m.set_selection([f['key'] for f in m.eligible()]);self.assertTrue(m.generate(),m.error)
        self.assertEqual(m.result['after']['summary']['problem_count'],0);self.assertEqual(first.read_bytes(),first_bytes)
        self.assertEqual(m.result['round']['input_round'],1);self.assertEqual(first_cfg,m.result['round']['rules_sha256'])
        second=WebSession(self.dir);second.load(CORRECT.read_bytes(),CORRECT.name);second.set_body(BODY);second.confirm();second.check()
        self.assertEqual(second.report['summary']['problem_count'],0);self.assertIsNone(second.active_template);self.assertEqual(second.norm_text,'')
        second.reset();self.assertTrue(first.exists());self.assertTrue(m.downloads())
        self.assertEqual(m.original_bytes,CORRECT.read_bytes());self.assertEqual(load_rules()['sha256'],hashlib.sha256(DEFAULT_RULES.read_bytes()).hexdigest())

    def test_norm_file_tamper_and_repeat_parse_fail_closed(self):
        self.norm();m=self.model;calls=[]
        self.assertTrue(m.parse_template(parser=lambda s:(calls.append(1) or fake_parse(s))))
        self.assertFalse(m.parse_template(parser=lambda s:(calls.append(1) or fake_parse(s))));self.assertEqual(len(calls),1)
        m.norm_path.write_bytes(b'changed');m.set_norm_text('正文四号。');m.confirm_norm_text()
        self.assertFalse(m.parse_template(parser=lambda s:(calls.append(1) or fake_parse(s))));self.assertEqual(len(calls),1)


    def test_font_slots_are_separate_and_unselected_western_is_retained(self):
        m=self.applied();m.set_selection([f['key'] for f in m.eligible() if f['rule_id']=='R03']);self.assertTrue(m.generate(),m.error)
        from core.formats import Resolver
        old=Package(CORRECT);new=Package(m.result['output']);a=Resolver(old);b=Resolver(new)
        for n in BODY:
            for r1,r2 in zip(old.paragraphs[n-1].findall(q('r')),new.paragraphs[n-1].findall(q('r'))):
                x=a.attrs(a.layers(old.paragraphs[n-1],r1),'rFonts');y=b.attrs(b.layers(new.paragraphs[n-1],r2),'rFonts')
                self.assertEqual(x.get('ascii'),y.get('ascii'));self.assertEqual(x.get('hAnsi'),y.get('hAnsi'));self.assertEqual(x.get('cs'),y.get('cs'))
        self.assertTrue(all(f['status']=='pass' for f in m.result['after']['findings'] if f['rule_id']=='R03' and f['status']!='not_applicable'))
        self.assertTrue(any(f['status']=='problem' for f in m.result['after']['findings'] if f['rule_id']=='R04'))

    def test_a4_landscape_attribute_detected_and_only_selected_page_fixed(self):
        m=self.applied('A',paper=False);pkg=Package(CORRECT)
        pg=pkg.sections[0].find(q('pgSz'));pg.set(q('orient'),'landscape')
        path=self.dir/'landscape.docx';write_new(path,pkg.bytes(True))
        before=check_document(path,BODY,custom_rules_path=m.custom_rules_path)
        finding=next(f for f in before['findings'] if f['key']=='R01:section=1')
        self.assertEqual(finding['status'],'problem');self.assertEqual(finding['actual']['orientation'],'landscape')
        result=fix_document(path,self.dir/'portrait.docx',BODY,selected_keys=[finding['key']],custom_rules_path=m.custom_rules_path)
        self.assertEqual(result['after']['summary']['problem_count'],0)
        self.assertTrue(result['preservation']['document_except_selected_format_attributes_equal'])
        self.assertEqual(path.read_bytes(),pkg.bytes(True))

    def test_caption_context_and_local_page_scope_cannot_supply_body_rules(self):
        for name,value,text,quote in [('font_size',12,'正文小四号，题注12pt。','12pt'),('margin_left',30,'摘要左边距3cm。','左边距3cm'),('margin_top',30,'标题上边距3cm。','上边距3cm'),('font_cn','宋体','正文西文字体为宋体。','宋体'),('font_western','Arial','正文中文字体为Arial。','Arial'),('margin_left',30,'正文段落缩进：左3cm。','左3cm')]:
            s=selected_segments(text);d={n:None for n in FIELDS};d['unsupported']=[]
            d[name]={'value':value,'unit':FIELDS[name][2],'quote':quote,'location':'摘录:L1','support':'supported'}
            with self.subTest(name=name),self.assertRaises(DemoError):validate_draft(d,s)
        from core.templates import local_review_notes
        notes=local_review_notes(selected_segments('正文固定20磅行距。\n正文使用微软雅黑。\n纸张横向。\n正文首行缩进2cm。'))
        self.assertTrue({'fixed_spacing','unknown_font','landscape','indent_unit'}<= {n['reason'] for n in notes})

    def test_applied_template_tamper_invalidates_and_never_exports_changed_file(self):
        m=self.applied();raw=m.template_download();data=json.loads(raw)
        data['final_values']['font_size']=13.5;data['field_sources']['font_size']['kind']='user_confirmed'
        m.custom_rules_path.write_text(json.dumps(data,ensure_ascii=False),encoding='utf-8')
        m.sync_rules();self.assertIsNone(m.report);self.assertIsNone(m.rules_hash);self.assertIsNone(m.confirmed)
        self.assertIsNone(m.template_download())
        with self.assertRaises(DemoError):m.downloads()
        m.use_default_rules();m.confirm();self.assertTrue(m.check());self.assertEqual(m.report['summary']['problem_count'],0)

    def test_docx_line_break_tabs_and_complex_norm_objects_reported(self):
        doc=Document();p=doc.add_paragraph();p.add_run('正文').add_tab();p.add_run('小四号').add_break();p.add_run('正文1.5倍行距')
        table=doc.add_table(rows=1,cols=1);table.cell(0,0).text='左右边距3cm'
        raw=BytesIO();doc.save(raw);path=self.dir/'breaks.docx';path.write_bytes(raw.getvalue())
        inv=extract_norm(raw.getvalue(),'docx',path=path)
        self.assertEqual(inv['segments'][0]['text'],'正文\t小四号\n正文1.5倍行距')
        self.assertEqual(inv['segments'][1]['location'],'DOCX:T1:R1:C1:P1')
        pkg=Package(path);_ensure(pkg.paragraphs[0].find(q('r')),'fldChar').set(q('fldCharType'),'begin')
        complex_path=self.dir/'field.docx';write_new(complex_path,pkg.bytes(True))
        inv=extract_norm(complex_path.read_bytes(),'docx',path=complex_path)
        self.assertEqual(len(inv['segments']),1);self.assertTrue(any('DOCX:P1' in note for note in inv['notes']))

    def test_web_import_parse_confirm_and_reload_uses_same_engine(self):
        app=AppTest.from_file(str(ROOT/'app.py'),default_timeout=10).run()
        self.assertFalse(app.exception)
        def button(label):return next(b for b in app.button if b.label==label)
        button('导入演示规范A（TXT）').click().run();self.assertFalse(app.exception)
        m=app.session_state.workflow
        self.assertEqual(m.norm_inventory['kind'],'txt');self.assertIn('小四号',m.norm_text)
        button('已核对，确认待发送文本').click().run()
        with patch('ai.template_spec.parse_template_requirements',side_effect=lambda s:fake_parse(s,'A')):
            button('解析R01–R07要求').click().run()
        self.assertFalse(app.exception);self.assertIsNotNone(m.template_draft)
        button('确认并应用完整模板').click().run();self.assertFalse(app.exception)
        self.assertEqual(m.active_template['final_values'],EXPECTED['A'])
        button('使用正确对照').click().run()
        button('确认正文范围').click().run();button('开始检查').click().run()
        self.assertFalse(app.exception);self.assertEqual(m.report['summary']['problem_count'],0)
        # User edit becomes manual; applying invalidates actual report and original download binding.
        value=next(w for w in app.number_input if w.label=='正文普通文字字号值')
        value.set_value(14.0).run();self.assertFalse(app.exception)
        self.assertEqual(m.active_template['final_values']['font_size'],12)
        button('确认并应用完整模板').click().run();self.assertIsNone(m.report);self.assertFalse(m.downloads())
        m.reset()

if __name__=='__main__':unittest.main()
