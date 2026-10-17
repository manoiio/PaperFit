"""All automatic API tests mocked; original scope and repair remain authoritative."""
from pathlib import Path
from tempfile import TemporaryDirectory
from copy import deepcopy
from unittest import TestCase
from unittest.mock import patch
import json
from ai.structure import validate,redact,MAX_CANDIDATES,MAX_INPUT
from ai.short_spec import Settings
from web.structure_session import prepare,run,merge
from web.smart_session import suggest,accept_scope,execute
from web.session import WebSession
from core.package import DemoError,sha256
ROOT=Path(__file__).resolve().parents[1]
CONFIG=Settings('synthetic-placeholder')

class StructureAcceptance(TestCase):
 def setUp(self):self.tmp=TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.root=Path(self.tmp.name)
 def session(self,name='s',mode='academic'):
  s=WebSession(self.root/name);s.set_mode(mode);s.load((ROOT/'samples/phase41/structure.docx').read_bytes(),'合成.docx');return s
 def response(self,p,change=None):
  labels=json.loads((ROOT/'samples/phase41/labels.json').read_text())['paragraphs']
  rows=[{'id':c['id'],'role':labels[int(c['id'][2:])-1]['role'],'evidence':c['text'],'reason':'合成模拟：对应独立作者标注与局部章节结构'} for c in p['request']['candidates']]
  if change:change(rows)
  return json.dumps({'choices':[{'finish_reason':'stop','message':{'content':json.dumps({'suggestions':rows},ensure_ascii=False)}}],'usage':{'prompt_tokens':120,'completion_tokens':80,'total_tokens':200}})
 def test_local_protection_and_exact_sent_content(self):
  s=self.session();s.exclude_materials([17],'作者代码保护');s.batch_types([13],'heading3');old=deepcopy(s._scope_snapshot());p=prepare(s)
  self.assertEqual(s._scope_snapshot(),old);ids={c['id'] for c in p['request']['candidates']}
  self.assertTrue({'p:9','p:10','p:11','p:12','p:14','p:15','p:16'}<=ids)
  self.assertFalse(ids&{'p:2','p:13','p:17','p:18','p:19'});self.assertNotIn('20260001234',json.dumps(p));self.assertIsNone(s.confirmed)
  self.assertEqual(redact('姓名：小明 邮箱：ab@host.example /Users/someone/secret.docx'), '姓名：[个人信息已隐藏]')
 def test_one_post_cache_merge_and_original_bytes(self):
  s=self.session();p=prepare(s);calls=[];store={}
  def transport(req,timeout):
   calls.append(req);d=json.loads(req.data);self.assertEqual(d['thinking'],{'type':'disabled'});self.assertEqual(d['response_format'],{'type':'json_object'});return self.response(p)
  result=run(s,p,store,transport=transport,config=CONFIG);again=run(s,p,store,transport=transport,config=CONFIG)
  self.assertEqual(len(calls),1);self.assertTrue(store['cache_hit']);self.assertEqual(again,result);self.assertNotIn('input_text',result['receipt']);self.assertEqual(result['receipt']['usage']['total_tokens'],200)
  proposal=merge(s,p,result);self.assertGreater(proposal['structure_ai']['accepted'],2);self.assertIsNone(s.confirmed);self.assertIsNone(s.report)
  local=suggest(s);self.assertLess(sum(v not in ('other','pending') for v in local['draft']['paragraph_types'].values()),sum(v not in ('other','pending') for v in proposal['draft']['paragraph_types'].values()))
  plan=accept_scope(s,proposal)
  ref=self.session('reference');ref.set_types(proposal['draft']['paragraph_types']);ref.confirm_and_check();ref.set_selection(plan['keys']);self.assertTrue(ref.generate(),ref.error)
  execute(s,plan,{});self.assertEqual(s.downloads()['fixed.docx'],ref.downloads()['fixed.docx'])
 def test_bad_evidence_isolated_and_role_conflicts_review(self):
  s=self.session();p=prepare(s);raw=json.loads(json.loads(self.response(p))['choices'][0]['message']['content'])
  raw['suggestions'][0]['evidence']='概括拼接不在输入中';raw['suggestions'][-1]['role']='heading1'
  rows=validate(raw,p['request']);self.assertFalse(rows[0]['accepted']);self.assertFalse(rows[-1]['accepted']);self.assertTrue(any(r['accepted'] for r in rows))
 def test_protocol_json_truncation_tool_calls_and_no_retry(self):
  s=self.session();p=prepare(s)
  bad=['not json','{"x":1,"x":2}',json.dumps({'choices':[{'finish_reason':'length','message':{'content':'{}'}}]}),json.dumps({'choices':[{'finish_reason':'stop','message':{'content':'{}','tool_calls':[{}]}}]})]
  for text in bad:
   calls=[]
   def transport(*args):calls.append(1);return text
   with self.assertRaises(DemoError):run(s,p,{},transport=transport,config=CONFIG)
   self.assertEqual(calls,[1]);self.assertIsNone(s.report)
 def test_unknown_extra_ids_roles_and_fields_rejected(self):
  s=self.session();p=prepare(s);raw=json.loads(json.loads(self.response(p))['choices'][0]['message']['content'])
  for mutate in [lambda d:d['suggestions'].append(deepcopy(d['suggestions'][0])),lambda d:d['suggestions'][0].update(id='p:999'),lambda d:d['suggestions'][0].update(role='arbitrary'),lambda d:d['suggestions'][0].update(confirmed=True)]:
   d=deepcopy(raw);mutate(d)
   with self.assertRaises(DemoError):validate(d,p['request'])
 def test_inflight_file_rules_mode_and_manual_draft_changes_rejected(self):
  for mutation in [lambda s:s.set_types({**s.paragraph_types,'9':'other'}),lambda s:s.set_mode('basic'),lambda s:setattr(s,'rules_hash','stale'),lambda s:s.input_path.write_bytes(b'changed')]:
   s=self.session(str(id(mutation)));p=prepare(s)
   def transport(*args):mutation(s);return self.response(p)
   with self.assertRaises(DemoError):run(s,p,{},transport=transport,config=CONFIG)
   self.assertIsNone(s.report)
 def test_cross_session_old_seal_and_restored_authorization(self):
  a=self.session('a');b=self.session('b');p=prepare(a);r=run(a,p,{},transport=lambda *_:self.response(p),config=CONFIG)
  with self.assertRaises(DemoError):merge(b,prepare(b),r)
  proposal=merge(a,p,r);accept_scope(a,proposal);raw=a.prepare_project();b.stage_project(raw);restored=b.restore_project(acknowledged=True)
  self.assertIsNotNone(restored.report);self.assertFalse(hasattr(restored,'structure_ai'));self.assertIsNone(restored.rehearsal)
  self.assertNotIn(b'input_text',raw);self.assertEqual(prepare(restored)['request']['candidates'],[])
 def test_budgets_basic_extended_and_zero_candidate_no_post(self):
  from docx import Document
  from io import BytesIO
  d=Document()
  for i in range(70):d.add_paragraph(f'未知短段落{i}')
  buf=BytesIO();d.save(buf);s=WebSession(self.root/'large');s.load(buf.getvalue(),'合成.docx');p=prepare(s)
  self.assertEqual(len(p['request']['candidates']),MAX_CANDIDATES);self.assertEqual(p['remaining'],30);self.assertLess(len(json.dumps(p['request'],ensure_ascii=False)),MAX_INPUT)
  self.assertEqual(p['request']['allowed_roles'],['body','other','pending'])
  s.set_body([1]);p=prepare(s)
  with self.assertRaises(DemoError):run(s,p,{},transport=lambda *_:self.fail('must not POST'),config=CONFIG)
  e=self.session('extended','extended');self.assertNotIn('abstract_cn_body',prepare(e)['request']['allowed_roles'])
 def test_configuration_failure_and_network_failure_preserve_local(self):
  s=self.session();p=prepare(s);old=deepcopy(s.__dict__)
  with self.assertRaises(DemoError):run(s,p,{},transport=lambda *_:(_ for _ in ()).throw(TimeoutError()),config=CONFIG)
  self.assertEqual(s.paragraph_types,old['paragraph_types']);self.assertIsNone(s.report);self.assertTrue(suggest(s)['preflight']['ready'])

 def test_inherited_special_text_is_withheld_in_candidates_and_context(self):
  from docx import Document
  from docx.enum.style import WD_STYLE_TYPE
  from io import BytesIO
  d=Document();style=d.styles.add_style('SyntheticHidden',WD_STYLE_TYPE.CHARACTER);style.font.hidden=True
  d.add_paragraph('普通短文本');r=d.add_paragraph().add_run('hidden-personal-content');r.style=style;d.add_paragraph('另一个普通短文本')
  buf=BytesIO();d.save(buf);s=WebSession(self.root/'hidden');s.load(buf.getvalue(),'合成.docx');p=prepare(s)
  self.assertNotIn('p:2',{c['id'] for c in p['request']['candidates']});self.assertNotIn('hidden-personal-content',json.dumps(p['request']))
