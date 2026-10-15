"""第八阶段最小合成主题/继承/语言/编号验收，无真实论文内容。"""
from pathlib import Path
from tempfile import TemporaryDirectory
from copy import deepcopy
import json,unittest
from docx import Document
from docx.shared import Pt
from lxml import etree as E
from core.package import Package,q,NS,sha256,write_new
from core.engine import _ensure,check_document,fix_document
from core.formats import Resolver,Unresolved
from core.numbering import numbering_info
from web.presentation import finding_table
A='http://schemas.openxmlformats.org/drawingml/2006/main'
def aq(t):return '{'+A+'}'+t

class ResolutionAcceptance(unittest.TestCase):
 def setUp(self):
  self.tmp=TemporaryDirectory();self.dir=Path(self.tmp.name);d=Document();d.add_paragraph('软件演示样本');p=d.add_paragraph('正文 ABC café');p.runs[0].font.size=Pt(12);d.save(self.dir/'base.docx');self.pkg=Package(self.dir/'base.docx')
  default=self.pkg.styles.find('w:docDefaults/w:rPrDefault/w:rPr',NS);rf=_ensure(default,'rFonts');rf.attrib.clear()
  for slot,font in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('eastAsia','宋体')]:rf.set(q(slot),font)
  self.settings=E.Element(q('settings'));E.SubElement(self.settings,q('themeFontLang'),{q('val'):'en-US',q('eastAsia'):'zh-CN'})
  self.theme=E.Element(aq('theme'));fs=E.SubElement(E.SubElement(self.theme,aq('themeElements')),aq('fontScheme'))
  for group,latin,cn in [('major','Arial','黑体'),('minor','Times New Roman','宋体')]:
   coll=E.SubElement(fs,aq(group+'Font'));E.SubElement(coll,aq('latin'),typeface=latin);E.SubElement(coll,aq('ea'),typeface='');E.SubElement(coll,aq('font'),script='Hans',typeface=cn);E.SubElement(coll,aq('font'),script='Hant',typeface='新細明體')
 def tearDown(self):self.tmp.cleanup()
 def runprops(self):return _ensure(self.pkg.paragraphs[1].find(q('r')),'rPr')
 def fonts(self):return _ensure(self.runprops(),'rFonts')
 def write(self,name='case.docx'):
  self.pkg.parts['word/styles.xml']=E.tostring(self.pkg.styles);self.pkg.parts['word/settings.xml']=E.tostring(self.settings);self.pkg.parts['word/theme/theme1.xml']=E.tostring(self.theme);p=self.dir/name;write_new(p,self.pkg.bytes(True));return p
 def info(self,slot='ascii'):
  res=Resolver(self.pkg);return res.font_info(res.layers(self.pkg.paragraphs[1],self.pkg.paragraphs[1].find(q('r'))),slot)
 def load(self):p=self.write();self.pkg=Package(p);return p
 def bodycheck(self,p):return check_document(p,[2])
 def numbering(self):
  num=E.Element(q('numbering'));ab=E.SubElement(num,q('abstractNum'),{q('abstractNumId'):'40'});E.SubElement(ab,q('multiLevelType'),{q('val'):'multilevel'});lvl=E.SubElement(ab,q('lvl'),{q('ilvl'):'0'});E.SubElement(lvl,q('numFmt'),{q('val'):'decimal'});E.SubElement(lvl,q('lvlText'),{q('val'):'%1.'});rp=E.SubElement(lvl,q('rPr'));E.SubElement(rp,q('rFonts'),{q('ascii'):'Times New Roman'});E.SubElement(rp,q('sz'),{q('val'):'36'});instance=E.SubElement(num,q('num'),{q('numId'):'50'});E.SubElement(instance,q('abstractNumId'),{q('val'):'40'});self.pkg.parts['word/numbering.xml']=E.tostring(num)
  pp=_ensure(self.pkg.paragraphs[1],'pPr');np=_ensure(pp,'numPr');_ensure(np,'numId').set(q('val'),'50');_ensure(np,'ilvl').set(q('val'),'0');return num,ab,lvl,instance
 def headingcheck(self,p):
  m={'input_sha256':sha256(p.read_bytes()),'paragraph_types':{'1':'other','2':'heading1'},'confirmed':True,'confirmation':'人工合成范围'};return check_document(p,mode='extended',classification_manifest=m),m
 def find(self,r,prop):return next(f for f in r['findings'] if f['locator'].get('property')==prop)
 def test_major_ascii_and_minor_hansi(self):
  self.fonts().set(q('asciiTheme'),'majorAscii');self.fonts().set(q('hAnsiTheme'),'minorHAnsi');self.load();self.assertEqual(self.info()['font'],'Arial');self.assertEqual(self.info('hAnsi')['font'],'Times New Roman');self.assertEqual(self.info()['theme_basis']['theme_font_lang'],'en-US')
 def test_chinese_major_minor_distinct_from_western(self):
  self.fonts().set(q('eastAsiaTheme'),'majorEastAsia');self.load();self.assertEqual(self.info('eastAsia')['font'],'黑体');self.assertEqual(self.info()['font'],'Times New Roman')
 def test_traditional_theme_language(self):
  self.settings.find(q('themeFontLang')).set(q('eastAsia'),'zh-TW');self.fonts().set(q('eastAsiaTheme'),'minorEastAsia');self.load();self.assertEqual(self.info('eastAsia')['font'],'新細明體')
 def test_settings_language_not_overridden_by_proofing_language(self):
  _ensure(self.runprops(),'lang').set(q('eastAsia'),'zh-TW');self.fonts().set(q('eastAsiaTheme'),'minorEastAsia');self.load();self.assertEqual(self.info('eastAsia')['font'],'宋体');self.assertEqual(self.info('eastAsia')['run_language']['eastAsia'],'zh-TW')
 def test_default_region_without_theme_language(self):
  self.settings.remove(self.settings.find(q('themeFontLang')));self.theme.find('.//'+aq('minorFont')+'/'+aq('ea')).set('typeface','楷体');self.fonts().set(q('eastAsiaTheme'),'minorEastAsia');self.load();self.assertEqual(self.info('eastAsia')['font'],'楷体')
 def test_missing_region_does_not_infer_chinese_from_run(self):
  self.settings.remove(self.settings.find(q('themeFontLang')));_ensure(self.runprops(),'lang').set(q('eastAsia'),'zh-CN');self.fonts().set(q('eastAsiaTheme'),'minorEastAsia');self.load();self.assertIsNone(self.info('eastAsia')['font'])
 def test_direct_literal_overrides_lower_theme(self):
  default=self.pkg.styles.find('w:docDefaults/w:rPrDefault/w:rPr/w:rFonts',NS);default.set(q('asciiTheme'),'majorAscii');self.fonts().set(q('ascii'),'Calibri');self.load();self.assertEqual(self.info()['font'],'Calibri');self.assertEqual(self.info()['declared_at'],'直接run格式')
 def test_same_layer_theme_over_literal(self):
  rf=self.fonts();rf.set(q('ascii'),'Times New Roman');rf.set(q('asciiTheme'),'majorAscii');self.load();self.assertEqual(self.info()['font'],'Arial');self.assertEqual(self.info()['source_type'],'theme')
 def test_character_style_basedon_and_paragraph_order(self):
  normal=self.pkg.styles.find("w:style[@w:styleId='Normal']",NS);_ensure(_ensure(normal,'rPr'),'rFonts').set(q('ascii'),'Arial')
  for sid,base,font in [('BaseChar',None,'Calibri'),('ChildChar','BaseChar',None)]:
   st=E.SubElement(self.pkg.styles,q('style'),{q('styleId'):sid,q('type'):'character'})
   if base:E.SubElement(st,q('basedOn'),{q('val'):base})
   if font:_ensure(_ensure(st,'rPr'),'rFonts').set(q('ascii'),font)
  _ensure(self.runprops(),'rStyle').set(q('val'),'ChildChar');self.load();self.assertEqual(self.info()['font'],'Calibri');self.assertIn('character样式:BaseChar',self.info()['inheritance_path']);self.assertIn('paragraph样式:Normal',self.info()['inheritance_path'])
 def test_cycles_duplicate_style_and_default_rejected(self):
  normal=self.pkg.styles.find("w:style[@w:styleId='Normal']",NS);_ensure(normal,'basedOn').set(q('val'),'Normal');self.load()
  with self.assertRaises(Unresolved):self.info()
 def test_duplicate_default_styles_not_guessed(self):
  st=E.SubElement(self.pkg.styles,q('style'),{q('styleId'):'OtherDefault',q('type'):'paragraph',q('default'):'1'});self.load()
  with self.assertRaises(Unresolved):self.info()
 def test_missing_font_never_default_pass(self):
  rf=self.pkg.styles.find('w:docDefaults/w:rPrDefault/w:rPr/w:rFonts',NS);rf.attrib.clear();p=self.load();f=next(x for x in self.bodycheck(p)['findings'] if x['rule_id']=='R04');self.assertEqual(f['status'],'review');self.assertIn('不补默认值',f['detail'])
 def test_missing_theme_relationship(self):
  self.fonts().set(q('asciiTheme'),'majorAscii');from core.package import parse
  rel=parse(self.pkg.parts['word/_rels/document.xml.rels']);[rel.remove(e) for e in list(rel) if e.get('Type','').endswith('/theme')];self.pkg.parts['word/_rels/document.xml.rels']=E.tostring(rel);self.load();self.assertIn('主题关系缺失',self.info()['reason'])
 def test_malformed_theme_review_not_crash(self):
  self.fonts().set(q('asciiTheme'),'majorAscii');p=self.write();pkg=Package(p);pkg.parts['word/theme/theme1.xml']=b'<broken';bad=self.dir/'bad.docx';write_new(bad,pkg.bytes(True));self.pkg=Package(bad);self.assertIn('XML损坏',self.info()['reason'])
 def test_duplicate_theme_script_review(self):
  E.SubElement(self.theme.find('.//'+aq('minorFont')),aq('font'),script='Hans',typeface='黑体');self.fonts().set(q('eastAsiaTheme'),'minorEastAsia');self.load();self.assertIsNone(self.info('eastAsia')['font'])
 def test_unknown_theme_language_and_cross_slot(self):
  self.settings.find(q('themeFontLang')).set(q('val'),'ru-RU');self.fonts().set(q('asciiTheme'),'minorAscii');self.load();self.assertIn('主题语言',self.info()['reason']);self.assertFalse(self.info()['safe_to_fix_slot'])
 def test_ascii_eastasiahint_safe_latin_extension_review(self):
  self.fonts().set(q('hint'),'eastAsia');self.fonts().set(q('ascii'),'Arial');self.load();self.assertEqual(self.info()['font'],'Arial');self.assertIsNone(self.info('hAnsi')['font']);self.assertIn('Latin扩展',self.info('hAnsi')['reason'])
 def test_unknown_hint_review(self):
  self.fonts().set(q('hint'),'invalid');self.load();self.assertIsNone(self.info()['font']);self.assertIn('未支持',self.info()['reason'])
 def test_cjk_timesnewroman_fallback_readonly(self):
  rf=self.fonts()
  for slot,font in [('eastAsia','Times New Roman'),('ascii','Arial'),('hAnsi','Arial')]:rf.set(q(slot),font)
  self.load();i=self.info('eastAsia');self.assertEqual(i['font'],'Arial');self.assertFalse(i['safe_to_fix_slot']);self.assertIn('回退',i['fallback_basis'])
 def test_declared_theme_font_does_not_check_computer_installation(self):
  self.theme.find('.//'+aq('minorFont')+'/'+aq('latin')).set('typeface','Synthetic Uninstalled Font');self.fonts().set(q('asciiTheme'),'minorAscii');p=self.load();f=next(x for x in self.bodycheck(p)['findings'] if x['rule_id']=='R04');self.assertEqual(f['status'],'problem');self.assertFalse(f['font_evidence']['ascii']['computer_font_installation_checked'])
 def test_theme_fix_selected_slot_preserves_chinese_and_package(self):
  self.fonts().set(q('asciiTheme'),'majorAscii');p=self.load();old=Package(p);r=self.bodycheck(p);f=next(x for x in r['findings'] if x['rule_id']=='R04');out=fix_document(p,self.dir/'out.docx',[2],selected_keys=[f['key']]);new=Package(out['output']);self.assertEqual(Resolver(new).font(Resolver(new).layers(new.paragraphs[1],new.paragraphs[1].find(q('r'))),'eastAsia'),'宋体');self.assertEqual(new.parts['word/theme/theme1.xml'],old.parts['word/theme/theme1.xml']);self.assertEqual(new.parts['word/numbering.xml'],old.parts['word/numbering.xml']);self.assertTrue(out['preservation']['original_sha256_unchanged']);self.assertEqual(next(x for x in out['after']['findings'] if x['key']==f['key'])['status'],'pass')
 def test_numbered_heading_text_checks_not_whole_skipped(self):
  self.numbering();p=self.load();r,m=self.headingcheck(p);self.assertEqual(self.find(r,'size')['status'],'problem');self.assertEqual(self.find(r,'font_eastAsia')['status'],'problem');self.assertEqual(self.find(r,'alignment')['status'],'problem');self.assertEqual(self.find(r,'line')['status'],'problem');self.assertEqual(self.find(r,'numbering_symbol')['status'],'pass')
 def test_real_heading_style_verified_even_with_numbering(self):
  self.numbering();_ensure(_ensure(self.pkg.paragraphs[1],'pPr'),'pStyle').set(q('val'),'Heading1');p=self.load();r,m=self.headingcheck(p);f=next(x for x in r['findings'] if x['rule_id']=='R11');self.assertEqual(f['status'],'pass')
 def test_missing_symbol_font_not_borrowed_from_text(self):
  num,ab,lvl,instance=self.numbering();lvl.remove(lvl.find(q('rPr')));self.pkg.parts['word/numbering.xml']=E.tostring(num);p=self.load();r,m=self.headingcheck(p);self.assertEqual(self.find(r,'numbering_symbol')['status'],'review');self.assertFalse(self.find(r,'numbering_symbol')['fixable']);self.assertEqual(self.find(r,'size')['status'],'problem')
 def test_bad_symbol_font_problem_never_fixable(self):
  num,ab,lvl,instance=self.numbering();lvl.find('w:rPr/w:rFonts',NS).set(q('ascii'),'Arial');self.pkg.parts['word/numbering.xml']=E.tostring(num);p=self.load();r,m=self.headingcheck(p);self.assertEqual(self.find(r,'numbering_symbol')['status'],'problem');self.assertFalse(self.find(r,'numbering_symbol')['fixable'])
 def test_unknown_style_link_keeps_text_checks_and_review_paragraph(self):
  num,ab,lvl,instance=self.numbering();E.SubElement(ab,q('numStyleLink'),{q('val'):'Unknown'});self.pkg.parts['word/numbering.xml']=E.tostring(num);p=self.load();r,m=self.headingcheck(p);self.assertEqual(self.find(r,'numbering_symbol')['status'],'review');self.assertEqual(self.find(r,'size')['status'],'problem');self.assertEqual(self.find(r,'alignment')['status'],'review')
 def test_level_override_and_start_readonly(self):
  num,ab,lvl,instance=self.numbering();ov=E.SubElement(instance,q('lvlOverride'),{q('ilvl'):'0'});E.SubElement(ov,q('startOverride'),{q('val'):'4'});rep=deepcopy(lvl);rep.find('w:rPr/w:rFonts',NS).set(q('ascii'),'Arial');ov.append(rep);self.pkg.parts['word/numbering.xml']=E.tostring(num);p=self.load();r,m=self.headingcheck(p);f=self.find(r,'numbering_symbol');self.assertEqual(f['actual']['font'],'Arial');self.assertEqual(f['numbering_evidence']['start_override'],'4');self.assertTrue(f['numbering_evidence']['level_overridden'])
 def test_duplicate_overrides_and_picture_bullet_review(self):
  num,ab,lvl,instance=self.numbering();E.SubElement(lvl,q('lvlPicBulletId'),{q('val'):'3'});self.pkg.parts['word/numbering.xml']=E.tostring(num);p=self.load();r,m=self.headingcheck(p);self.assertIn('图片编号',self.find(r,'numbering_symbol')['detail'])
 def test_numbering_conflicting_paragraph_attrs_not_assumed(self):
  num,ab,lvl,instance=self.numbering();E.SubElement(E.SubElement(lvl,q('pPr')),q('jc'),{q('val'):'center'});self.pkg.parts['word/numbering.xml']=E.tostring(num);p=self.load();r,m=self.headingcheck(p);self.assertEqual(self.find(r,'alignment')['status'],'review');self.assertEqual(self.find(r,'line')['status'],'review');self.assertEqual(self.find(r,'size')['status'],'problem')
 def test_R11_mismatch_problem_but_unsafe_correction_refused(self):
  self.numbering();p=self.load();r,m=self.headingcheck(p);f=next(x for x in r['findings'] if x['rule_id']=='R11');self.assertEqual(f['status'],'problem');self.assertFalse(f['fixable']);out=fix_document(p,self.dir/'out.docx',None,mode='extended',classification_manifest=m,selected_keys=[f['key']]);self.assertEqual(out['changes'],[]);self.assertEqual(Path(out['output']).read_bytes(),p.read_bytes())
 def test_numbered_title_size_fix_does_not_touch_numbering_or_style(self):
  self.numbering();p=self.load();r,m=self.headingcheck(p);f=self.find(r,'size');raw=p.read_bytes();old=Package(p);out=fix_document(p,self.dir/'out.docx',None,mode='extended',classification_manifest=m,selected_keys=[f['key']]);new=Package(out['output']);self.assertEqual(new.parts['word/numbering.xml'],old.parts['word/numbering.xml']);self.assertEqual(new.parts['word/styles.xml'],old.parts['word/styles.xml']);self.assertEqual(p.read_bytes(),raw);self.assertEqual(self.find(out['after'],'size')['status'],'pass');self.assertTrue(out['preservation']['document_except_selected_format_attributes_equal'])
 def test_font_and_numbering_reasons_visible_and_escaped(self):
  self.fonts().set(q('asciiTheme'),'invalid');p=self.load();r=self.bodycheck(p);inventory={'paragraphs':[{'text':'软件演示'},{'text':'<script>合成</script>'}]};html=finding_table(r['findings'],inventory);self.assertIn('主题枚举',html);self.assertIn('&lt;script&gt;',html);self.assertNotIn('<script>',html);self.assertIn('其余可判定属性',html)
 def test_cross_slot_theme_reference_not_guessed(self):
  self.fonts().set(q('asciiTheme'),'majorEastAsia');self.load();self.assertIsNone(self.info()['font']);self.assertIn('槽位不对应',self.info()['reason'])
 def test_duplicate_theme_language_review(self):
  E.SubElement(self.settings,q('themeFontLang'),{q('val'):'fr-FR'});self.fonts().set(q('asciiTheme'),'minorAscii');self.load();self.assertIn('重复themeFontLang',self.info()['reason'])
 def test_duplicate_overrides_not_guessed(self):
  num,ab,lvl,instance=self.numbering()
  for _ in range(2):E.SubElement(instance,q('lvlOverride'),{q('ilvl'):'0'})
  self.pkg.parts['word/numbering.xml']=E.tostring(num);p=self.load();r,m=self.headingcheck(p);self.assertIn('重复编号层级覆盖',self.find(r,'numbering_symbol')['detail']);self.assertEqual(self.find(r,'size')['status'],'problem')
 def test_malformed_numbering_preserves_independent_text_check(self):
  self.numbering();self.pkg.parts['word/numbering.xml']=b'<broken';p=self.load();r,m=self.headingcheck(p);self.assertEqual(self.find(r,'numbering_symbol')['status'],'review');self.assertEqual(self.find(r,'size')['status'],'problem');self.assertEqual(self.find(r,'line')['status'],'review')
 def test_complex_numbering_symbol_not_passed_from_ascii_font(self):
  num,ab,lvl,instance=self.numbering();E.SubElement(lvl.find(q('rPr')),q('cs'));self.pkg.parts['word/numbering.xml']=E.tostring(num);p=self.load();r,m=self.headingcheck(p);self.assertEqual(self.find(r,'numbering_symbol')['status'],'review');self.assertIn('复杂文字',self.find(r,'numbering_symbol')['detail'])
 def test_duplicate_character_style_reference_conflict(self):
  rp=self.runprops()
  for _ in range(2):E.SubElement(rp,q('rStyle'),{q('val'):'DefaultParagraphFont'})
  self.load()
  with self.assertRaisesRegex(Unresolved,'重复字符样式'):self.info()
 def test_numbering_symbol_character_style_dependency_review(self):
  num,ab,lvl,instance=self.numbering();E.SubElement(lvl.find(q('rPr')),q('rStyle'),{q('val'):'Unknown'});self.pkg.parts['word/numbering.xml']=E.tostring(num);p=self.load();r,m=self.headingcheck(p);self.assertEqual(self.find(r,'numbering_symbol')['status'],'review');self.assertIn('字符样式依赖',self.find(r,'numbering_symbol')['detail'])
 def test_localized_chinese_font_name_comparison_keeps_slots(self):
  self.theme.find('.//'+aq('minorFont')+'/'+aq('font')+'[@script="Hans"]').set('typeface','SimSun');self.fonts().set(q('eastAsiaTheme'),'minorEastAsia');self.fonts().set(q('ascii'),'Arial');p=self.load();r=self.bodycheck(p);cn=next(f for f in r['findings'] if f['rule_id']=='R03');western=next(f for f in r['findings'] if f['rule_id']=='R04');self.assertEqual(cn['status'],'pass');self.assertEqual(cn['font_evidence']['eastAsia']['font'],'SimSun');self.assertEqual(western['status'],'problem')
 def test_malformed_numbering_classification_does_not_abort_loading(self):
  from core.classification import suggest_paragraph_types
  from web.session import WebSession
  self.numbering();self.pkg.parts['word/numbering.xml']=b'<broken';p=self.load();self.assertEqual(len(suggest_paragraph_types(p)['paragraphs']),2);s=WebSession(self.dir/'sessions');s.set_mode('extended');self.assertTrue(s.load(p.read_bytes(),'合成.docx'));s.set_types({'1':'other','2':'heading1'});s.confirm();self.assertTrue(s.check());self.assertEqual(self.find(s.report,'size')['status'],'problem');self.assertEqual(self.find(s.report,'numbering_symbol')['status'],'review')
 def test_known_ascii_error_in_mixed_pending_run_can_be_safely_fixed(self):
  self.fonts().set(q('hint'),'eastAsia');self.fonts().set(q('ascii'),'Arial');p=self.load();r=self.bodycheck(p);pending=next(f for f in r['findings'] if f['key']=='R04:paragraph=2:run=1');known=next(f for f in r['findings'] if f['key']=='R04:paragraph=2:run=1:property=font_ascii');self.assertEqual(pending['status'],'review');self.assertEqual(known['status'],'problem');self.assertTrue(known['fixable']);out=fix_document(p,self.dir/'out.docx',[2],selected_keys=[known['key']]);self.assertEqual(next(f for f in out['after']['findings'] if f['key']==known['key'])['status'],'pass');self.assertEqual(next(f for f in out['after']['findings'] if f['key']==pending['key'])['status'],'review');self.assertEqual(out['changes'][0]['after'],{'ascii':'Times New Roman'})
 def test_western_fix_does_not_change_unselected_cjk_fallback(self):
  rf=self.fonts()
  for slot,font in [('eastAsia','Times New Roman'),('ascii','Arial'),('hAnsi','Arial')]:rf.set(q(slot),font)
  p=self.load();r=self.bodycheck(p);f=next(f for f in r['findings'] if f['rule_id']=='R04');self.assertEqual(f['status'],'problem');self.assertFalse(f['fixable']);out=fix_document(p,self.dir/'out.docx',[2],selected_keys=[f['key']]);self.assertEqual(out['changes'],[]);self.assertEqual(Path(out['output']).read_bytes(),p.read_bytes())
 def test_latin_primary_language_with_unknown_script_not_assumed_latin(self):
  self.settings.find(q('themeFontLang')).set(q('val'),'en-Arab');self.fonts().set(q('asciiTheme'),'minorAscii');self.load();self.assertIsNone(self.info()['font']);self.assertIn('主题语言',self.info()['reason'])
