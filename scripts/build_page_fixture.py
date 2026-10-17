"""Author-marked public multi-page fixture; expected visible numbers are independent."""
from pathlib import Path
import sys,json
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from docx import Document
from docx.shared import Pt
from lxml import etree as E
from core.package import q,NS,Package
from copy import deepcopy

def build(out):
 out=Path(out);out.mkdir(parents=True,exist_ok=True)
 d=Document();d.sections[0].header.paragraphs[0].text='公开合成页眉 — 保持'
 f=d.sections[0].footer;f.paragraphs[0].text='原有页脚说明（保持）';p=f.add_paragraph();field=E.SubElement(p._p,q('fldSimple'));field.set(q('instr'),' PAGE ')
 # Existing first/even variants are active and must remain active after processing.
 d.sections[0].different_first_page_header_footer=True;d.settings.odd_and_even_pages_header_footer=True
 for footer in (d.sections[0].first_page_footer,d.sections[0].even_page_footer):
  footer.paragraphs[0].text='原有变体说明（保持）';E.SubElement(footer.add_paragraph()._p,q('fldSimple'),{q('instr'):' PAGE '})
 definitions=[('other','封面 — 公开合成论文'),('other','封面第二页 — 不应显示页码'),('abstract_cn_title','摘要'),('other','目录 — 前置第二页'),('heading1','第1章 合成研究'),('body','依据图1-2和图2-3核对纯文本引用。'),('figure_caption','图1-2 合成曲线'),('body','第一章第二页 — 正文页码应连续。'),('heading1','第2章 合成结果'),('body','第二章第一页 — 不应重新从1编号。'),('figure_caption','图2-3 合成结果曲线'),('body','第二章第二页 — 正文末页。')]
 for i,(role,tx) in enumerate(definitions,1):
  p=d.add_paragraph(tx,style='Heading 1' if role=='heading1' else None)
  if i in (2,3,4,5,8,9,12):p.paragraph_format.page_break_before=True
  if role=='body':p.paragraph_format.line_spacing=1;p.runs[0].font.size=Pt(14)
 d.save(out/'single.docx');pkg=Package(out/'single.docx');sect=deepcopy(pkg.sections[-1]);ps=pkg.paragraphs
 for start in (3,5,9):
  pp=ps[start-2].find(q('pPr'))
  if pp is None:pp=E.Element(q('pPr'));ps[start-2].insert(0,pp)
  pp.append(deepcopy(sect))
 (out/'existing.docx').write_bytes(pkg.bytes(True))
 roles={str(i):r for i,(r,_) in enumerate(definitions,1)}
 (out/'roles.json').write_text(json.dumps(roles,ensure_ascii=False,indent=2)+'\n')
 (out/'expected.json').write_text(json.dumps({'kind':'author labelled before rendering/model response','physical_pages':8,'visible_numbers':[None,None,'i','ii','1','2','3','4'],'chapter_footers':{'1':'第1章 合成研究','2':'第2章 合成结果'},'digits':{'6':'依据图1-1和图2-1核对纯文本引用。','7':'图1-1 合成曲线','11':'图2-1 合成结果曲线'},'unchanged':['all nonnumeric body text','run properties during structure transaction','header parts','note and numbering settings','original footer parts','other relationships']},ensure_ascii=False,indent=2)+'\n')
 return out/'existing.docx',roles
if __name__=='__main__':
 import argparse
 p=argparse.ArgumentParser();p.add_argument('--output',required=True);a=p.parse_args();build(a.output)
