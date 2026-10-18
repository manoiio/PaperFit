"""Small wholly synthetic three-level TOC, independent of local demonstration papers."""
from pathlib import Path
import sys,json
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
def build(folder):
 from docx import Document
 from docx.shared import Pt,Cm
 from docx.oxml import OxmlElement
 from docx.oxml.ns import qn
 out=Path(folder);out.mkdir(parents=True,exist_ok=True)
 d=Document();d.core_properties.author='PaperFit public synthetic fixture';sec=d.sections[0];sec.left_margin=Cm(2.6);sec.right_margin=Cm(3.1)
 normal=d.styles['Normal'];normal.font.name='宋体';normal.font.size=Pt(12);normal.paragraph_format.line_spacing=1.5
 d.add_paragraph('封面 — 公开合成目录外观验证','Title');d.add_paragraph('摘要').paragraph_format.page_break_before=True;d.add_paragraph('本文只使用公开编写的短句，验证目录点状引导线、宽度与层级。');d.add_paragraph('关键词：合成；目录')
 d.add_paragraph('目录');p=d.add_paragraph();f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),' TOC \\o "1-3" \\h \\z \\u ');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='旧目录与旧页码99';r.append(t);f.append(r);p._p.append(f)
 d.add_paragraph('')
 headings=[('第1章 目录外观验证',1),('1.1 长标题验证普通标题换行后页码仍然位于同一个右侧边界并且不依靠连续句点或空格模拟引导线以及保持内部链接可用',2),('1.1.1 参数与条件',3),('第2章 核对结果',1)]
 for title,level in headings:
  par=d.add_paragraph(title,'Heading '+str(level))
  if level==1:par.paragraph_format.page_break_before=True
  for _ in range(3):d.add_paragraph('这是独立编写的公开合成正文。目录外观只影响获授权目录缓存，正文文字、编号和其他部件保持。')
 path=out/'source.docx';d.save(path)
 (out/'expected.json').write_text(json.dumps({'headings':[{'text':t,'level':l} for t,l in headings],'level_left_twips':[0,360,720],'right_stop':'section page width minus left/right margins and gutter; subtract paragraph right indent','leader':'dot','contents_must_not_change':True,'school_default':False},ensure_ascii=False,indent=2)+'\n');return path
if __name__=='__main__':print(build(sys.argv[1]))
