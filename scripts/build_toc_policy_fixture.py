"""Independent small public three-level/long-title fixture; no local paper data."""
from pathlib import Path
import json

def build(folder):
 from docx import Document
 from docx.shared import Pt
 from docx.enum.section import WD_SECTION_START
 from docx.oxml import OxmlElement
 from docx.oxml.ns import qn
 out=Path(folder);out.mkdir(parents=True,exist_ok=True)
 d=Document();d.core_properties.author='PaperFit public synthetic fixture'
 normal=d.styles['Normal'];normal.font.name='宋体';normal.font.size=Pt(12)
 def p(text,style=None,newpage=False):
  r=d.add_paragraph(text,style)
  if newpage:r.paragraph_format.page_break_before=True
  return r
 p('摘要');p('这是独立编写的公开合成摘要，用于核对目录收录，不是真实论文。')
 p('Abstract',newpage=True);p('A public synthetic abstract for bounded table of contents verification.')
 p('目录',newpage=True);r=p('');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),' TOC \\o "1-2" \\h \\z \\u ');rr=OxmlElement('w:r');t=OxmlElement('w:t');t.text='过时条目\t99';rr.append(t);fld.append(rr);r._p.append(fld)
 # Deliberately stale direct and inherited geometry must not win over reference mode.
 r.paragraph_format.line_spacing=1;r.paragraph_format.left_indent=Pt(3);r.paragraph_format.space_after=Pt(8)
 d.add_section(WD_SECTION_START.NEW_PAGE)
 p('第1章 合成方法','Heading 1');p('正文内容保持，用于核对非目录文字与格式不变。')
 p('1.1 处理过程','Heading 2');p('说明内容保持。')
 p('1.1.1 '+('连续文本核对长标题的换行与页码保留区域' * 4),'Heading 3');p('三级标题后正文保持。')
 p('参考文献');p('[1] Public Synthetic Reference. 2026.')
 for i,s in enumerate(d.sections):
  s.page_width=__import__('docx.shared',fromlist=['Cm']).Cm(21)
  s.left_margin=s.right_margin=__import__('docx.shared',fromlist=['Cm']).Cm(2.8)
  x=OxmlElement('w:pgNumType');x.set(qn('w:fmt'),'lowerRoman' if i==0 else 'decimal');x.set(qn('w:start'),'1' if i==0 else '5');s._sectPr.append(x)
 path=out/'source.docx';d.save(path)
 roles={str(n):'other' for n in range(1,len(d.paragraphs)+1)}
 roles.update({'1':'abstract_cn_title','2':'abstract_cn_body','3':'abstract_en_title','4':'abstract_en_body','8':'heading1','9':'body','10':'heading2','11':'body','12':'heading3','13':'body','14':'reference_title','15':'reference_item'})
 (out/'roles.json').write_text(json.dumps(roles,ensure_ascii=False,indent=2)+'\n')
 (out/'expected.json').write_text(json.dumps({'target_paragraphs':[1,3,5,8,10,12,14],'target_levels':[1,1,1,1,2,3,1],'included_toc_title':5,'line_twips':474,'lineRule':'exact','left_twips':[0,480,960],'before':'0','after':'0','first_abstract_pages':['i','ii'],'body_start':'5'},ensure_ascii=False,indent=2)+'\n')
 return path
if __name__=='__main__':
 import sys
 print(build(sys.argv[1]))
