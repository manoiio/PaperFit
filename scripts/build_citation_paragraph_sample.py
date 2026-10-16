"""Independent paragraph-property fixture; no exception checker or document fixer."""
from pathlib import Path
import sys,json
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
ROOT=Path(__file__).resolve().parents[1]
def build(output):
 out=Path(output);out.mkdir(parents=True,exist_ok=True);d=Document(ROOT/'samples/phase30/errors.docx');targets=[]
 texts={3:[('已有普通正文',False,12),('[1]',True,9),('继续正文。',False,12)],4:[('列表正文',False,12),('[1,3]',True,8),('重复',False,12),('[1,3]',True,8),('跨run',False,12),('[1–',True,9),('3,5]',True,9),('结束。',False,12)],5:[('未确认上标正文',False,12),('[2]',True,9),('保留。',False,12)],6:[('额外上标x',False,12),('2',True,10),('和引文',False,12),('[1]',True,9)],7:[('[1]',True,9)]}
 roles={str(i):'reference_item' if 8<=i<=17 else 'body' if i in texts else 'other' for i in range(1,len(d.paragraphs)+1)}
 for n,parts in texts.items():
  p=d.paragraphs[n-1]
  for c in list(p._p):p._p.remove(c)
  p.paragraph_format.line_spacing=2;p.paragraph_format.space_before=Pt(12);p.paragraph_format.space_after=Pt(0);p.alignment=0
  pp=p._p.get_or_add_pPr();ind=pp.get_or_add_ind();ind.set(qn('w:firstLineChars'),'100' if n==3 else '200' if n in (5,6) else '0');ind.set(qn('w:firstLine'),'360' if n==4 else '240' if n==7 else '0')
  if n in (4,7):ind.attrib.pop(qn('w:firstLineChars'),None)
  pos=0
  for tx,super_,size in parts:
   r=p.add_run(tx);r.font.name='Times New Roman';r.font.size=Pt(size);r.font.superscript=super_;r._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
   pos+=len(tx)
  full=''.join(t for t,_,_ in parts)
  for raw in ('[1]','[1,3]','[1–3,5]','[2]'):
   start=0
   while raw in full[start:]:
    a=full.index(raw,start);targets.append({'paragraph':n,'start':a,'end':a+len(raw),'raw':raw,'purpose':'pending' if n==5 else 'citation'});start=a+len(raw)
 d.core_properties.author='软件演示';d.core_properties.last_modified_by='软件演示';d.paragraphs[0].runs[0].text='确认引文上标后的段落属性演示'
 d.save(out/'errors.docx')
 # Correct only the independently authorized safe pPr. Protected paragraphs5/6 unchanged.
 for n in (3,4,7):
  p=d.paragraphs[n-1];p.paragraph_format.line_spacing=1.5;p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=None;p.alignment=3
  spacing=p._p.find(qn('w:pPr')).find(qn('w:spacing'));spacing.set(qn('w:after'),'0');spacing.set(qn('w:afterLines'),'50');spacing.set(qn('w:afterAutospacing'),'0');spacing.set(qn('w:beforeAutospacing'),'0')
  if n!=7:
   ind=p._p.find(qn('w:pPr')).find(qn('w:ind'));ind.set(qn('w:firstLineChars'),'200');ind.set(qn('w:firstLine'),'0')
 d.save(out/'correct.docx')
 meta={'roles':roles,'scan':list(texts),'citations':targets,'complete':False,'references_complete':False,'material_exclusions':{'18':'人工排除：合成代码，保留空白与换行'},'expected':{'3':['R06','R07','R13.before','R13.after','R13.alignment'],'4':['R06','R07','R13.before','R13.after','R13.alignment'],'5':'引文用途pending，段落属性review','6':'引文区间外上标，段落属性review','7':['R06','R13.before','R13.after','R13.alignment'],'7_R07':'仅引文无可靠普通正文，点值缩进review'},'target_properties':{'line':360,'lineRule':'auto','firstLineChars':200,'firstLine':0,'before':120,'beforeAutospacing':0,'after':0,'afterLines':50,'afterAutospacing':0,'jc':'both'}}
 (out/'independent-errors.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2)+'\n')
if __name__=='__main__':build(ROOT/'samples/phase31')
