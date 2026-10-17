"""Small author-labelled structural ambiguities, no historical or personal input."""
from pathlib import Path
import sys,json
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))

def build(path):
    from docx import Document
    from docx.shared import Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    data=json.loads((ROOT/'samples/phase41/labels.json').read_text())
    d=Document()
    for row in data['paragraphs']:
        p=d.add_paragraph(style=row.get('style','Normal'));r=p.add_run(row['text']);r.font.size=Pt(14)
        fonts=OxmlElement('w:rFonts');fonts.set(qn('w:eastAsia'),'宋体');fonts.set(qn('w:ascii'),'Times New Roman');r._r.get_or_add_rPr().insert(0,fonts)
        p.paragraph_format.line_spacing=1.5
    d.core_properties.author='Synthetic acceptance';d.core_properties.title='Structure acceptance'
    d.save(path);return path
if __name__=='__main__':
    import argparse
    p=argparse.ArgumentParser();p.add_argument('--output',required=True);a=p.parse_args();build(Path(a.output))
