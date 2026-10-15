"""Synthetic known-dependency probe; no private input or external service."""
from pathlib import Path
from io import BytesIO
import sys,json,base64
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
from docx.shared import Mm
from core.package import Package,q,NS
from lxml import etree as E
ROOT=Path(__file__).resolve().parents[1]
def build(output):
 output=Path(output);output.mkdir(parents=True,exist_ok=True)
 d=Document(ROOT/'samples/phase19/errors.docx')
 d.paragraphs[0].runs[0].text='分步修正计划软件演示样本'
 code=d.add_paragraph('    result = values[0]\t# 人工排除的合成代码\n    print(result)')
 table=d.add_table(rows=2,cols=2);table.cell(0,0).text='保持表格';table.cell(0,1).text='值';table.cell(1,0).text='合成数据';table.cell(1,1).text='28'
 img=base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII=')
 d.add_paragraph('下方小图仅验证媒体保留，不作论文插图。');d.add_picture(BytesIO(img),width=Mm(8))
 temp=output/'author.docx';d.save(temp);pkg=Package(temp)
 def ensure(e,n):
  x=e.find(q(n))
  if x is None:
   x=E.Element(q(n));e.insert(0,x) if n in ('rPr','pPr') else e.append(x)
  return x
 for n in (3,4):
  run=pkg.paragraphs[n-1].findall(q('r'))[0];rp=ensure(run,'rPr');ensure(rp,'sz').set(q('val'),'28')
  fonts=ensure(rp,'rFonts');fonts.set(q('eastAsia'),'黑体');fonts.set(q('ascii'),'Arial');fonts.set(q('hAnsi'),'Arial')
 for run in pkg.paragraphs[2].findall(q('r')):ensure(ensure(run,'rPr'),'sz').set(q('val'),'28')
 # At14pt,28pt was two characters.12pt correction introduces an UNSELECTED indent risk.
 pp=ensure(pkg.paragraphs[2],'pPr');ind=ensure(pp,'ind');ind.attrib.clear();ind.set(q('firstLine'),'560')
 (output/'errors.docx').write_bytes(pkg.bytes(True));temp.unlink()
 roles={str(i):'body' if i in (3,4,5,6) else 'reference_item' if 8<=i<=17 else 'other' for i in range(1,len(pkg.paragraphs)+1)}
 roles['2']='heading1';roles['7']='reference_title'
 scope={'roles':roles,'scan':[3,4,5,6],'excluded_citation_text':'[8]','material_exclusions':{'18':'人工排除：合成程序代码，保留制表与换行'},'selected_ordinary':{'paragraphs':[3,4],'runs':[1,2,3],'rules':['R03','R04','R05']},'selected_citations':[{'paragraph':3,'start':7,'end':10},{'paragraph':3,'start':21,'end':24},{'paragraph':4,'start':12,'end':15},{'paragraph':5,'start':20,'end':23}]}
 (output/'scope.json').write_text(json.dumps(scope,ensure_ascii=False,indent=2)+'\n')
 (output/'independent-errors.json').write_text(json.dumps({'author':'合成注入清单，先于检查结果制定','ordinary':{'p3r1':['eastAsia黑体→宋体','14pt→12pt'],'p3r2r3':['14pt→12pt'],'p4r1':['eastAsia黑体→宋体','14pt→12pt']},'citations':scope['selected_citations'],'not_selected':['p3首行28pt（14pt时等效两字符；变12pt后不等效）','其他全部格式问题','复杂域只读引文','数组[8]','代码p18、图片、表格、分节、所有非主文档部件']},ensure_ascii=False,indent=2)+'\n')
if __name__=='__main__':build(ROOT/'samples/phase28')
