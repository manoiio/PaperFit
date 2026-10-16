"""Independently authored literal spans; no citation parser/checker/fixer."""
from pathlib import Path
from copy import deepcopy
import sys,json
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
def build(output):
 out=Path(output);out.mkdir(parents=True,exist_ok=True);d=Document(ROOT/'samples/phase28/errors.docx')
 d.paragraphs[0].runs[0].text='组合与区间引文 · 软件演示（无真实作者）'
 payloads={3:['组合[1,3]与重复[1,3]。'],4:['区间[','1–','3]和混合[1-3,5]结束。'],5:['已上标','[2,1]','与单引[4]。数组[8]。']}
 literals={3:['[1,3]','[1,3]'],4:['[1–3]','[1-3,5]'],5:['[2,1]','[4]']};targets=[]
 for n,texts in payloads.items():
  p=d.paragraphs[n-1]
  for child in list(p._p):
   if child.tag!=qn('w:pPr'):p._p.remove(child)
  for i,tx in enumerate(texts):
   r=p.add_run(tx);r.font.name='Times New Roman';r.font.size=Pt(14 if n in (3,4) else 12);r._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'黑体' if n in (3,4) else '宋体')
   if n==4 and i==1 or n==5 and i==1:r.font.superscript=True
  start=0;whole=''.join(texts)
  for raw in literals[n]:
   a=whole.index(raw,start);targets.append({'paragraph':n,'start':a,'end':a+len(raw),'raw':raw,'expected_numbers':{'[1,3]':[1,3],'[1–3]':[1,2,3],'[1-3,5]':[1,2,3,5],'[2,1]':[2,1],'[4]':[4]}[raw]});start=a+len(raw)
 d.save(out/'errors.docx')
 # Independent correct control: ordinary six font/size attributes + literal spans only.
 for n in (3,4):
  for r in d.paragraphs[n-1].runs:
   r.font.size=Pt(12);r._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
 for n in payloads:
  p=d.paragraphs[n-1]._p;spans=[x for x in targets if x['paragraph']==n];offset=0
  for r in list(p.findall(qn('w:r'))):
   tx=''.join(r.itertext()) if False else ''.join(t.text or '' for t in r.findall(qn('w:t')));end=offset+len(tx)
   cuts=sorted({offset,end}|{v for x in spans for v in (x['start'],x['end']) if offset<v<end})
   pos=p.index(r);p.remove(r)
   for a,b in zip(cuts,cuts[1:]):
    c=deepcopy(r);t=c.find(qn('w:t'));t.text=tx[a-offset:b-offset];t.set(qn('xml:space'),'preserve')
    if any(x['start']<=a and b<=x['end'] for x in spans):
     rp=c.find(qn('w:rPr'))
     if rp is None:rp=OxmlElement('w:rPr');c.insert(0,rp)
     va=rp.find(qn('w:vertAlign'))
     if va is None:va=OxmlElement('w:vertAlign');rp.append(va)
     va.set(qn('w:val'),'superscript')
    p.insert(pos,c);pos+=1
   offset=end
 d.save(out/'correct.docx')
 scope=json.loads((ROOT/'samples/phase28/scope.json').read_text());scope['selected_citations']=targets;scope['excluded_citation_text']='[8]'
 (out/'scope.json').write_text(json.dumps(scope,ensure_ascii=False,indent=2)+'\n')
 (out/'independent-errors.json').write_text(json.dumps({'source':'合成作者预先定义的字面区间，不是检查器输出','selected_spans':targets,'ordinary_paragraphs':[3,4],'ordinary_properties':['eastAsia→宋体','sz→24（12pt）'],'expected_unselected':['p3首行28pt','p6 REF字段引文','数组[8]','代码p18','其余所有属性与对象'],'control':'correct.docx由字面区间独立设置，未调用修正引擎'},ensure_ascii=False,indent=2)+'\n')
if __name__=='__main__':build(ROOT/'samples/phase30')
