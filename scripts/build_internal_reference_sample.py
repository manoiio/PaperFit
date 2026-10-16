"""Author-owned synthetic targets/errors. No real paper or environment input."""
from pathlib import Path
from io import BytesIO
import sys,json,zipfile
from copy import deepcopy
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
from docx.shared import Pt,Mm
from docx.enum.style import WD_STYLE_TYPE
from lxml import etree as E
from core.package import Package,q,NS
from core.engine import _ensure
ROOT=Path(__file__).resolve().parents[1]
R='http://schemas.openxmlformats.org/officeDocument/2006/relationships';REL='http://schemas.openxmlformats.org/package/2006/relationships';CT='http://schemas.openxmlformats.org/package/2006/content-types'

def run(tx,bad=False):
 r=E.Element(q('r'));rp=E.SubElement(r,q('rPr'));E.SubElement(rp,q('rStyle')).set(q('val'),'DemoLink')
 if bad:E.SubElement(rp,q('sz')).set(q('val'),'30')
 t=E.SubElement(r,q('t'));t.text=tx;t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');return r

def link(target,tx,bad=False):
 h=E.Element(q('hyperlink'));h.set(q('anchor'),target);h.set(q('tooltip'),'本地合成内部目标');h.set(q('history'),'1');h.append(run(tx,bad));return h

def bookmark(p,name,ident):
 a=E.Element(q('bookmarkStart'));a.set(q('name'),name);a.set(q('id'),str(ident));p.insert(1 if p.find(q('pPr')) is not None else 0,a)
 b=E.SubElement(p,q('bookmarkEnd'));b.set(q('id'),str(ident))

def write_parts(path,parts):
 with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
  for n,b in parts.items():z.writestr(n,b)

def build(output):
 out=Path(output);out.mkdir(parents=True,exist_ok=True);d=Document();d.core_properties.author='';d.core_properties.last_modified_by=''
 normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12);normal.element.get_or_add_rPr().get_or_add_rFonts().set(q('eastAsia'),'宋体');normal.paragraph_format.line_spacing=1.5
 s=d.styles.add_style('DemoLink',WD_STYLE_TYPE.CHARACTER);s.font.name='Times New Roman';s.font.size=Pt(12);s.font.color.rgb=__import__('docx').shared.RGBColor.from_string('0563C1');s.font.underline=True;s.element.get_or_add_rPr().get_or_add_rFonts().set(q('eastAsia'),'宋体')
 for sec in d.sections:sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Mm(25);sec.page_width=Mm(210);sec.page_height=Mm(297)
 d.add_paragraph('内部书签与引用完整性软件演示',style='Title')
 p=d.add_paragraph('1 合成目标标题',style='Heading 1');bookmark(p._p,'_TocTarget',1)
 p=d.add_paragraph('正文前文 ');h=link('_TocTarget','正确 Link ');h.append(run('后半文字 Link [1]',True));p._p.append(h);p.add_run(' 后文保持。')
 p=d.add_paragraph('重复显示文字与两个目标：');p._p.append(link('_TocTarget','相同文字',True));p.add_run(' / ');p._p.append(link('TargetTwo','相同文字'))
 p=d.add_paragraph('第二个可靠目标，包含普通文字。');bookmark(p._p,'TargetTwo',2)
 p=d.add_paragraph('缺失目标的两个链接：');p._p.append(link('MissingTarget','显示文字看似正常'));p.add_run(' / ');p._p.append(link('MissingTarget','另一个位置'))
 p=d.add_paragraph('缓存仍显示数字：');f=E.SubElement(p._p,q('fldSimple'));f.set(q('instr'),' PAGEREF MissingTarget \\h ');f.append(run('7'))
 p=d.add_paragraph('拆分复杂域：')
 for tag,value in [('fldChar','begin'),('instrText',' RE'),('instrText','F _TocTarget \\h '),('fldChar','separate'),('t','合成目标标题'),('fldChar','end')]:
  r=E.SubElement(p._p,q('r'));e=E.SubElement(r,q(tag));e.set(q('fldCharType'),value) if tag=='fldChar' else setattr(e,'text',value)
 p=d.add_paragraph('其他字段不作为书签引用：')
 for code in ('PAGE','SEQ Figure','STYLEREF 1'):
  f=E.SubElement(p._p,q('fldSimple'));f.set(q('instr'),code);f.append(run('1 '))
 p=d.add_paragraph('[1] 公开合成资料。');p._p.append(link('_TocTarget','参考 Link',True))
 p=d.add_paragraph('脚注与尾注正文引用：')
 for kind in ('footnote','endnote'):
  r=E.SubElement(p._p,q('r'));rp=E.SubElement(r,q('rPr'));E.SubElement(rp,q('vertAlign')).set(q('val'),'superscript');E.SubElement(r,q(kind+'Reference')).set(q('id'),'1')
 p=d.add_paragraph('代码材料保持：  a = [1]\t# 不纳入正文\n  print(a)')
 # Target in a verified header. Text support for this container remains closed.
 hp=d.sections[0].header.paragraphs[0];hp.add_run('合成页眉目标');bookmark(hp._p,'HeaderTarget',9)
 temp=out/'author.docx';d.save(temp);pkg=Package(temp);parts=dict(pkg.parts);rels=E.fromstring(parts['word/_rels/document.xml.rels']);ct=E.fromstring(parts['[Content_Types].xml'])
 for kind in ('footnote','endnote'):
  member='word/'+kind+'s.xml';E.SubElement(rels,'{'+REL+'}Relationship',Id='demo-'+kind,Type=R+'/'+kind+'s',Target=kind+'s.xml');E.SubElement(ct,'{'+CT+'}Override',PartName='/'+member,ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.'+kind+'s+xml')
  root=E.Element(q(kind+'s'),nsmap={'w':NS['w']})
  for ident,typ,marker in [('-1','separator','separator'),('0','continuationSeparator','continuationSeparator')]:
   n=E.SubElement(root,q(kind));n.set(q('id'),ident);n.set(q('type'),typ);p=E.SubElement(n,q('p'));E.SubElement(E.SubElement(p,q('r')),q(marker))
  n=E.SubElement(root,q(kind));n.set(q('id'),'1');p=E.SubElement(n,q('p'));pp=E.SubElement(p,q('pPr'));sp=E.SubElement(pp,q('spacing'));sp.set(q('line'),'360');sp.set(q('lineRule'),'auto')
  r=E.SubElement(p,q('r'));E.SubElement(r,q(kind+'Ref'));p.append(run(' 普通注释文字 '));p.append(link('_TocTarget','注释内部 Link',True));p.append(link('NoteTarget',' 注释目标 Link'))
  if kind=='footnote':bookmark(p,'NoteTarget',21)
  parts[member]=E.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
 parts['word/_rels/document.xml.rels']=E.tostring(rels,encoding='UTF-8',xml_declaration=True,standalone=True);parts['[Content_Types].xml']=E.tostring(ct,encoding='UTF-8',xml_declaration=True,standalone=True)
 parts['word/document.xml']=E.tostring(pkg.root,encoding='UTF-8',xml_declaration=True,standalone=True);write_parts(out/'errors.docx',parts);temp.unlink()
 # Hand-written oracle: locations and desired font sizes, independent of R44.
 repairs=[{'part':'word/document.xml','paragraph':3,'link':1,'link_run':2,'pt':12},{'part':'word/document.xml','paragraph':4,'link':1,'link_run':1,'pt':12},{'part':'word/document.xml','paragraph':10,'link':1,'link_run':1,'pt':12},*[{'part':'word/'+k+'s.xml','note_id':'1','note_paragraph':1,'link':1,'link_run':1,'pt':12} for k in ('footnote','endnote')]]
 plan={'roles':{str(i):'body' if i in (3,4,5,6,7,8,9) else 'reference_item' if i==10 else 'other' for i in range(1,len(pkg.paragraphs)+1)},'selected_notes':['footnote:1:3','endnote:1:3'],'repairs':repairs,'targets':{'_TocTarget':'valid','TargetTwo':'valid','HeaderTarget':'valid','NoteTarget':'valid','MissingTarget':'missing'},'missing_reference_positions':3,'protected_material':{'12':'合成代码：不检查、不改写空白'},'not_selected':['全部段落属性','未选链接run','书签/字段/缓存/颜色/下划线/字符样式/关系']}
 (out/'independent-plan.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2)+'\n')
 cp=dict(parts);roots={}
 for e in repairs:
  root=roots.setdefault(e['part'],E.fromstring(cp[e['part']]));p=root.find(q('body')).findall(q('p'))[e['paragraph']-1] if e['part']=='word/document.xml' else next(n for n in root if n.get(q('id'))==e['note_id']).findall(q('p'))[e['note_paragraph']-1]
  r=p.findall(q('hyperlink'))[e['link']-1].findall(q('r'))[e['link_run']-1];r.find('w:rPr/w:sz',NS).set(q('val'),'24')
 for n,r in roots.items():cp[n]=E.tostring(r,encoding='UTF-8',xml_declaration=True,standalone=True)
 write_parts(out/'correct.docx',cp)
 # Deliberately invalid boundaries/names only in protection fixture, not legal Word demo.
 bad=dict(parts);root=E.fromstring(bad['word/document.xml']);body=root.find(q('body'));ps=body.findall(q('p'));bookmark(ps[4],'_TocTarget',30);ps[4].findall(q('bookmarkEnd'))[-1].set(q('id'),'99');bad['word/document.xml']=E.tostring(root);write_parts(out/'protected.docx',bad)
if __name__=='__main__':build(ROOT/'samples/phase32')
