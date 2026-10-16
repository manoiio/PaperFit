"""Portable synthetic hyperlink objects, independent manually authored error plan."""
from pathlib import Path
from copy import deepcopy
import sys,json,zipfile
from io import BytesIO
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from docx import Document
from docx.shared import Pt
from docx.enum.style import WD_STYLE_TYPE
from lxml import etree as E
from core.package import Package,q,NS
from core.engine import _ensure
ROOT=Path(__file__).resolve().parents[1];R='http://schemas.openxmlformats.org/officeDocument/2006/relationships';REL='http://schemas.openxmlformats.org/package/2006/relationships'
def run(value,*,style=None,size=None,cn=None,western=None,theme=False):
 r=E.Element(q('r'));rp=_ensure(r,'rPr')
 if style:_ensure(rp,'rStyle').set(q('val'),style)
 if cn:_ensure(rp,'rFonts').set(q('eastAsia'),cn)
 if western:
  for slot in ('ascii','hAnsi'):_ensure(rp,'rFonts').set(q(slot),western)
 if size:_ensure(rp,'sz').set(q('val'),str(round(size*2)))
 if theme:
  for slot,token in [('ascii','minorAscii'),('hAnsi','minorHAnsi'),('eastAsia','minorEastAsia')]:_ensure(rp,'rFonts').set(q(slot+'Theme'),token)
 t=E.SubElement(r,q('t'));t.text=value;t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');return r

def link(rid,children):
 e=E.Element(q('hyperlink'));e.set('{'+R+'}id',rid);e.set(q('tooltip'),'公开合成目标，不自动打开');e.set(q('history'),'1');e.set(q('tgtFrame'),'_blank')
 for r in children:e.append(r)
 return e

def relationships(rows):
 root=E.Element('{'+REL+'}Relationships',nsmap={None:REL})
 for rid,target in rows:E.SubElement(root,'{'+REL+'}Relationship',Id=rid,Type=R+'/hyperlink',Target=target,TargetMode='External')
 return E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)

def authored_bytes(pkg):
 raw=pkg.bytes(True);buf=BytesIO()
 with zipfile.ZipFile(BytesIO(raw)) as src,zipfile.ZipFile(buf,'w',zipfile.ZIP_DEFLATED) as dst:
  for item in src.infolist():dst.writestr(item,src.read(item.filename))
  for name,value in pkg.parts.items():
   if name not in src.namelist():dst.writestr(name,value)
 return buf.getvalue()

def build(output):
 output=Path(output);output.mkdir(parents=True,exist_ok=True);d=Document(ROOT/'samples/phase21/correct.docx');d.paragraphs[0].runs[0].text='普通超链接格式软件演示样本'
 if 'DemoBadLink' not in d.styles:
  s=d.styles.add_style('DemoBadLink',WD_STYLE_TYPE.CHARACTER);s.font.name='Arial';s.font.size=Pt(15);s.font.color.rgb=__import__('docx').shared.RGBColor.from_string('0563C1');s.font.underline=True;s.element.get_or_add_rPr().get_or_add_rFonts().set(q('eastAsia'),'黑体')
 for _ in range(6):d.add_paragraph()
 d.paragraphs[6]._p.append(run('链接前文 ',cn='黑体',size=14));d.paragraphs[6]._p.append(link('demoLink',[run('公开 Link ',style='DemoBadLink',size=12,cn='宋体',western='Times New Roman'),run('后半 Link [1]',style='DemoBadLink')]))
 d.paragraphs[6]._p.append(run(' 链接后文 Word',cn='宋体',western='Arial',size=14))
 p=d.paragraphs[7]._p;p.append(link('demoLink',[run('重复 Link',style='DemoBadLink')]));p.append(run(' 与 '));p.append(link('demoLink',[run('重复 Link',style='DemoBadLink',cn='宋体',western='Times New Roman',size=12)]));p.append(link('themeLink',[run('主题 Link',style='DemoBadLink',size=12,theme=True)]))
 d.paragraphs[8].add_run('参考文献');p=d.paragraphs[9]._p;p.append(run('[1] 公开演示资料 '));p.append(link('demoLink',[run('网址 Guide',style='DemoBadLink')]));p.append(run('，2026。'))
 p=d.paragraphs[10]._p;p.append(run('[2] 正确网址 '));p.append(link('demoLink',[run('正确 Guide',style='DemoBadLink',cn='宋体',western='Times New Roman',size=12)]))
 d.paragraphs[11]._p.append(run('普通段落引文[1]，后文保留。',size=14,cn='宋体',western='Times New Roman'))
 d.add_paragraph('    values = data[0]\t# 人工排除的代码\n    print(values)')
 temp=output/'author.docx';d.save(temp);pkg=Package(temp)
 rel=E.fromstring(pkg.parts['word/_rels/document.xml.rels'])
 for e in E.fromstring(relationships([('demoLink','https://example.invalid/guide'),('themeLink','https://example.invalid/theme')])):rel.append(e)
 pkg.parts['word/_rels/document.xml.rels']=E.tostring(rel,xml_declaration=True,encoding='UTF-8',standalone=True)
 # The SAME r:id has a DIFFERENT target in each independent story.
 for kind in ('footnote','endnote'):
  part='word/'+kind+'s.xml';root=E.fromstring(pkg.parts[part]);note=next(e for e in root if e.get(q('id'))=='1');p=note.findall(q('p'))[0]
  p.append(link('demoLink',[run('注释链接 Note',style='DemoBadLink'),run(' 后段 Note',style='DemoBadLink',cn='宋体',size=12)]))
  pkg.story_roots[part]=root;pkg.parts['word/_rels/'+kind+'s.xml.rels']=relationships([('demoLink','https://example.invalid/'+kind)])
 # Give every theme slot a uniquely interpretable specification; no font installation inference.
 a='http://schemas.openxmlformats.org/drawingml/2006/main';theme=E.fromstring(pkg.parts['word/theme/theme1.xml']);minor=theme.find('.//{'+a+'}minorFont')
 minor.find('{'+a+'}latin').set('typeface','Times New Roman')
 hans=[e for e in minor if e.get('script')=='Hans'];hans[0].set('typeface','宋体')
 pkg.parts['word/theme/theme1.xml']=E.tostring(theme,xml_declaration=True,encoding='UTF-8',standalone=True)
 settings=E.fromstring(pkg.parts['word/settings.xml']);lang=settings.find(q('themeFontLang'))
 if lang is None:lang=E.SubElement(settings,q('themeFontLang'))
 lang.set(q('val'),'en-US');lang.set(q('eastAsia'),'zh-CN');pkg.parts['word/settings.xml']=E.tostring(settings,xml_declaration=True,encoding='UTF-8',standalone=True)
 (output/'errors.docx').write_bytes(authored_bytes(pkg));temp.unlink()
 roles={str(i):'body' if i in (7,8,12) else 'reference_title' if i==9 else 'reference_item' if i in (10,11) else 'other' for i in range(1,len(pkg.paragraphs)+1)}
 scope={'roles':roles,'materials':{'13':'合成程序代码；保留文字、空格、制表和换行'},'citation_scan':[7,8,12],'protected_citation_paragraphs':[7],'selected_notes':['footnote:1:3','endnote:1:3','endnote:3:5']}
 (output/'scope.json').write_text(json.dumps(scope,ensure_ascii=False,indent=2)+'\n')
 # Independent authorized repair list; NOT copied from checker findings.
 errors=[{'part':'word/document.xml','paragraph':7,'run':1,'property':'font_eastAsia','value':'宋体'}, {'part':'word/document.xml','paragraph':7,'run':1,'property':'size','value':12},
 {'part':'word/document.xml','paragraph':7,'link':1,'link_run':2,'property':'font_eastAsia','value':'宋体'},{'part':'word/document.xml','paragraph':7,'link':1,'link_run':2,'property':'font_ascii','value':'Times New Roman'},{'part':'word/document.xml','paragraph':7,'link':1,'link_run':2,'property':'size','value':12},
 {'part':'word/document.xml','paragraph':7,'run':2,'property':'font_ascii','value':'Times New Roman'},{'part':'word/document.xml','paragraph':7,'run':2,'property':'size','value':12},
 *[{'part':'word/document.xml','paragraph':n,'link':1,'link_run':1,'property':prop,'value':val} for n in (8,10) for prop,val in [('font_eastAsia','宋体'),('font_ascii','Times New Roman'),('size',12)] if not (n==10 and prop=='font_ascii')],
 *[{'part':'word/'+k+'s.xml','note_id':'1','note_paragraph':1,'link':1,'link_run':1,'property':prop,'value':val} for k in ('footnote','endnote') for prop,val in [('font_eastAsia','宋体'),('size',12)]]]
 (output/'independent-errors.json').write_text(json.dumps({'kind':'作者按注入位置独立制定，非程序检查结果','allowed_repairs':errors,'not_selected':['正文p12的字号与R33','所有段落属性','注释编号、未选注释、西文字体、颜色、下划线、样式','链接目标及全部关系','人工排除代码和表格图片']},ensure_ascii=False,indent=2)+'\n')
 # Correct control applies the authored whitelist directly, independently of checker/fixer.
 base=Package(output/'errors.docx');roots={'word/document.xml':base.root}
 for e in errors:
  part=e['part'];root=roots.setdefault(part,E.fromstring(base.parts[part]));p=root.find(q('body')).findall(q('p'))[e['paragraph']-1] if part=='word/document.xml' else next(x for x in root if x.get(q('id'))==e['note_id']).findall(q('p'))[e['note_paragraph']-1]
  r=p.findall(q('hyperlink'))[e['link']-1].findall(q('r'))[e['link_run']-1] if 'link' in e else p.findall(q('r'))[e['run']-1];rp=_ensure(r,'rPr')
  if e['property']=='size':_ensure(rp,'sz').set(q('val'),str(round(e['value']*2)))
  else:
   slot=e['property'][5:];rf=_ensure(rp,'rFonts');rf.set(q(slot),e['value']);rf.attrib.pop(q(slot+'Theme'),None)
 for part,root in roots.items():
  if part!='word/document.xml':base.story_roots[part]=root
 (output/'correct.docx').write_bytes(base.bytes(True))
 # Complex/invalid relationships fixture is intentionally separate from legal visual output.
 bad=Package(output/'errors.docx');bad.paragraphs[7].findall(q('hyperlink'))[0].set('{'+R+'}id','missing');nested=bad.paragraphs[7].findall(q('hyperlink'))[1];nested.append(E.Element(q('sdt')));(output/'protected.docx').write_bytes(bad.bytes(True))
if __name__=='__main__':build(ROOT/'samples/phase29')
