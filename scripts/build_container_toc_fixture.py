"""Public author-labelled SDT/decimal-heading integration fixture; no model data."""
from pathlib import Path
import sys,json
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
def build(folder):
 from docx import Document
 from docx.shared import Pt
 from lxml import etree as E
 from core.package import q,NS,Package
 folder=Path(folder);folder.mkdir(parents=True,exist_ok=True)
 d=Document();d.add_paragraph('公开合成论文 封面','Title');d.add_paragraph('摘要');d.add_paragraph('此处是可公开的结构与目录更新验证摘要。')
 d.add_paragraph('目录');d.add_paragraph('旧目录99')
 for title,level in [('研究方法',1),('实现步骤',2),('参数设置',3),('研究方法',1)]:
  d.add_paragraph(title,'Heading '+str(level))
  for i in range(9):d.add_paragraph('这是一段公开合成的论文内容，用于验证最终字号行距变化后的真实分页、目录编号与精确目标。'*3)
 path=folder/'authored.docx';d.save(path);p=Package(path)
 root=E.fromstring(p.parts['word/numbering.xml']);a=E.SubElement(root,q('abstractNum'));a.set(q('abstractNumId'),'42');E.SubElement(a,q('multiLevelType')).set(q('val'),'multilevel')
 for i in range(3):
  lv=E.SubElement(a,q('lvl'));lv.set(q('ilvl'),str(i));E.SubElement(lv,q('start')).set(q('val'),'1');E.SubElement(lv,q('numFmt')).set(q('val'),'decimal');E.SubElement(lv,q('pStyle')).set(q('val'),'Heading'+str(i+1));E.SubElement(lv,q('lvlText')).set(q('val'),'.'.join('%'+str(j+1) for j in range(i+1)));E.SubElement(lv,q('suff')).set(q('val'),'tab')
 for ident,start in ((42,1),(43,5)):
  n=E.SubElement(root,q('num'));n.set(q('numId'),str(ident));E.SubElement(n,q('abstractNumId')).set(q('val'),'42')
  if start!=1:ov=E.SubElement(n,q('lvlOverride'));ov.set(q('ilvl'),'0');E.SubElement(ov,q('startOverride')).set(q('val'),str(start))
 p.parts['word/numbering.xml']=E.tostring(root)
 roles={};count=0
 for n,par in enumerate(p.paragraphs,1):
  st=par.find('w:pPr/w:pStyle',NS);sid=st.get(q('val')) if st is not None else ''
  if n>5 and sid.startswith('Heading'):
   pr=par.find(q('pPr'));num=E.SubElement(pr,q('numPr'));E.SubElement(num,q('ilvl')).set(q('val'),str(int(sid[-1])-1));E.SubElement(num,q('numId')).set(q('val'),'43' if n>30 else '42');roles[str(n)]='heading'+sid[-1];count+=1
  else:roles[str(n)]='body' if n>5 else 'other'
 # Genuine block-level TOC wrapper, title remains outside the field and inside shell.
 body=p.root.find(q('body'));title=p.paragraphs[3];cache=p.paragraphs[4];at=body.index(title);body.remove(title);body.remove(cache)
 sdt=E.Element(q('sdt'));pr=E.SubElement(sdt,q('sdtPr'));E.SubElement(pr,q('id')).set(q('val'),'52001');E.SubElement(pr,q('alias')).set(q('val'),'公开合成自动目录');E.SubElement(pr,q('tag')).set(q('val'),'paperfit-public-toc');gal=E.SubElement(pr,q('docPartObj'));E.SubElement(gal,q('docPartGallery')).set(q('val'),'Table of Contents');E.SubElement(gal,q('docPartUnique'));content=E.SubElement(sdt,q('sdtContent'));content.append(title)
 cache=E.SubElement(content,q('p'));f=E.SubElement(cache,q('fldSimple'));f.set(q('instr'),' TOC \\o "1-3" \\h \\z \\u ');r=E.SubElement(f,q('r'));E.SubElement(r,q('t')).text='旧目录99';body.insert(at,sdt)
 source=folder/'source.docx';source.write_bytes(p.bytes(True))
 pkg=Package(source);types={str(i):('heading'+par.find('w:pPr/w:pStyle',NS).get(q('val'))[-1] if par.find('w:pPr/w:pStyle',NS) is not None and par.find('w:pPr/w:pStyle',NS).get(q('val'),'').startswith('Heading') and i>3 else 'body' if i>3 else 'other') for i,par in enumerate(pkg.paragraphs,1)}
 (folder/'roles.json').write_text(json.dumps(types,ensure_ascii=False,indent=2)+'\n')
 return source
if __name__=='__main__':print(build(sys.argv[1]))
