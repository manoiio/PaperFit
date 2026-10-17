"""Public author-labelled combined fixture, independent of classifiers/planners."""
from pathlib import Path
from io import BytesIO
import sys,json
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from docx import Document
from docx.shared import Mm,Pt
from lxml import etree as E
from core.package import Package,q,NS

def build(folder,ambiguous=False):
 folder=Path(folder);folder.mkdir(parents=True,exist_ok=True);d=Document();d.sections[0].left_margin=Mm(10);d.sections[0].right_margin=Mm(10)
 roles=[]
 def p(tx,role='body',style=None):
  x=d.add_paragraph(tx,style=style);roles.append(role);return x
 p('封面 — 公开合成图表论文','other','Title')
 cover=d.add_table(rows=1,cols=2);cover.cell(0,0).text='姓名';cover.cell(0,1).text='公开合成作者'
 p('封面 — 合成身份表结束','other')
 p('摘要','abstract_cn_title').paragraph_format.page_break_before=True
 p('本公开合成论文用于验证图表、分页与目录处理。全部文字和数值均为合成，不含真实论文。','abstract_cn_body')
 p('目录','other');f=E.SubElement(p('','other')._p,q('fldSimple'),{q('instr'):' TOC \\o "1-2" \\h '});E.SubElement(E.SubElement(f,q('r')),q('t')).text='旧目录……99'
 p('','other')
 p('第1章 合成方法','heading1','Heading 1')
 p('本章介绍公开合成的比较方法，图1-2展示有限样本趋势。以下图表文字保持不变，安全辅助只改变已授权位置和尺寸。')
 p('图1-2 合成趋势','figure_caption')
 media=next(b for n,b in Package(ROOT/'samples/extended-correct.docx').parts.items() if n.startswith('word/media/'))
 image=d.add_picture(BytesIO(media),width=Mm(180));roles.append('other')
 p('图片之后继续叙述，验证未选内容及全部文字保持。这里不会为了适配尺寸删去任何数据或改变图片像素。')
 p('表1-1 合成结果','table_caption')
 t=d.add_table(rows=37,cols=2);t.autofit=False
 for row in t.rows:
  for i,c in enumerate(row.cells):c.width=Mm(88);c.text='方案' if row is t.rows[0] else '公开样本数值及说明'
 for i,row in enumerate(t.rows):row.cells[0].text='方案' if i==0 else f'方案{i}';row.cells[1].text='耗时（秒）' if i==0 else f'{i+10} 合成测量值，独立保留。'
 # Existing tblHeader is structural evidence, not inferred from labels.
 E.SubElement(t.rows[0]._tr.get_or_add_trPr(),q('tblHeader'),{q('val'):'1'})
 p('表格后保留的叙述段落，不参与任意对象移动。安全重复表头不会生成续表文字或编号。')
 p('表1-2 复杂对象保持','table_caption');complex_=d.add_table(rows=2,cols=2);complex_.cell(0,0).merge(complex_.cell(0,1));complex_.cell(0,0).text='合并保护';complex_.cell(1,0).text='保持';complex_.cell(1,1).text='原样'
 p('复杂对象之后的普通正文用于验证保留和分节边界。其内容、链接及未选格式不因图表辅助而变化。')
 p('第2章 合成结果','heading1','Heading 1')
 target=p('本章为目录和页码引用提供一个明确目标。安全PAGE与PAGEREF是不同类型的域，不能混淆缓存与实际页码。')
 E.SubElement(target._p,q('bookmarkStart'),{q('id'):'20',q('name'):'ResultTarget'});E.SubElement(target._p,q('bookmarkEnd'),{q('id'):'20'})
 field=E.SubElement(p('引用结果：','pending')._p,q('fldSimple'),{q('instr'):' PAGEREF ResultTarget \\h '});E.SubElement(E.SubElement(field,q('r')),q('t')).text='99'
 p('结尾普通正文保留，所有中间轮次不覆盖原稿。本文为功能验收合成内容，并不代表任何真实论文全部排版合格。')
 if ambiguous:
  unknown=d.add_table(rows=2,cols=2)
  for i,r in enumerate(unknown.rows):
   for j,c in enumerate(r.cells):c.text=f'条目{i}{j}'
 d.save(folder/'source.docx');pkg=Package(folder/'source.docx')
 t=pkg.root.find(q('body')).findall(q('tbl'))[1];twips=9978;cols=[4989,4989]
 for col,c in zip(t.findall('w:tblGrid/w:gridCol',NS),cols):col.set(q('w'),str(c))
 t.find('w:tblPr/w:tblW',NS).set(q('w'),str(twips));t.find('w:tblPr/w:tblW',NS).set(q('type'),'dxa')
 for row in t.findall(q('tr')):
  for cell,c in zip(row.findall(q('tc')),cols):cell.find('w:tcPr/w:tcW',NS).set(q('w'),str(c))
 (folder/'source.docx').write_bytes(pkg.bytes(True))
 types={str(i):r for i,r in enumerate(roles,1)};(folder/'roles.json').write_text(json.dumps(types,ensure_ascii=False,indent=2)+'\n')
 expected={'author':'hand-labelled before execution','table_roles':{'1':'cover','2':'content','3':'pending'},'move':{'caption':10,'image':11,'target':'image then caption'},'width_tables':[2],'width_image':11,'table_header':{'2':1},'protected_tables':[1,3],'TOC_titles':['第1章 合成方法','第2章 合成结果'],'PAGEREF_target':'ResultTarget','all_original_text_preserved':True}
 (folder/'expected.json').write_text(json.dumps(expected,ensure_ascii=False,indent=2)+'\n');return folder/'source.docx',types
if __name__=='__main__':build(sys.argv[1])
