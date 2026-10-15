"""Build a public synthetic composite; never reads private papers, work/ or outputs/."""
from pathlib import Path
from io import BytesIO
import json,sys
from copy import deepcopy
from docx import Document
from docx.shared import Mm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from core.package import Package,write_new

def build(output):
    output=Path(output);output.mkdir(parents=True,exist_ok=True)
    d=Document(ROOT/'samples/phase24/combined-errors.docx')
    # Original13 paragraph roles remain intact; new objects appended explicitly.
    section=d.add_section();section.header.paragraphs[0].text='软件演示页眉（合成资料）'
    d.add_paragraph('参考文献')
    d.add_paragraph('[1] 软件演示资料，2026。')
    d.add_paragraph('这是一段仅供软件演示的正文，采用普通手工引文[1]，不涉及真实作者资料。')
    from samples.build_samples import png
    image=d.add_paragraph();image.add_run().add_picture(BytesIO(png()),width=Mm(35))
    d.add_paragraph('图1-1 软件演示结构')
    # A content-control TOC clue, deliberately unconfirmed by the manifest.
    sdt=OxmlElement('w:sdt');pr=OxmlElement('w:sdtPr');sdt.append(pr);content=OxmlElement('w:sdtContent');sdt.append(content)
    p=OxmlElement('w:p');pp=OxmlElement('w:pPr');style=OxmlElement('w:pStyle');style.set(qn('w:val'),'TOC1');pp.append(style);p.append(pp)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),' TOC \\o "1-3" ');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='1 软件演示章节\t1';r.append(t);field.append(r);p.append(field);content.append(p)
    body=d._element.body;body.insert(len(body)-1,sdt)
    raw=BytesIO();d.save(raw)
    # python-docx preserves legal note parts and their relationships on save.
    path=write_new(output/'composite.docx',raw.getvalue());pkg=Package(path)
    roles=json.loads((ROOT/'samples/phase24/independent-scope.json').read_text())['roles']
    # add_section adds a top-level empty paragraph at14; remaining roles fixed here independently.
    roles.update({'14':'other','15':'reference_title','16':'reference_item','17':'body','18':'other','19':'figure_caption'})
    ds={'description':'软件演示综合样本，独立人工定义；不是程序分类结果。编号含空分节段，表格另计。',
        'roles':roles,'material_exclusions':{'13':'合成程序代码，保留制表符和空白，不套普通文字'},
        'table_types':{'1':'content'},'cover_objects':{'p:1':'title'},'section_roles':{'1':'mixed','2':'mixed'},
        'toc_objects':{'d:3':'title','d:4':'level1'},'toc_complete':False,'headings_complete':False,
        'citation_scan':[8,9,17],'citation_use':{'paragraph':17,'text':'[1]'},
        'caption_roles':{'10':'normal','19':'normal'},'chapter_map':{'7':1},'associations':{'10':'table:1','19':'image:p:18'},
        'notes':'明确选择全部6条普通注释（其中复杂链接/域继续保护）；系统对象不选择',
        'excluded_toc':'d:20 内容控件自动目录线索留待复核；不据此判断整篇缺项'}
    (output/'scope.json').write_text(json.dumps(ds,ensure_ascii=False,indent=2)+'\n')
    assert len(pkg.paragraphs)==19
    return path
if __name__=='__main__':build(sys.argv[1] if len(sys.argv)>1 else ROOT/'samples/phase26')
