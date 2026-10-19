"""Small independently authored public cases; no local paper derivatives."""
from pathlib import Path
import json

def build(folder):
 from docx import Document
 from docx.shared import Pt
 from docx.oxml import OxmlElement
 from docx.oxml.ns import qn
 out=Path(folder);out.mkdir(parents=True,exist_ok=True)
 d=Document();d.core_properties.author='PaperFit public synthetic fixture'
 d.styles['Normal'].font.name='宋体';d.styles['Normal'].font.size=Pt(12)
 d.add_heading('第1章 公开合成参数核对',1)
 d.add_paragraph('这是独立编写的普通正文，用来验证安全题注属性与只读结构核对，不包含真实论文内容。')
 for text,alignment in [('表1-1 合成参数与单位',0),('表1-2 已正确的参数表',1),('表1-3 有偏移的合成题注',0)]:
  p=d.add_paragraph(text);p.alignment=alignment
  if '偏移' in text:p.paragraph_format.first_line_indent=Pt(12)
  t=d.add_table(rows=2,cols=2);b=OxmlElement('w:tblBorders')
  for side in ('left','right','top','bottom','insideH','insideV'):
   e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');b.append(e)
  t._tbl.tblPr.append(b);t.cell(0,0).text='参数';t.cell(0,1).text='单位';t.cell(1,0).text='长度';t.cell(1,1).text='毫米'
 from PIL import Image,ImageDraw
 from io import BytesIO
 im=Image.new('RGB',(300,60),'white');draw=ImageDraw.Draw(im);draw.rectangle((5,5,295,55),outline='black',width=2);draw.text((70,22),'PUBLIC SYNTHETIC FLOW',fill='black');data=BytesIO();im.save(data,format='PNG');data.seek(0);d.add_picture(data,width=__import__('docx.shared',fromlist=['Cm']).Cm(6))
 p=d.add_paragraph('图1-1 公开合成流程');p.alignment=0
 # The preceding image is public synthetic content; association still uses original validators.
 p=d.add_paragraph('图1-2 受保护的字段题注');f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),' SEQ Figure ');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='1';r.append(t);f.append(r);p._p.append(f)
 p=d.add_paragraph('缓存看似正常，实际目标不存在：');f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),' REF MissingPublicTarget ');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='1';r.append(t);f.append(r);p._p.append(f)
 path=out/'source.docx';d.save(path)
 expected={'captions':{'3':'center','4':'center','5':'center','7':'center'},'protected_paragraph':8,'offset_paragraph':5,'missing_target':'MissingPublicTarget','originally_centered':4,'text_and_non_document_parts_must_equal':True}
 (out/'expected.json').write_text(json.dumps(expected,ensure_ascii=False,indent=2)+'\n')
 return path
if __name__=='__main__':
 import sys
 print(build(sys.argv[1]))
