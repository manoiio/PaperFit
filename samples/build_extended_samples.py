"""独立扩展样本，保留原阶段样本；所有新输出拒绝覆盖。"""
from pathlib import Path
from io import BytesIO
import json,hashlib
from docx import Document
from docx.shared import Mm,Pt
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .build_samples import element, fonts, png
ROOT=Path(__file__).resolve().parent
TYPES={i:'other' for i in range(1,25)}
TYPES.update({7:'heading1',8:'body',9:'heading2',10:'body',11:'heading3',12:'body',13:'heading2',14:'body',16:'figure_caption',17:'table_caption',19:'heading1',20:'body'})

def build(error=False):
    d=Document();d.core_properties.author='软件演示';d.core_properties.last_modified_by='论文有格'
    d.core_properties.title='论文有格标题题注软件演示样本'
    normal=d.styles['Normal'];fonts(normal.element.get_or_add_rPr());element(d.styles.element.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}docDefaults').find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPrDefault').find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rPr'),'szCs',val=24)
    normal.paragraph_format.line_spacing=1.5;normal.paragraph_format.space_after=Pt(4)
    element(normal.element.get_or_add_pPr(),'ind',firstLineChars=200)
    for level,size in ((1,18),(2,15),(3,14)):
        s=d.styles[f'Heading {level}'];fonts(s.element.get_or_add_rPr(),cn='黑体');element(s.element.get_or_add_rPr(),'sz',val=round(size*2))
        element(s.element.get_or_add_rPr(),'color',val='000000');s.paragraph_format.line_spacing=1
        s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER if level==1 else WD_ALIGN_PARAGRAPH.LEFT
        s.paragraph_format.space_before=Pt(4);s.paragraph_format.space_after=Pt(4)
        element(s.element.get_or_add_pPr(),'ind',firstLineChars=0,firstLine=0)
    caption=d.styles['Caption'];fonts(caption.element.get_or_add_rPr(),cn='宋体',en='宋体');element(caption.element.get_or_add_rPr(),'sz',val=21)
    element(caption.element.get_or_add_pPr(),'ind',firstLineChars=0,firstLine=0);caption.paragraph_format.line_spacing=1
    title=d.styles['Title'];fonts(title.element.get_or_add_rPr(),cn='黑体');element(title.element.get_or_add_rPr(),'sz',val=36)
    element(title.element.get_or_add_pPr(),'ind',firstLineChars=0,firstLine=0)
    title.paragraph_format.line_spacing=1;title.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    border=title.element.get_or_add_pPr().find(qn('w:pBdr'))
    if border is not None:title.element.get_or_add_pPr().remove(border)
    for s in (title,caption):
        color=element(s.element.get_or_add_rPr(),'color',val='000000')
        color.attrib.clear();color.set(qn('w:val'),'000000')
    for section in d.sections:
        section.page_width,section.page_height=Mm(210),Mm(297)
        for side in ('top','bottom','left','right'):setattr(section,side+'_margin',Mm(25))
    d.add_paragraph('论文有格标题题注软件演示样本','Title') #1
    d.add_paragraph('本文件是脱敏合成软件演示样本，作者、单位及实验数据均为虚构。') #2
    d.add_paragraph('摘要') #3
    d.add_paragraph('这是一段软件演示摘要，不应批量识别成正文。') #4
    d.add_paragraph('目录') #5
    d.add_paragraph('第一章 本地格式流程 ........ 1') #6
    p=d.add_heading('第一章 本地格式流程',level=1) #7
    if error:
        element(p.runs[0]._r.get_or_add_rPr(),'rFonts',eastAsia='宋体');p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.line_spacing=1.5
    d.add_paragraph('本文使用合成内容验证段落类型建议，所有检查范围均须由用户确认。') #8
    p=d.add_paragraph(style='Heading 2');p.add_run('1.1 ');r=p.add_run('段落类型核对') #9
    cb=d.styles.add_style('DemoHeadingCharacterBase',WD_STYLE_TYPE.CHARACTER)
    element(cb.element.get_or_add_rPr(),'sz',val=28 if error else 30)
    cc=d.styles.add_style('DemoHeadingCharacter',WD_STYLE_TYPE.CHARACTER);cc.base_style=cb;r.style=cc
    d.add_paragraph('1.5倍行距是本段讨论的格式值，不是标题编号；程序应依据完整句子判断。') #10
    p=d.add_paragraph(style='Heading 3');p.add_run('1.1.1 ');r=p.add_run('普通片段格式') #11
    if error:r.font.size=Pt(12)
    d.add_paragraph('2026年仅为软件演示中的数字，不是标题编号，也不包含真实作者信息。') #12
    p=d.add_paragraph(style='Normal' if error else 'Heading 2');p.add_run('1.2 手工格式标题') #13
    if error:
        fonts(p.runs[0]._r.get_or_add_rPr(),cn='黑体');p.runs[0].font.size=Pt(15)
        p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.line_spacing=1
        p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4)
        element(p._p.get_or_add_pPr(),'ind',firstLineChars=0,firstLine=0);element(p._p.get_or_add_pPr(),'outlineLvl',val=9)
    point=d.styles.add_style('DemoExtendedPointBody',WD_STYLE_TYPE.PARAGRAPH);fonts(point.element.get_or_add_rPr());point.paragraph_format.line_spacing=1.5;point.paragraph_format.first_line_indent=Pt(24)
    d.add_paragraph('本段使用二十四磅点值缩进，十二磅字号时等效两个汉字；改变字号须复核。','DemoExtendedPointBody') #14
    d.add_paragraph().add_run().add_picture(BytesIO(png()),width=Mm(75)) #15
    p=d.add_paragraph(style='Caption');p.add_run('图1-1 ');r=p.add_run('本地检查流程（软件演示）') #16
    if error:element(r._r.get_or_add_rPr(),'rFonts',eastAsia='黑体');r.font.size=Pt(12)
    p=d.add_paragraph('表1-1 合成测试项目（软件演示）','Caption') #17
    if error:p.runs[0].font.size=Pt(12)
    t=d.add_table(rows=2,cols=2);t.style='Table Grid'
    for row,values in zip(t.rows,[('对象','保护要求'),('文本与图片','格式修正不重写内容')]):
        for cell,text in zip(row.cells,values):cell.text=text
    d.add_section(WD_SECTION_START.NEW_PAGE) #18
    d.add_heading('第二章 软件演示结果',level=1) #19
    d.add_paragraph('当前流程保留原稿、各轮副本和复查报告，未覆盖对象继续由操作者人工核对。') #20
    d.add_paragraph('参考文献') #21
    d.add_paragraph('[1] 演示资料编写组. 合成测试说明[M]. 软件演示, 2026.') #22
    d.add_paragraph('实验观察') #23 deliberately ambiguous under references
    p=d.add_paragraph() #24 field exclusion, survives other changes
    for tag,val in [('fldChar','begin'),('instrText','PAGE'),('fldChar','end')]:
        r=OxmlElement('w:r');e=OxmlElement('w:'+tag)
        if tag=='fldChar':e.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}fldCharType',val)
        else:e.text=val
        r.append(e);p._p.append(r)
    return d

if __name__=='__main__':
    for name,error in [('extended-correct',False),('extended-errors',True)]:
        path=ROOT/(name+'.docx')
        if path.exists() or path.with_suffix('.types.json').exists():raise SystemExit('输出已存在，不覆盖')
        build(error).save(path)
        manifest={'input_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'paragraph_types':{str(k):v for k,v in TYPES.items()},'confirmed':True,'confirmation':'合成样本作者人工核对；程序建议仍须网页用户再次确认'}
        path.with_suffix('.types.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
