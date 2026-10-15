"""Synthetic four-section cover/story fixtures with independent injected errors."""
from pathlib import Path
from hashlib import sha256
import json
from docx import Document
from docx.shared import Mm,Pt
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement as X
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parent

def el(parent,tag,**attrs):
 e=X('w:'+tag)
 for k,v in attrs.items():e.set(qn('w:'+k),str(v))
 parent.append(e);return e

def font(r,name='宋体',size=12,b=False):
 r.font.name='Times New Roman';r.font.size=Pt(size);r.bold=b;r._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),name)
def para(p,text,name='宋体',size=12,b=False,center=False):
 p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT;font(p.add_run(text),name,size,b);return p

def page(p,simple=True,result='1'):
 p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 if simple:
  f=el(p._p,'fldSimple',instr=' PAGE \\* MERGEFORMAT ');r=el(f,'r');rp=el(r,'rPr');el(rp,'sz',val=21);el(r,'t').text=result
 else:
  for typ,tx in [('begin',None),(None,' PAGE '),('separate',None),(None,result),('end',None)]:
   r=p.add_run();font(r,size=10.5)
   if typ:el(r._r,'fldChar',fldCharType=typ)
   elif tx.strip()=='PAGE':el(r._r,'instrText').text=tx
   else:r.add_text(tx)

def build(errors=False):
 d=Document();d.core_properties.author='软件演示';d.core_properties.title='封面分节软件演示'
 normal=d.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12);normal.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体');normal.paragraph_format.line_spacing=1.5
 d.settings.odd_and_even_pages_header_footer=True
 p=d.add_paragraph(style='Title');para(p,'软件演示论文封面',size=14,center=True)
 p=d.add_paragraph();para(p,'单文档格式',name='黑体',size=22,b=True,center=True);r=p.add_run('与分节验证');font(r,'宋体' if errors else '黑体',12 if errors else 22,not errors)
 if errors:p.alignment=WD_ALIGN_PARAGRAPH.LEFT
 t=d.add_table(3,2);t.style='Table Grid'
 labels=['学生姓名：','学生学号：','学院：']
 for i,label in enumerate(labels):
  para(t.cell(i,0).paragraphs[0],label);para(t.cell(i,1).paragraphs[0],'待填写' if i==0 else '软件演示字段')
 p=d.add_paragraph();para(p,'专业：软件演示专业')
 p=d.add_paragraph();para(p,'指导教师：待填写')
 p=d.add_paragraph();para(p,'提交日期：年 月 日')
 p=d.add_paragraph();para(p,'学生姓名：重复候选（软件演示）')
 p=d.add_paragraph();para(p,'本文件仅用于软件演示，不是真实论文或个人信息。',size=10.5)
 # Distinct front section; first-page different off, safe ordinary header separate from PAGE footer.
 s=d.add_section(WD_SECTION_START.NEW_PAGE);s.header.is_linked_to_previous=False;s.footer.is_linked_to_previous=False
 para(s.header.paragraphs[0],'软件演示前置部分',size=10.5,center=not errors);page(s.footer.paragraphs[0],True,'I')
 el(s._sectPr,'pgNumType',fmt='upperRoman',start=1)
 para(d.add_paragraph(),'摘 要',name='黑体',size=18,center=True)
 para(d.add_paragraph(),'软件演示摘要内容用于检查前置部分页码字段。保留内容，不自动扩写摘要。')
 s=d.add_section(WD_SECTION_START.NEW_PAGE);s.header.is_linked_to_previous=False;s.footer.is_linked_to_previous=False
 para(s.header.paragraphs[0],'第一章 软件演示',size=10.5,center=True)
 s.even_page_header.is_linked_to_previous=False;para(s.even_page_header.paragraphs[0],'华南理工大学学士学位论文',size=10.5,center=True)
 page(s.footer.paragraphs[0],False,'1')
 for old in list(s._sectPr.findall(qn('w:pgNumType'))):s._sectPr.remove(old)
 el(s._sectPr,'pgNumType',fmt='decimal',start=1)
 para(d.add_paragraph(),'第一章 软件演示正文',name='黑体',size=18,center=True)
 for i in range(3):para(d.add_paragraph(),'这是软件演示正文。此处仅用于核对分节、普通页眉和复杂PAGE字段；不包含真实作者资料。')
 s=d.add_section(WD_SECTION_START.NEW_PAGE)
 # Inherits section3 stories deliberately; role mixed blocks ordinary shared formatting.
 para(d.add_paragraph(),'混合用途演示',name='黑体',size=14)
 para(d.add_paragraph(),'本节故意不拆分，用于说明用途混合时页码判断需要人工核对。')
 for i,s in enumerate(d.sections,1):
  s.page_width=Mm(210);s.page_height=Mm(297);s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Mm(25)
  s.header_distance=Mm(9 if errors and i==2 else 15);s.footer_distance=Mm(8 if errors and i==3 else 15)
 return d

def main():
 for bad,name in [(False,'front-correct'),(True,'front-errors')]:
  dest=ROOT/(name+'.docx')
  if dest.exists():raise RuntimeError('拒绝覆盖样本')
  d=build(bad);d.save(dest)
  from core.front import make_scope
  from core.package import Package
  pkg=Package(dest);s=make_scope(dest);s.update(confirmed=True,confirmation='合成样本人工独立定位；用户使用仍须核对')
  s['template']['name']='封面分节软件演示补充要求'
  s['template']['sources']['header_text__alignment']={'kind':'user_confirmed','quote':'','location':None}
  s['objects']={'p:2':'title','t:1:1:1:1':'name','t:1:1:2:1':'name','t:1:2:2:1':'student_id','t:1:3:2:1':'college','p:3':'major','p:4':'supervisor','p:5':'date'}
  s['section_roles']={'1':'cover','2':'front','3':'body','4':'mixed'}
  m={'input_sha256':pkg.hash,'paragraph_types':{str(i):'other' for i in range(1,len(pkg.paragraphs)+1)},'confirmed':True,'confirmation':'软件演示，只确认封面/分节新能力，其他类型排除','front_scope':s}
  dest.with_suffix('.types.json').write_text(json.dumps(m,ensure_ascii=False,indent=2)+'\n')
 gold={'description':'独立人工注入；不含程序检查结果。其他结构/页码只读待复核不应清零。','errors':[{'rule':'R24','locator':'p:2/run:2','property':'font_cn','actual':'宋体','expected':'黑体'},{'rule':'R24','locator':'p:2/run:2','property':'font_size','actual':12,'expected':22},{'rule':'R24','locator':'p:2/run:2','property':'bold','actual':False,'expected':True},{'rule':'R24','locator':'p:2','property':'alignment','actual':'left','expected':'center'},{'rule':'R26','section':2,'property':'header_distance_mm','actual':9,'expected':15},{'rule':'R26','section':3,'property':'footer_distance_mm','actual':8,'expected':15},{'rule':'R27','section':2,'story':'header:default','property':'alignment','actual':'left','expected':'center'}]}
 (ROOT/'front-expected-errors.json').write_text(json.dumps(gold,ensure_ascii=False,indent=2)+'\n')
if __name__=='__main__':main()
