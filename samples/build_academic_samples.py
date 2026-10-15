"""Stage15 independent synthetic fixtures; no personal document data."""
from pathlib import Path
from io import BytesIO
import json,hashlib
from docx import Document
from docx.shared import Mm,Pt
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from samples.build_samples import element,fonts,png
ROOT=Path(__file__).resolve().parent
TYPES={1:'other',2:'other',3:'abstract_cn_title',4:'abstract_cn_body',5:'keywords_cn',6:'abstract_en_title',7:'abstract_en_body',8:'keywords_en',9:'heading1',10:'body',11:'heading2',12:'body',13:'heading3',14:'body',15:'other',16:'figure_caption',17:'table_caption',18:'reference_title',19:'reference_item',20:'reference_item',21:'reference_item'}
CN_TEXT=('本软件演示通过合成段落核对格式属性，所有文字不含真实作者或课程资料。处理过程保留文档内容与对象，操作者明确确认类型后才检查所选范围。段落间距使用原单位，不将行单位直接换成磅值。关键词与参考文献仅检查确定的格式和数量，不生成学术结论。')*4
CN_TEXT=CN_TEXT[:450]
TYPES.update({i:'reference_item' for i in range(22,29)})

def build(error=False):
 d=Document();d.core_properties.author='软件演示';d.core_properties.last_modified_by='论文有格';d.core_properties.title='单文档前后置软件演示'
 for s in d.sections:
  s.page_width,s.page_height=Mm(210),Mm(297)
  for side in ('top','bottom','left','right'):setattr(s,side+'_margin',Mm(25))
 normal=d.styles['Normal'];fonts(normal.element.get_or_add_rPr());normal.paragraph_format.line_spacing=1.5;normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;normal.paragraph_format.space_before=Pt(0);normal.paragraph_format.space_after=Pt(0)
 element(normal.element.get_or_add_pPr(),'ind',firstLineChars=200)
 for lvl,size in [(1,18),(2,15),(3,14)]:
  s=d.styles[f'Heading {lvl}'];fonts(s.element.get_or_add_rPr(),cn='黑体');element(s.element.get_or_add_rPr(),'sz',val=size*2);element(s.element.get_or_add_rPr(),'color',val='000000');s.paragraph_format.line_spacing=1;s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER if lvl==1 else WD_ALIGN_PARAGRAPH.LEFT
  element(s.element.get_or_add_pPr(),'spacing',before=0,after=0,beforeLines=50,afterLines=50,beforeAutospacing=0,afterAutospacing=0);element(s.element.get_or_add_pPr(),'ind',firstLineChars=0,firstLine=0)
 title=d.styles['Title'];fonts(title.element.get_or_add_rPr(),cn='黑体');element(title.element.get_or_add_rPr(),'sz',val=36);title.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
 cap=d.styles['Caption'];fonts(cap.element.get_or_add_rPr(),cn='宋体',en='宋体');element(cap.element.get_or_add_rPr(),'sz',val=21);cap.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;cap.paragraph_format.line_spacing=1;element(cap.element.get_or_add_pPr(),'ind',firstLineChars=0,firstLine=0)
 def special_title(value,en=False):
  p=d.add_paragraph(value);r=p.runs[0];fonts(r._r.get_or_add_rPr(),cn='黑体',en='Times New Roman');r.font.size=Pt(18);r.font.bold=False;p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1;element(p._p.get_or_add_pPr(),'ind',firstLineChars=0,firstLine=0);element(p._p.get_or_add_pPr(),'spacing',before=0,after=0,beforeLines=50,afterLines=50,beforeAutospacing=0,afterAutospacing=0);return p
 d.add_paragraph('论文有格单文档软件演示样本','Title')
 d.add_paragraph('脱敏合成内容；非正式论文、学校规范或真实参考文献。')
 p=special_title('摘 要')
 if error:
  e=p._p.get_or_add_pPr().find(qn('w:spacing'));e.attrib.pop(qn('w:beforeLines'));e.set(qn('w:before'),'120');p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
 p=d.add_paragraph();p.add_run(CN_TEXT[:225]);r=p.add_run(CN_TEXT[225:]);
 if error:r.font.size=Pt(14)
 p=d.add_paragraph();p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.line_spacing=1;element(p._p.get_or_add_pPr(),'ind',firstLineChars=0)
 for value in ('关','键词：','合成文档；格式属性；人工确认'):
  r=p.add_run(value);fonts(r._r.get_or_add_rPr(),cn='黑体' if value!='合成文档；格式属性；人工确认' and not error else '宋体');r.font.size=Pt(12)
 if error:p.runs[-1].font.size=Pt(14)
 special_title('Abstract',True)
 p=d.add_paragraph('This software demonstration checks explicit formatting attributes without rewriting academic content. '+('Synthetic examples keep text, drawings and tables while documenting the selected changes. '*5));p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT if error else WD_ALIGN_PARAGRAPH.JUSTIFY
 p=d.add_paragraph();p.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.line_spacing=1;element(p._p.get_or_add_pPr(),'ind',firstLineChars=0)
 if error:
  r=p.add_run('Keywords: synthetic document; paragraph format; human review');fonts(r._r.get_or_add_rPr());r.font.size=Pt(12);r.font.bold=False
 else:
  for value,b in [('Keywords: ',True),('synthetic document; paragraph format; human review',False)]:
   r=p.add_run(value);fonts(r._r.get_or_add_rPr());r.font.size=Pt(12);r.font.bold=b
 p=d.add_heading('第一章 软件演示',1)
 if error:
  e=p._p.get_or_add_pPr().find(qn('w:spacing'))
  if e is None:e=element(p._p.get_or_add_pPr(),'spacing',before=0,beforeLines=25,after=0)
  e.set(qn('w:beforeLines'),'25')
 d.add_paragraph('本文使用合成内容验证正文规则；正文提及摘要不应自动变成摘要标题。')
 d.add_heading('1.1 普通段落属性',2)
 d.add_paragraph('1.5倍行距是一个格式术语，而不是标题；类型建议仍需用户核对。')
 d.add_heading('1.1.1 安全修正',3)
 d.add_paragraph('2026年是软件演示中的数字句子，未选中的格式和所有正文均应保留。')
 d.add_paragraph().add_run().add_picture(BytesIO(png()),width=Mm(55))
 d.add_paragraph('图1-1 软件演示流程','Caption');d.add_paragraph('表1-1 软件演示项目','Caption')
 t=d.add_table(rows=2,cols=2);t.style='Table Grid'
 for row,values in zip(t.rows,[('项目','说明'),('格式检查','保持文本和表格结构')]):
  for cell,value in zip(row.cells,values):cell.text=value
 special_title('参考文献')
 style=d.styles.add_style('DemoReference',WD_STYLE_TYPE.PARAGRAPH);style.base_style=normal
 # Explicit source-independent standard answer: Chinese12pt /1.5, no imposed indent.
 fonts(style.element.get_or_add_rPr(),cn='黑体' if error else '宋体');element(style.element.get_or_add_rPr(),'sz',val=20 if error else 24)
 style.paragraph_format.line_spacing=Pt(20) if error else 1.5
 for i in range(1,11):
  value=f'[{i}] 软件演示资料. 合成格式核对说明（虚构条目{i}）. 2026.' if i<9 else f'[{i}] Software demonstration. Synthetic reference {i} (fictional). 2026.'
  d.add_paragraph(value,'DemoReference' if i<=3 else 'Normal')
 return d

if __name__=='__main__':
 for name,error in [('academic-correct',False),('academic-errors',True)]:
  path=ROOT/(name+'.docx')
  if path.exists():raise SystemExit('拒绝覆盖样本')
  build(error).save(path)
  manifest={'input_sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'paragraph_types':{str(k):v for k,v in TYPES.items()},'confirmed':True,'confirmation':'合成样本作者独立核对；网页仍需用户确认'}
  path.with_suffix('.types.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
