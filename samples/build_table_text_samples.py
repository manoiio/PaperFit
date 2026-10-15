"""Independent synthetic table fixtures, no real thesis content; refuse overwrite."""
from pathlib import Path
import json,hashlib
from docx import Document
from docx.shared import Mm,Pt
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from samples.build_samples import fonts,element
ROOT=Path(__file__).resolve().parent

def build(errors=False):
 d=Document(ROOT/'academic-correct.docx');d.core_properties.title='表格普通文字软件演示';d.core_properties.author='软件演示'
 def format_run(r,cn='宋体',size=10.5):fonts(r._r.get_or_add_rPr(),cn=cn,en='Arial');r.font.size=Pt(size)
 def table(rows,cols):
  t=d.add_table(rows,cols);t.autofit=False
  for c in t.columns:c.width=Mm(120/cols)
  element(t._tbl.tblPr,'tblW',type='dxa',w=round(120*1440/25.4))
  for i,row in enumerate(t.rows):
   for j,c in enumerate(row.cells):
    c.width=Mm(120/cols);p=c.paragraphs[0];p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=1;p.add_run(f'合成数据{i+1}-{j+1}')
  return t
 # Cover uses a true table object and deliberately different font/size; not content.
 cover=table(1,2);cover._tbl.getparent().remove(cover._tbl);d.paragraphs[0]._p.addnext(cover._tbl)
 for c in cover.rows[0].cells:
  for r in c.paragraphs[0].runs:format_run(r,'黑体',14)
 content=d.tables[1];content.autofit=False
 for row in content.rows:
  for c in row.cells:
   for p in c.paragraphs:
    for r in p.runs:format_run(r)
 p=content.rows[-1].cells[-1].paragraphs[0];late=p.add_run(' 后半片段');format_run(late,'黑体' if errors else '宋体',12 if errors else 10.5)
 p=content.rows[-1].cells[-1].add_paragraph();r=p.add_run('同一单元格第二段');format_run(r)
 st=d.styles.add_style('DemoTableText',WD_STYLE_TYPE.PARAGRAPH);fonts(st.element.get_or_add_rPr());element(st.element.get_or_add_rPr(),'sz',val=21)
 bad=d.styles.add_style('DemoTableWrong',WD_STYLE_TYPE.CHARACTER);fonts(bad.element.get_or_add_rPr(),cn='楷体');element(bad.element.get_or_add_rPr(),'sz',val=24)
 inherited=table(2,2)
 for row in inherited.rows:
  for c in row.cells:c.paragraphs[0].style=st
 if errors:inherited.rows[1].cells[1].paragraphs[0].runs[0].style=bad
 merged=table(2,2);merged.cell(0,0).merge(merged.cell(0,1))
 for row in merged._tbl.findall(qn('w:tr')):
  for cell in row.findall(qn('w:tc')):
   for p in cell.findall(qn('w:p')):
    for rr in p.findall(qn('w:r')):fonts(rr.get_or_add_rPr());element(rr.get_or_add_rPr(),'sz',val=21)
 r=merged.cell(0,0).paragraphs[0].add_run(' 合并格后半段');format_run(r,'宋体',14 if errors else 10.5)
 special=table(1,2)
 for row in special.rows:
  for c in row.cells:
   for r in c.paragraphs[0].runs:format_run(r)
 r=special.cell(0,0).paragraphs[0].add_run('2');format_run(r,'黑体',8);r.font.superscript=True
 math=OxmlElement('m:oMath');mr=OxmlElement('m:r');mt=OxmlElement('m:t');mt.text='x';mr.append(mt);math.append(mr);special.cell(0,1).paragraphs[0]._p.append(math)
 nested=table(1,1);inner=nested.cell(0,0).add_table(1,1);inner.cell(0,0).text='嵌套软件演示'
 conditional=table(2,2);style=d.styles.add_style('DemoConditional',WD_STYLE_TYPE.TABLE)
 c=OxmlElement('w:tblStylePr');c.set(qn('w:type'),'firstRow');rp=OxmlElement('w:rPr');fonts(rp,cn='黑体');element(rp,'sz',val=28);c.append(rp);style.element.append(c);conditional.style=style
 # A direct override is uniquely known despite conditional fonts; other cells remain review.
 for r in conditional.cell(1,1).paragraphs[0].runs:format_run(r)
 return d

def main():
 for err,name in [(False,'table-text-correct'),(True,'table-text-errors')]:
  out=ROOT/(name+'.docx')
  if out.exists():raise RuntimeError('Refuse fixture overwrite')
  build(err).save(out)
  base=json.loads((ROOT/'academic-correct.types.json').read_text());base['input_sha256']=hashlib.sha256(out.read_bytes()).hexdigest()
  roles={'1':'cover',**{str(i):'content' for i in range(2,8)}}
  base['table_scope']={'input_sha256':base['input_sha256'],'table_types':roles,'confirmed':True,'confirmation':'样本作者独立核对用途'}
  out.with_suffix('.types.json').write_text(json.dumps(base,ensure_ascii=False,indent=2))
 # Authored from injection locations, not checker results.
 gold={'description':'独立合成注入清单；只列确定错误，复杂/特殊对象另待复核','expected_problem_count':5,'errors':[{'table':2,'row':2,'cell':2,'cell_paragraph':1,'run':2,'property':'font_cn','expected':'宋体'},{'table':2,'row':2,'cell':2,'cell_paragraph':1,'run':2,'property':'font_size','expected':10.5},{'table':3,'row':2,'cell':2,'cell_paragraph':1,'run':1,'property':'font_cn','expected':'宋体'},{'table':3,'row':2,'cell':2,'cell_paragraph':1,'run':1,'property':'font_size','expected':10.5},{'table':4,'row':1,'cell':1,'cell_paragraph':1,'run':2,'property':'font_size','expected':10.5}]}
 p=ROOT/'table-text-expected-errors.json'
 if p.exists():raise RuntimeError('Refuse gold overwrite')
 p.write_text(json.dumps(gold,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
