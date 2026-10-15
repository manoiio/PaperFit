"""只生成脱敏合成样本；不读取真实课程论文；输出已存在时拒绝覆盖。"""
from pathlib import Path
import hashlib
import json
import struct
import zlib
from io import BytesIO
from docx import Document
from docx.shared import Mm, Pt
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent

def element(parent, name, **attrs):
    e = parent.find(qn('w:' + name))
    if e is None:
        e = OxmlElement('w:' + name)
        parent.append(e)
    for k, v in attrs.items():
        e.set(qn('w:' + k), str(v))
    return e

def fonts(rp, cn='宋体', en='Times New Roman'):
    f = element(rp, 'rFonts', eastAsia=cn, ascii=en, hAnsi=en)
    for a in ('eastAsiaTheme', 'asciiTheme', 'hAnsiTheme'):
        f.attrib.pop(qn('w:' + a), None)
    element(rp, 'sz', val=24)

def png():
    # Standard-library PNG of a simple blue process graphic; no AI/media download dependency.
    width, height = 360, 90
    rows = []
    for y in range(height):
        row = bytearray()
        for x in range(width):
            block = (15 <= x < 105 or 135 <= x < 225 or 255 <= x < 345) and 20 <= y < 70
            arrow = 43 <= y < 48 and (105 <= x < 135 or 225 <= x < 255)
            row.extend((40, 85, 140) if block or arrow else (240, 245, 250))
        rows.append(b'\0' + bytes(row))
    def chunk(kind, data):
        return struct.pack('>I', len(data)) + kind + data + struct.pack('>I', zlib.crc32(kind + data))
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', width, height, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(b''.join(rows))) + chunk(b'IEND', b'')

def build(messy=False):
    d = Document()
    d.core_properties.author = '课程软件演示'
    d.core_properties.last_modified_by = '论文有格'
    d.core_properties.title = '论文有格脱敏合成软件演示样本'
    normal = d.styles['Normal']
    fonts(normal.element.get_or_add_rPr())
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    element(normal.element.get_or_add_pPr(), 'ind', firstLineChars=200)
    inherited = d.styles.add_style('DemoBodyBase', WD_STYLE_TYPE.PARAGRAPH)
    inherited.base_style = normal
    fonts(inherited.element.get_or_add_rPr())
    if messy:
        element(inherited.element.get_or_add_rPr(), 'rFonts', eastAsia='黑体')
    child = d.styles.add_style('DemoBody', WD_STYLE_TYPE.PARAGRAPH)
    child.base_style = inherited
    char_base = d.styles.add_style('DemoCharacterBase', WD_STYLE_TYPE.CHARACTER)
    fonts(char_base.element.get_or_add_rPr())
    if messy:
        element(char_base.element.get_or_add_rPr(), 'rFonts', ascii='Arial')
        element(char_base.element.get_or_add_rPr(), 'sz', val=28)
    char_child = d.styles.add_style('DemoCharacter', WD_STYLE_TYPE.CHARACTER)
    char_child.base_style = char_base
    for section in d.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        for side in ('top', 'bottom', 'left', 'right'):
            setattr(section, side + '_margin', Mm(25))
    title = d.add_paragraph('论文有格脱敏合成软件演示样本', 'Title') # p1
    element(title._p.get_or_add_pPr(), 'ind', firstLineChars=0, firstLine=0)
    for name, size in [('Title', 20), ('Heading 1', 18), ('Heading 2', 15), ('Caption', 10.5)]:
        style = d.styles[name]
        pp = style.element.get_or_add_pPr()
        border = pp.find(qn('w:pBdr'))
        if border is not None:
            pp.remove(border)
        fonts(style.element.get_or_add_rPr())
        element(style.element.get_or_add_rPr(), 'sz', val=int(size * 2))
        color = element(style.element.get_or_add_rPr(), 'color', val='000000')
        for key in list(color.attrib):
            if key != qn('w:val'):
                del color.attrib[key]
        element(style.element.get_or_add_pPr(), 'ind', firstLineChars=0, firstLine=0)
    # Labels and table cells are demonstration objects, not confirmed body paragraphs.
    d.add_paragraph('本文件仅用于软件演示。作者、单位和测量内容均为虚构，不是实际课程论文。') # p2
    d.add_heading('第一章 格式检查流程', level=1) # p3
    p = d.add_paragraph() # p4
    p.add_run('本文以简短论文片段验证文档格式处理。')
    r = p.add_run('后半段文字需要逐片段检查，不能遗漏字号错误。')
    if messy:
        r.font.size = Pt(10.5)
    p.paragraph_format.line_spacing = 1.2 if messy else 1.5
    p = d.add_paragraph('本段通过父样式继承中文字体，验证样式中的设置能否被读取。', 'DemoBody') # p5
    p.paragraph_format.line_spacing = Pt(18) if messy else 1.5
    p = d.add_paragraph('字符缩进示例保留两个汉字，编号值 Sample 2026 使用独立西文字体。') # p6
    element(p._p.get_or_add_pPr(), 'ind', firstLineChars=100 if messy else 200)
    d.add_heading('1.1 图表与保留对象', level=2) # p7
    d.add_paragraph().add_run().add_picture(BytesIO(png()), width=Mm(105)) # p8
    d.add_paragraph('图1-1 本地处理流程示意（软件演示）', 'Caption') # p9
    d.add_paragraph('表1-1 合成验证项目（软件演示）', 'Caption') # p10
    table = d.add_table(rows=3, cols=2)
    table.style = 'Table Grid'
    for row, values in zip(table.rows, [('项目', '说明'), ('检查', '固定七条规则'), ('输出', '新文件与复查报告')]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    p = d.add_paragraph('点值缩进示例按统一的十二磅字号等效缩进两个汉字。') # p11
    # Override inherited character indentation: direct chars=0 would supersede point indentation,
    # so use a separate paragraph style without firstLineChars.
    point = d.styles.add_style('DemoPointBody', WD_STYLE_TYPE.PARAGRAPH)
    fonts(point.element.get_or_add_rPr())
    point.paragraph_format.line_spacing = 1.5
    point.paragraph_format.first_line_indent = Pt(36 if messy else 24)
    p.style = point
    d.add_section(WD_SECTION_START.NEW_PAGE) # p12 section break
    d.add_heading('第二章 合成结果说明', level=1) # p13
    p = d.add_paragraph() # p14
    p.add_run('本地处理后保留原有文字、图片和表格。 ')
    p.add_run('Demo validation 2026.').style = char_child
    d.add_paragraph('本文只展示已确认正文段落的格式检查，不声称覆盖所有论文对象。') # p15
    if messy:
        d.sections[0].page_width, d.sections[0].page_height = Mm(216), Mm(279)
        d.sections[0].left_margin = Mm(20)
        d.sections[1].bottom_margin = Mm(30)
    return d

if __name__ == '__main__':
    paths = [ROOT / 'correct.docx', ROOT / 'errors.docx', ROOT / 'correct.body.json', ROOT / 'errors.body.json']
    if any(p.exists() for p in paths):
        raise SystemExit('样本已存在；为保护已有文件，本脚本拒绝覆盖。')
    for name, messy in [('correct', False), ('errors', True)]:
        path = ROOT / (name + '.docx')
        build(messy).save(path)
        data = {'input_sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                'body_paragraphs': [4, 5, 6, 11, 14, 15],
                'confirmation': '合成样本作者人工确认；顶层段落1起始，包括空段和分节段；表格段落不编号。'}
        (ROOT / (name + '.body.json')).write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        print(path)
