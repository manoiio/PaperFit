"""重建第五阶段公开软件演示规范；只创建新文件，拒绝覆盖。"""
from pathlib import Path
from io import BytesIO
import json
from docx import Document
from docx.shared import Pt, Mm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from core.package import write_new

ROOT=Path(__file__).resolve().parent
A=['软件演示规范A','仅用于测试论文有格软件，不是学校正式规范。','纸张为A4竖向。','上下左右边距均为2.5cm。','正文中文字体为宋体。','正文西文字体为Times New Roman。','正文普通文字字号为小四号。','正文采用1.5倍行距。','正文首行缩进2字。','请双面打印。']
B=['软件演示规范B','仅用于测试论文有格软件，不是学校正式规范。','纸张为A4竖向。','上边距3cm。','下边距2cm。','左边距3.5cm。','右边距3cm。','正文中文字体为黑体。','正文西文字体为Arial。','正文普通文字字号为四号。','正文采用2倍行距。','正文首行缩进1字符。','标题采用固定20磅行距。','请双面打印。']
VALUES_A={'paper':'A4 portrait','margin_top':25,'margin_bottom':25,'margin_left':25,'margin_right':25,'font_cn':'宋体','font_western':'Times New Roman','font_size':12,'line_spacing':1.5,'first_line_indent':2}
VALUES_B={**VALUES_A,'margin_top':30,'margin_bottom':20,'margin_left':35,'margin_right':30,'font_cn':'黑体','font_western':'Arial','font_size':14,'line_spacing':2,'first_line_indent':1}

def build():
    write_new(ROOT/'norm-A.txt',('\n'.join(A)+'\n').encode())
    doc=Document()
    sec=doc.sections[0];sec.page_width=Mm(210);sec.page_height=Mm(297)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Mm(25)
    for name in ('Normal','Title'):
        style=doc.styles[name];style.font.name='Times New Roman';style.font.size=Pt(12 if name=='Normal' else 18)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
        style.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
    for style in doc.styles:
        for border in style.element.xpath('./w:pPr/w:pBdr'):border.getparent().remove(border)
    doc.add_paragraph(B[0],style='Title');doc.add_paragraph(B[1]);doc.add_paragraph(B[2])
    table=doc.add_table(rows=0,cols=2);table.style='Table Grid'
    for label,text in zip(['上边距','下边距','左边距','右边距','中文字体','西文字体','字号','倍数行距','首行缩进'],B[3:12]):
        cells=table.add_row().cells;cells[0].text=label;cells[1].text=text
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','bottom','left','right','insideH','insideV'):
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    table._tbl.tblPr.append(borders)
    for text in B[12:]:doc.add_paragraph(text)
    stream=BytesIO();doc.save(stream);write_new(ROOT/'norm-B.docx',stream.getvalue())
    write_new(ROOT/'norm-expected.json',(json.dumps({'description':'独立人工期望清单；不是引擎生成结果','A':VALUES_A,'B':VALUES_B,'unsupported':['请双面打印。','标题采用固定20磅行距。']},ensure_ascii=False,indent=2)+'\n').encode())

if __name__=='__main__':build()
